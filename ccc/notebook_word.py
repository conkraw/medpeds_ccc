"""The program's own Jupyter Word-formatting functions.

Formatting helpers below are copied from the supplied report notebook without
redesigning tables, spacing, colors, section order, or scoring labels. The
notebook's CSV discovery, batch loop, folder creation, and ZIP writes are not
part of this module. ``create_doc`` returns the Word document for Streamlit.
"""
import pandas as pd
import json
import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START
from difflib import SequenceMatcher
from docx.oxml.ns import qn
from docx.shared import RGBColor

TABLE_WIDTH = 7.6
MAIN_WIDTHS_4COL = [1.8, 2.0, 1.8, 2.0]


def get_initials(name):
    if not name:
        return "UNK"

    parts = re.split(r"[,\s]+", str(name).strip())
    parts = [p for p in parts if p]

    initials = "".join(p[0].upper() for p in parts)

    return initials


def clean_filename(text):
    text = str(text).strip()
    text = re.sub(r'[\\/*?:"<>|]', "", text)
    return text[:100] if text else "unknown_resident"


def add_page_break(doc):
    doc.add_page_break()


def has_value(row, field):
    return field in row.index and pd.notna(row[field]) and str(row[field]).strip() != ""


def get_val(row, field, default=""):
    if has_value(row, field):
        return str(row[field]).strip()
    return default


def find_resident_target(row):
    review_order = [
        "r4eoy", "r4my",
        "r3eoy", "r3my",
        "r2eoy", "r2my",
        "r1eoy", "r1my",
    ]

    for target in review_order:
        if has_value(row, f"review_period_{target}"):
            return target

    return None


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=7, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]

    if align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    
    r = run._element
    r.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    r.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_font_color(cell, hex_color):
    """
    Changes font color in a cell after set_cell_text() has already run.
    This avoids needing color= inside set_cell_text().
    """
    rgb = RGBColor.from_string(hex_color)

    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = rgb


def set_table_widths(table, widths):
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)


def set_row_height(table, height=0.2):
    for row in table.rows:
        row.height = Inches(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def get_clinic_targets(target):
    mapping = {
        "r1my":  ["r1my"],
        "r1eoy": ["r1eoy"],
        "r2my":  ["r1eoy", "r2my"],
        "r2eoy": ["r1eoy", "r2eoy"],
        "r3my":  ["r1eoy", "r2eoy", "r3my"],
        "r3eoy": ["r1eoy", "r2eoy", "r3eoy"],
        "r4my":  ["r1eoy", "r2eoy", "r3eoy", "r4my"],
        "r4eoy": ["r1eoy", "r2eoy", "r3eoy", "r4eoy"],
    }
    return mapping.get(target, [])


def add_clinic_tally_table(doc, row, target):
    clinic_targets = get_clinic_targets(target)

    table = doc.add_table(rows=1 + len(clinic_targets), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    force_table_width(table, MAIN_WIDTHS_4COL)
    set_row_height(table, 0.2)

    header = table.cell(0, 0)
    for c in range(1, 4):
        header = header.merge(table.cell(0, c))

    shade_cell(header, "D9EAF7")
    set_cell_text(
        header,
        "ADULT AND PEDIATRIC CLINIC PATIENT TALLY",
        bold=True,
        size=7
    )

    for r, clinic_target in enumerate(clinic_targets, start=1):
        pgy = clinic_target[:2].upper()

        set_cell_text(table.cell(r, 0), f"ADULT PGY{pgy[-1]}", bold=True)
        set_cell_text(table.cell(r, 1), get_val(row, f"ad_clinic_num_{clinic_target}"))

        set_cell_text(table.cell(r, 2), f"PEDIATRIC PGY{pgy[-1]}", bold=True)
        set_cell_text(table.cell(r, 3), get_val(row, f"ped_clinic_num_{clinic_target}"))


def add_single_value_table(doc, title, value):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    table.cell(0, 0).width = Inches(7.6)
    table.cell(1, 0).width = Inches(7.6)

    table.rows[0].height = Inches(0.2)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    table.rows[1].height = Inches(0.3)
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    shade_cell(table.cell(0, 0), "D9EAF7")

    set_cell_text(
        table.cell(0, 0),
        title,
        bold=True,
        size=7
    )

    set_cell_text(
        table.cell(1, 0),
        value,
        size=7,
        align="left"
    )

    doc.add_paragraph()


def remove_subject_name(text, subject_name):
    if not subject_name or not text:
        return text

    cleaned = text

    # Remove full subject_name exactly
    cleaned = re.sub(re.escape(subject_name), "", cleaned, flags=re.IGNORECASE)

    # Try removing parts of name: "Last, First Middle"
    parts = re.split(r"[,\s]+", subject_name)
    parts = [p.strip() for p in parts if len(p.strip()) > 2]

    for part in parts:
        cleaned = re.sub(rf"\b{re.escape(part)}\b", "", cleaned, flags=re.IGNORECASE)

    return cleaned


def clean_eval_html(value, subject_name=""):
    if value is None or pd.isna(value):
        return ""

    value = str(value)

    # Convert HTML to text
    soup = BeautifulSoup(value, "html.parser")

    # Remove any text that is specifically a resident evaluator name
    # Example: Davis, Garrison Michael [Resident]
    text = soup.get_text("\n")

    text = re.sub(
        r"\b[A-Z][A-Za-z'\-]+,\s+[A-Z][A-Za-z'\-]+(?:\s+[A-Z][A-Za-z'\-]+)*\s+\[Resident\]",
        "",
        text
    )

    # Remove subject name from comments
    text = remove_subject_name(text, subject_name)

    # Clean extra blank lines/spaces
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()

    return text


def clean_eval_html_keep_formatting(html, subject_name=""):
    if html is None or pd.isna(html):
        return ""

    html = str(html)

    # Remove any evaluator names ending in [Resident]
    # Handles middle names, hyphenated names, apostrophes, periods, and suffix-ish initials
    html = re.sub(
        r"\b[A-Z][A-Za-z.'\-]+,\s+[A-Z][A-Za-z.'\-]+(?:\s+[A-Z](?:[A-Za-z.'\-]*)?)*\s*\[Resident\]",
        "[Resident]",
        html,
        flags=re.IGNORECASE
    )
    # Remove subject_name exactly
    if subject_name:
        html = re.sub(re.escape(subject_name), "", html, flags=re.IGNORECASE)

        # If subject_name is "Last, First Middle", also remove:
        # First, Last, Middle, First Last, Last First, etc.
        parts = re.split(r"[,\s]+", subject_name)
        parts = [p.strip() for p in parts if len(p.strip()) > 2]

        for part in parts:
            html = re.sub(rf"\b{re.escape(part)}\b", "", html, flags=re.IGNORECASE)

    # Clean extra spaces inside cells
    html = re.sub(r">\s+<", "><", html)
    html = re.sub(r"\s{2,}", " ", html)

    return html


def html_to_clean_text(html):
    if html is None or pd.isna(html):
        return ""

    soup = BeautifulSoup(str(html), "html.parser")

    for tag in soup.find_all(["h1", "h2", "h3", "h4"]):
        tag.insert_before("\n\n")
        tag.insert_after("\n")

    for tag in soup.find_all("li"):
        tag.insert_before("\n• ")

    text = soup.get_text(" ", strip=True)

    text = re.sub(r"\s*•\s*", "\n• ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)

    return text.strip()


def add_html_table_to_doc(
    doc,
    html,
    table_width=7.6,
    force_center=False,
    force_header_gray=False,
    force_header_bold=False
):
    
    soup = BeautifulSoup(html, "html.parser")
    html_table = soup.find("table")

    if html_table is None:
        text = soup.get_text("\n").strip()

        # create a blank table row if no data
        blank_table = doc.add_table(rows=1, cols=1)
        blank_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        blank_table.style = "Table Grid"
        blank_table.autofit = False
        blank_table.allow_autofit = False

        force_table_width(blank_table, [table_width])
        set_row_height(blank_table, 0.2)

        set_cell_text(
            blank_table.cell(0, 0),
            text if text else "",
            size=7,
            align="left"
        )

        return

    rows = html_table.find_all("tr")
    max_cols = 0

    parsed_rows = []
    for tr in rows:
        cells = tr.find_all(["td", "th"])
        parsed_cells = []
        col_count = 0

        for cell in cells:
            colspan = int(cell.get("colspan", 1))
            text = cell.get_text("\n").strip()
            if text == "":
                text = "XX"
                
            parsed_cells.append((text, colspan, cell))
            col_count += colspan

        parsed_rows.append(parsed_cells)
        max_cols = max(max_cols, col_count)

    word_table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    word_table.style = "Table Grid"
    word_table.autofit = False

    col_width = table_width / max_cols
    force_table_width(word_table, [col_width] * max_cols)
    
    set_row_height(word_table, 0.2)

    for r_idx, parsed_cells in enumerate(parsed_rows):
        c_idx = 0

        for text, colspan, html_cell in parsed_cells:
            word_cell = word_table.cell(r_idx, c_idx)

            if colspan > 1:
                for merge_idx in range(1, colspan):
                    word_cell = word_cell.merge(word_table.cell(r_idx, c_idx + merge_idx))

            style = html_cell.get("style", "")

            if "background-color" in style:
                if "#bfedd2" in style.lower():
                    shade_cell(word_cell, "BFEDD2")
                elif "#ffffff" in style.lower():
                    shade_cell(word_cell, "FFFFFF")
                else:
                    shade_cell(word_cell, "D9EAF7")

            align = "center" if force_center or "text-align: center" in style else "left"

            is_first_row = r_idx == 0

            is_header = (
                force_header_bold and is_first_row
            ) or (
                "MEDICINE EVALUATION SUMMARY" in text.upper()
                or "PEDIATRIC EVALUATION SUMMARY" in text.upper()
            )
            
            if force_header_gray and is_first_row:
                shade_cell(word_cell, "D9D9D9")
    
            set_cell_text(
                word_cell,
                text,
                bold=is_header,
                size=6,
                align=align
            )

            c_idx += colspan


def add_html_section_with_title(doc, title, html, table_width=7.6, header_gray=True):
    # Section title table
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.style = "Table Grid"
    title_table.autofit = False
    set_row_height(title_table, 0.2)
    force_table_width(title_table, [table_width])
    
    shade_cell(title_table.cell(0, 0), "D9EAF7")
    set_cell_text(title_table.cell(0, 0), title, bold=True, size=7, align="center")

    # HTML table
    add_html_table_to_doc(
        doc,
        html,
        table_width=table_width,
        force_center=True,
        force_header_gray=header_gray,
        force_header_bold=True
    )


def force_table_width(table, widths):
    table.autofit = False
    table.allow_autofit = False

    tbl = table._tbl
    tblPr = tbl.tblPr

    # Fixed table layout
    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")

    # Set total table width
    total_width = sum(widths)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(total_width * 1440)))
    tblW.set(qn("w:type"), "dxa")

    # Remove existing grid
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)

    # Add exact grid
    tblGrid = OxmlElement("w:tblGrid")
    for width in widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(int(width * 1440)))
        tblGrid.append(gridCol)
    tbl.insert(0, tblGrid)

    # Set each cell width
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue

            cell = row.cells[idx]
            cell.width = Inches(width)

            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)

            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def close_match(a, b, threshold=0.75):
    return SequenceMatcher(None, a.lower(), b.lower()).ratio() >= threshold


def scrub_dr_lastname_variants(text, subject_name):
    if not subject_name:
        return text

    name_parts = subject_name.strip().split()
    if not name_parts:
        return text

    last_name = name_parts[-1]

    def repl(match):
        possible_last = match.group(1)

        if close_match(possible_last, last_name):
            return "Dr."

        return match.group(0)

    return re.sub(
        r"\bDr\.?\s+([A-Z][A-Za-z.'\-]+)\b",
        repl,
        text
    )


def scrub_eval_text(text, subject_name=""):
    if not text:
        return "XX"

    text = str(text)

    # Replace resident evaluator names with [Resident]
    text = re.sub(
        r"\b[A-Z][A-Za-z.'\-]+,\s+[A-Z][A-Za-z.'\-]+(?:\s+[A-Z](?:[A-Za-z.'\-]*)?)*\s*\[Resident\]",
        "[Resident]",
        text,
        flags=re.IGNORECASE
    )

    if subject_name:
        name_parts = re.split(r"[,\s]+", subject_name.strip())
        name_parts = [p.strip() for p in name_parts if p.strip()]
        
        text = scrub_dr_lastname_variants(text, subject_name)

        if name_parts:
            last_name = name_parts[-1]

            # Remove Dr Lastname FIRST
            text = re.sub(
                rf"\bDr\.?\s*{re.escape(last_name)}\b",
                "Dr.",
                text,
                flags=re.IGNORECASE
            )

        # Remove exact full subject name
        text = re.sub(re.escape(subject_name), "", text, flags=re.IGNORECASE)

        # Remove individual name parts
        parts = [p for p in name_parts if len(p) > 2]

        for part in parts:
            text = re.sub(
                rf"\b{re.escape(part)}\b",
                "",
                text,
                flags=re.IGNORECASE
            )

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()

    return text if text else "XX"


def extract_comments_only_from_eval_html(html, subject_name=""):
    if html is None or pd.isna(html):
        return ["XX"]

    html = clean_eval_html_keep_formatting(html, subject_name)

    soup = BeautifulSoup(str(html), "html.parser")
    table = soup.find("table")

    if table is None:
        return [scrub_eval_text(soup.get_text("\n").strip(), subject_name)]

    comments = []

    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"])

        # Comment rows usually have one cell spanning all columns
        if len(cells) == 1:
            text = cells[0].get_text("\n").strip()
            upper = text.upper()

            if (
                text
                and "EVALUATION SUMMARY" not in upper
                and not re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", text)
            ):
                comments.append(scrub_eval_text(text, subject_name))

    return comments if comments else ["XX"]


def add_eval_comments_only_section(doc, title, comments):
    table = doc.add_table(rows=1 + len(comments), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    shade_cell(table.cell(0, 0), "D9EAF7")
    set_cell_text(table.cell(0, 0), title, bold=True, size=7, align="center")

    for i, comment in enumerate(comments, start=1):
        set_cell_text(
            table.cell(i, 0),
            comment if comment.strip() else "XX",
            bold=False,
            size=7,
            align="left"
        )

    force_table_width(table, [TABLE_WIDTH])
    doc.add_paragraph()


def add_observed_encounters_table(doc, row):
    # Build only encounter rows that have at least one logged X
    encounter_rows = []

    for encounter_num in range(1, 5):
        displays = []

        for pgy in range(1, 5):
            field = f"obenc{encounter_num}_{pgy}"
            value = get_val(row, field, default="")
            display = "X" if str(value).strip() == "1" else ""
            displays.append(display)

        # Keep row if any PGY has an X
        if any(displays):
            encounter_rows.append((encounter_num, displays))

    # If no encounters logged, keep one blank row for attention
    if not encounter_rows:
        encounter_rows = [(1, ["", "", "", ""])]

    table = doc.add_table(rows=2 + len(encounter_rows), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    # =========================================
    # TITLE ROW
    # =========================================

    header = table.cell(0, 0).merge(table.cell(0, 4))
    shade_cell(header, "D9EAF7")

    set_cell_text(
        header,
        "OBSERVED ENCOUNTERS COMPLETION",
        bold=True,
        size=7
    )

    # =========================================
    # COLUMN HEADERS
    # =========================================

    headers = ["NUMBER", "PGY1", "PGY2", "PGY3", "PGY4"]

    for col_idx, text in enumerate(headers):
        shade_cell(table.cell(1, col_idx), "EDEDED")
        set_cell_text(
            table.cell(1, col_idx),
            text,
            bold=True,
            size=7
        )

    # =========================================
    # DATA ROWS
    # =========================================

    for r_idx, (encounter_num, displays) in enumerate(encounter_rows, start=2):
        set_cell_text(
            table.cell(r_idx, 0),
            str(encounter_num),
            bold=True,
            size=7
        )

        for pgy_idx, display in enumerate(displays, start=1):
            set_cell_text(
                table.cell(r_idx, pgy_idx),
                display,
                bold=True,
                size=8
            )

    force_table_width(table, [1.5, 1.525, 1.525, 1.525, 1.525])


def add_html_summary_section(doc, title, html, fill_color="D9EAF7"):
    # =========================================
    # TITLE
    # =========================================

    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_table.style = "Table Grid"
    title_table.autofit = False
    title_table.allow_autofit = False

    set_row_height(title_table, 0.2)
    force_table_width(title_table, [TABLE_WIDTH])

    shade_cell(title_table.cell(0, 0), fill_color)

    set_cell_text(
        title_table.cell(0, 0),
        title,
        bold=True,
        size=7,
        align="center"
    )

    # =========================================
    # HTML CONTENT
    # =========================================

    soup = BeautifulSoup(str(html), "html.parser")

    content_table = doc.add_table(rows=1, cols=1)
    content_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    content_table.style = "Table Grid"
    content_table.autofit = False
    content_table.allow_autofit = False

    force_table_width(content_table, [TABLE_WIDTH])

    cell = content_table.cell(0, 0)

    # clear default paragraph
    cell.text = ""

    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        text = element.get_text(" ", strip=True)

        if not text:
            continue

        p = cell.add_paragraph()

        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1

        if element.name in ["h1", "h2", "h3", "h4"]:
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(8)

        elif element.name == "li":
            p.style = "List Bullet"

            run = p.add_run(text)
            run.font.size = Pt(7)

        else:
            run = p.add_run(text)
            run.font.size = Pt(7)


def target_to_pgy_period(target):
    # r3eoy -> 3, y
    # r3my  -> 3, mm
    pgy = int(target[1])
    period = "mm" if target.endswith("my") else "y"
    return pgy, period


def get_milestone_periods_for_target(target):
    current_pgy, current_period = target_to_pgy_period(target)

    periods = []

    for pgy in range(1, current_pgy + 1):
        if pgy < current_pgy:
            periods.append((pgy, "mm"))
            periods.append((pgy, "y"))
        else:
            if current_period == "mm":
                periods.append((pgy, "mm"))
            else:
                periods.append((pgy, "mm"))
                periods.append((pgy, "y"))

    return periods


def add_competency_by_training_year_table(doc, row, target):
    milestone_cols = [
        ("PC",  ["pc1", "pc2", "pc3", "pc4", "pc5", "pc6"]),
        ("MK",  ["mk1", "mk2", "mk3"]),
        ("SBP", ["sbp1", "sbp2", "sbp3"]),
        ("PBLI", ["pbli1", "pbli2"]),
        ("PRO", ["pro1", "pro2", "pro3", "pro4"]),
        ("ICS", ["ics1", "ics2", "ics3"]),
    ]

    flat_cols = []
    for group, cols in milestone_cols:
        for col in cols:
            flat_cols.append((group, col))

    periods = get_milestone_periods_for_target(target)

    # Build only rows with at least one value
    data_rows = []

    for pgy, period in periods:
        values = []

        for _, base in flat_cols:
            field = f"{base}{period}_{pgy}"
            values.append(get_val(row, field, default=""))

        if any(str(v).strip() != "" for v in values):
            label = f"PGY{pgy} {'Mid-Year' if period == 'mm' else 'End-of-Year'}"
            data_rows.append((label, values))

    if not data_rows:
        return

    table = doc.add_table(rows=3 + len(data_rows), cols=1 + len(flat_cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    total_cols = 1 + len(flat_cols)

    # Main title
    header = table.cell(0, 0).merge(table.cell(0, total_cols - 1))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, "MEDICINE MILESTONES", bold=True, size=7)

    # Group headers
    set_cell_text(table.cell(1, 0), "")
    shade_cell(table.cell(1, 0), "FFFFFF")

    col_idx = 1
    for group, cols in milestone_cols:
        start = col_idx
        end = col_idx + len(cols) - 1

        group_cell = table.cell(1, start).merge(table.cell(1, end))
        shade_cell(group_cell, "EDEDED")
        set_cell_text(group_cell, group, bold=True, size=6)

        col_idx = end + 1

    # Number headers
    set_cell_text(table.cell(2, 0), "")

    col_idx = 1
    for group, cols in milestone_cols:
        for i in range(1, len(cols) + 1):
            shade_cell(table.cell(2, col_idx), "000000")
            set_cell_text(table.cell(2, col_idx), str(i), bold=True, size=6)
            col_idx += 1

    # Data rows
    for r_idx, (label, values) in enumerate(data_rows, start=3):
        set_cell_text(table.cell(r_idx, 0), label, bold=True, size=6)

        for c_idx, value in enumerate(values, start=1):
            display = str(value).strip()
            set_cell_text(table.cell(r_idx, c_idx), display, bold=True, size=6)

    # 22 columns total: label + 21 milestone cells
    first_col = 0.75
    remaining = TABLE_WIDTH - first_col
    each_col = remaining / len(flat_cols)

    force_table_width(table, [first_col] + [each_col] * len(flat_cols))

    doc.add_paragraph()


def safe_float(value):
    try:
        if value is None:
            return None
        value = str(value).strip()
        if value == "":
            return None
        return float(value)
    except:
        return None


def get_milestone_config(domain):
    if domain == "medicine":
        title = "MEDICINE MILESTONE TREND SUMMARY"
        milestone_cols = [
            ("PC", ["pc1", "pc2", "pc3", "pc4", "pc5", "pc6"]),
            ("MK", ["mk1", "mk2", "mk3"]),
            ("SBP", ["sbp1", "sbp2", "sbp3"]),
            ("PBLI", ["pbli1", "pbli2"]),
            ("PRO", ["pro1", "pro2", "pro3", "pro4"]),
            ("ICS", ["ics1", "ics2", "ics3"]),
        ]
    else:
        title = "PEDIATRIC MILESTONE TREND SUMMARY"
        milestone_cols = [
            ("PC", ["pc1", "pc2", "pc3", "pc4", "pc5"]),
            ("MK", ["mk1", "mk2"]),
            ("SBP", ["sbp1", "sbp2", "sbp3", "sbp4", "sbp5", "sbp6"]),
            ("PBLI", ["pbli1", "pbli2"]),
            ("PRO", ["pro1", "pro2", "pro3", "pro4"]),
            ("ICS", ["ics1", "ics2", "ics3"]),
        ]

    return title, milestone_cols


def classify_milestone_trend(final_score, change):
    if final_score is None or change is None:
        return ""

    if final_score < 3.75 or change < 0:
        return "Needs review"
    elif change >= 0.50:
        return "Strong upward progression"
    elif change >= 0.20:
        return "Modest upward progression"
    elif final_score >= 4.50 and abs(change) < 0.20:
        return "Stable advanced performance"
    elif 3.75 <= final_score < 4.50 and abs(change) < 0.20:
        return "Stable expected performance"
    else:
        return "Needs review"


def average_domain_score(row, bases, suffix, pgy):
    scores = []

    for base in bases:
        field = f"{base}{suffix}_{pgy}"
        value = safe_float(get_val(row, field, ""))

        if value is not None:
            scores.append(value)

    if not scores:
        return None

    return round(sum(scores) / len(scores), 2)


def add_milestone_trend_summary_table(doc, row, target, domain):
    title, milestone_cols = get_milestone_config(domain)

    periods = get_milestone_rows_for_target(target, domain)

    # Keep rows that are actually present
    available_periods = []

    for pgy, suffix, label in periods:
        has_any_value = False

        for _, bases in milestone_cols:
            if average_domain_score(row, bases, suffix, pgy) is not None:
                has_any_value = True
                break

        if has_any_value:
            available_periods.append((pgy, suffix, label))

    if len(available_periods) < 2:
        return

    trend_rows = []

    for domain_name, bases in milestone_cols:
        scores = []

        for pgy, suffix, label in available_periods:
            avg = average_domain_score(row, bases, suffix, pgy)
            scores.append(avg)

        nonblank_scores = [s for s in scores if s is not None]

        if len(nonblank_scores) < 2:
            continue

        first_score = nonblank_scores[0]
        final_score = nonblank_scores[-1]
        change = round(final_score - first_score, 2)

        trend = classify_milestone_trend(final_score, change)

        trend_rows.append((domain_name, scores, change, trend))

    if not trend_rows:
        return

    table = doc.add_table(
        rows=2 + len(trend_rows),
        cols=1 + len(available_periods) + 2
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    total_cols = 1 + len(available_periods) + 2

    header = table.cell(0, 0).merge(table.cell(0, total_cols - 1))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, title, bold=True, size=7)

    headers = ["Domain"] + [label for _, _, label in available_periods] + ["Change", "Trend"]

    for c, h in enumerate(headers):
        shade_cell(table.cell(1, c), "EDEDED")
        set_cell_text(table.cell(1, c), h, bold=True, size=6)
        
    for r, (domain_name, scores, change, trend) in enumerate(trend_rows, start=2):
        set_cell_text(table.cell(r, 0), domain_name, bold=True, size=6)

        for c, score in enumerate(scores, start=1):
            display = "" if score is None else f"{score:.2f}"
            set_cell_text(table.cell(r, c), display, bold=True, size=6, align="center")

        change_display = f"{change:+.2f}"

        change_cell = table.cell(r, 1 + len(available_periods))
        trend_cell = table.cell(r, 2 + len(available_periods))

        set_cell_text(change_cell, change_display, bold=True, size=6, align="center")
        set_cell_text(trend_cell, trend, bold=False, size=6, align="center")

        if "needs review" in str(trend).lower():
            shade_cell(trend_cell, "F4CCCC")

    first_col = 0.75
    change_col = 0.65
    trend_col = 1.55
    remaining = TABLE_WIDTH - first_col - change_col - trend_col
    period_width = remaining / len(available_periods)

    widths = [first_col] + [period_width] * len(available_periods) + [change_col, trend_col]
    force_table_width(table, widths)


def add_duty_hour_deep_dive_table(doc, row, target):

    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    # Same alignment/width idea you already had
    force_table_width(table, [1.7, 2.95, 2.95])
    set_row_height(table, 0.2)

    # ------------------------------------------------
    # Title
    # ------------------------------------------------

    header = table.cell(0, 0)
    for c in range(1, 3):
        header = header.merge(table.cell(0, c))

    shade_cell(header, "D9EAF7")
    set_cell_text(
        header,
        "DUTY HOUR DEEP DIVE",
        bold=True,
        size=7
    )

    # ------------------------------------------------
    # Last duty hour log
    # ------------------------------------------------

    shade_cell(table.cell(1, 0), "EDEDED")
    set_cell_text(
        table.cell(1, 0),
        "LAST DUTY HOUR LOG:",
        bold=True,
        size=6
    )

    merged = table.cell(1, 1)
    merged = merged.merge(table.cell(1, 2))

    set_cell_text(
        merged,
        get_val(row, f"dh_last_log_{target}", default=""),
        size=6
    )

    # ------------------------------------------------
    # Minimum days off header row
    # ------------------------------------------------

    for c in range(3):
        shade_cell(table.cell(2, c), "000000")

    set_cell_text(
        table.cell(2, 0),
        "MINIMUM DAYS OFF",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(2, 0), "FFFFFF")

    set_cell_text(
        table.cell(2, 1),
        "DAY OFF WINDOW",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(2, 1), "FFFFFF")

    set_cell_text(
        table.cell(2, 2),
        "ROTATIONS",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(2, 2), "FFFFFF")

    # ------------------------------------------------
    # Minimum days off values
    # ------------------------------------------------

    set_cell_text(
        table.cell(3, 0),
        get_val(row, f"dh_min_days_off_{target}", default=""),
        size=6
    )

    set_cell_text(
        table.cell(3, 1),
        get_val(row, f"dayoff_window_{target}", default=""),
        size=6
    )

    set_cell_text(
        table.cell(3, 2),
        get_val(row, f"dayoff_rotations_{target}", default=""),
        size=6
    )

    # ------------------------------------------------
    # 80-hour header row
    # ------------------------------------------------

    for c in range(3):
        shade_cell(table.cell(4, c), "000000")

    set_cell_text(
        table.cell(4, 0),
        "MAX HOURS IN 4 WEEK PERIOD",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(4, 0), "FFFFFF")

    set_cell_text(
        table.cell(4, 1),
        "80 HOUR WINDOW",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(4, 1), "FFFFFF")

    set_cell_text(
        table.cell(4, 2),
        "ROTATIONS",
        bold=True,
        size=6
    )
    set_cell_font_color(table.cell(4, 2), "FFFFFF")

    # ------------------------------------------------
    # 80-hour values
    # ------------------------------------------------

    set_cell_text(
        table.cell(5, 0),
        get_val(row, f"eightyhr_max_{target}", default=""),
        size=6
    )

    set_cell_text(
        table.cell(5, 1),
        get_val(row, f"eightyhr_window_{target}", default=""),
        size=6
    )

    set_cell_text(
        table.cell(5, 2),
        get_val(row, f"eightyhr_rotations_{target}", default=""),
        size=6
    )


def get_training_year_from_target(target):
    target = str(target).lower()

    if "r1" in target:
        return 1
    elif "r2" in target:
        return 2
    elif "r3" in target:
        return 3
    elif "r4" in target:
        return 4
    return None


def readiness_label(score, change=0, target=None):
    pgy = get_training_year_from_target(target)

    if score is None:
        return "Developing with Additional Observation Needed"

    # PGY1: developmental framing
    if pgy == 1:
        if score >= 4.2:
            return "Advanced for Early Training"
        elif score >= 3.5:
            return "Progressing Strongly for Level of Training"
        elif score >= 2.8:
            return "Developing with Increasing Independence"
        else:
            return "Early Development with Ongoing Coaching"

    # PGY2
    if pgy == 2:
        if score >= 4.4:
            return "Advanced for Level of Training"
        elif score >= 3.8:
            return "Progressing Toward Practice Readiness"
        elif score >= 3.2:
            return "Developing with Increasing Consistency"
        else:
            return "Developing with Focused Coaching"

    # PGY3
    if pgy == 3:
        if score >= 4.5:
            return "Advanced"
        elif score >= 4.0:
            return "Demonstrating Practice Readiness in Many Settings"
        elif score >= 3.5:
            return "Approaching Practice Readiness with Continued Growth"
        else:
            return "Developing Toward Practice Readiness"

    # PGY4
    if pgy == 4:
        if score >= 4.6:
            return "Exemplary"
        elif score >= 4.3:
            return "Advanced"
        elif score >= 4.0:
            return "Ready for Independent Practice"
        elif score >= 3.6:
            return "Near Practice Readiness with Continued Growth"
        else:
            return "Progressing Toward Practice Readiness with Focused Coaching"

    # fallback
    if score >= 4.6:
        return "Exemplary"
    elif score >= 4.3:
        return "Advanced"
    elif score >= 4.0:
        return "Demonstrating Practice Readiness"
    elif score >= 3.5:
        return "Progressing Toward Practice Readiness"
    else:
        return "Developing with Continued Coaching"


def latest_domain_score(row, target, domain, domain_name):
    _, milestone_cols = get_milestone_config(domain)
    domain_map = dict(milestone_cols)

    if domain_name not in domain_map:
        return None, None

    periods = get_milestone_rows_for_target(target, domain)

    scores = []
    for pgy, suffix, label in periods:
        avg = average_domain_score(row, domain_map[domain_name], suffix, pgy)
        if avg is not None:
            scores.append(avg)

    if not scores:
        return None, None

    final_score = scores[-1]
    change = round(scores[-1] - scores[0], 2) if len(scores) >= 2 else 0

    return final_score, change


def combined_score(row, target, domain_names):
    scores = []
    changes = []

    for domain, milestone_domain in domain_names:
        score, change = latest_domain_score(row, target, domain, milestone_domain)

        if score is not None:
            scores.append(score)

        if change is not None:
            changes.append(change)

    if not scores:
        return None, None

    final_score = round(sum(scores) / len(scores), 2)
    avg_change = round(sum(changes) / len(changes), 2) if changes else 0

    return final_score, avg_change


def is_low_readiness_text(text):
    text = str(text).lower()

    low_phrases = [
        "early development",
        "focused coaching",
        "ongoing coaching",
        "additional observation",
        "developing toward practice readiness",
        "near practice readiness"
    ]

    return any(phrase in text for phrase in low_phrases)


def add_practice_readiness_snapshot(doc, row, target):
    readiness_domains = [
        (
            "Clinical Reasoning",
            [("medicine", "PC"), ("medicine", "MK"), ("pediatrics", "PC"), ("pediatrics", "MK")]
        ),
        (
            "Team Leadership",
            [("medicine", "SBP"), ("medicine", "ICS"), ("pediatrics", "SBP"), ("pediatrics", "ICS")]
        ),
        (
            "Learner Supervision",
            [("medicine", "PBLI"), ("medicine", "ICS"), ("pediatrics", "PBLI"), ("pediatrics", "ICS")]
        ),
        (
            "Communication",
            [("medicine", "ICS"), ("pediatrics", "ICS")]
        ),
        (
            "Professionalism",
            [("medicine", "PRO"), ("pediatrics", "PRO")]
        ),
        (
            "Systems-Based Practice",
            [("medicine", "SBP"), ("pediatrics", "SBP")]
        ),
    ]

    rows = []

    for label, mapped_domains in readiness_domains:
        score, change = combined_score(row, target, mapped_domains)
        assessment = readiness_label(score, change, target)

        if score is not None and change is not None:
            display = f"{assessment} ({score:.2f}, {change:+.2f})"
        else:
            display = assessment

        rows.append((label, display))

    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    header = table.cell(0, 0).merge(table.cell(0, 1))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, "PRACTICE READINESS AND DEVELOPMENT SNAPSHOT", bold=True, size=7)

    for r, (domain_label, assessment) in enumerate(rows, start=1):
        domain_cell = table.cell(r, 0)
        assessment_cell = table.cell(r, 1)

        shade_cell(domain_cell, "EDEDED")
        set_cell_text(domain_cell, domain_label, bold=True, size=7, align="center")

        set_cell_text(assessment_cell, assessment, bold=False, size=7, align="center")

        if is_low_readiness_text(assessment):
            shade_cell(assessment_cell, "F4CCCC")

        force_table_width(table, [2.2, 5.4])


def add_intraining_scores_table(doc, row):
    # =========================================
    # BUILD ONLY NON-BLANK ROWS
    # =========================================

    score_rows = []

    for r in range(1, 5):
        med_score = get_val(row, f"med_intran_r{r}", "")
        med_pct = get_val(row, f"med_intranp_r{r}", "")

        ped_score = get_val(row, f"peds_intran_r{r}", "")
        ped_pct = get_val(row, f"peds_intranp_r{r}", "")

        med_text = (
            f"{med_score} ({med_pct}%tile)"
            if med_score or med_pct else ""
        )

        ped_text = (
            f"{ped_score} ({ped_pct}%_PR)"
            if ped_score or ped_pct else ""
        )

        # only keep rows with at least one score
        if med_text or ped_text:
            score_rows.append((r, med_text, ped_text))

    # keep one blank row if absolutely none exist
    if not score_rows:
        score_rows = [(1, "", "")]

    # =========================================
    # CREATE TABLE
    # =========================================

    table = doc.add_table(rows=2 + len(score_rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    force_table_width(table, [0.55, 3.525, 3.525])
    set_row_height(table, 0.2)

    # =========================================
    # TITLE ROW
    # =========================================

    header = table.cell(0, 0).merge(table.cell(0, 2))

    shade_cell(header, "D9EAF7")

    set_cell_text(
        header,
        "IN-TRAINING EXAM SCORES",
        bold=True,
        size=7
    )

    # =========================================
    # COLUMN HEADERS
    # =========================================

    shade_cell(table.cell(1, 0), "EDEDED")
    shade_cell(table.cell(1, 1), "EDEDED")
    shade_cell(table.cell(1, 2), "EDEDED")

    set_cell_text(table.cell(1, 0), "")
    set_cell_text(
        table.cell(1, 1),
        "MED IN-TRAINING EXAM SCORE:",
        bold=True,
        size=7
    )

    set_cell_text(
        table.cell(1, 2),
        "PED IN-TRAINING EXAM SCORE:",
        bold=True,
        size=7
    )

    # =========================================
    # DATA ROWS
    # =========================================

    for idx, (pgy, med_text, ped_text) in enumerate(score_rows, start=2):
        set_cell_text(
            table.cell(idx, 0),
            f"PGY{pgy}",
            bold=True,
            size=7
        )

        set_cell_text(
            table.cell(idx, 1),
            med_text,
            bold=True,
            size=7
        )

        set_cell_text(
            table.cell(idx, 2),
            ped_text,
            bold=True,
            size=7
        )


def add_usmle_scores_table(doc, row):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    force_table_width(table, [3.8, 3.8])
    set_row_height(table, 0.2)

    header = table.cell(0, 0).merge(table.cell(0, 1))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, "USMLE SCORES", bold=True, size=7)

    rows = [
        ("STEP 1:", "step1"),
        ("STEP 2 CK:", "step2ck"),
        ("STEP 3:", "step3"),
    ]

    for r_idx, (label, field) in enumerate(rows, start=1):
        set_cell_text(table.cell(r_idx, 0), label, bold=True, size=7)
        set_cell_text(table.cell(r_idx, 1), get_val(row, field), bold=True, size=7)


def set_cell_text_preserve_paragraphs(cell, text, size=7, bold=False):
    cell.text = ""

    text = "" if pd.isna(text) else str(text).strip()
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    if not paragraphs:
        paragraphs = [""]

    for i, para_text in enumerate(paragraphs):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()

        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(6)

        run = p.add_run(para_text)
        run.bold = bold
        run.font.size = Pt(size)

        # ADD THIS
        run.font.name = "Arial"


def add_simple_summary_table(doc, title, value, fill_color="D9EAF7"):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    set_cell_text(
        table.cell(0, 0),
        title,
        bold=True,
        size=7,
        align="center"
    )

    shade_cell(table.cell(0, 0), fill_color)

    set_cell_text_preserve_paragraphs(
        table.cell(1, 0),
        value,
        size=7,
        bold=False
    )
    force_table_width(table, [TABLE_WIDTH])


def get_cumulative_targets(target):
    mapping = {
        "r1my":  ["r1my"],
        "r1eoy": ["r1my", "r1eoy"],
        "r2my":  ["r1my", "r1eoy", "r2my"],
        "r2eoy": ["r1my", "r1eoy", "r2my", "r2eoy"],
        "r3my":  ["r1my", "r1eoy", "r2my", "r2eoy", "r3my"],
        "r3eoy": ["r1my", "r1eoy", "r2my", "r2eoy", "r3my", "r3eoy"],
        "r4my":  ["r1my", "r1eoy", "r2my", "r2eoy", "r3my", "r3eoy", "r4my"],
        "r4eoy": ["r1my", "r1eoy", "r2my", "r2eoy", "r3my", "r3eoy", "r4my", "r4eoy"],
    }
    return mapping.get(target, [])


def add_scholarly_activities_table(doc, row, target):
    targets = get_cumulative_targets(target)

    abstract_values = [
        get_val(row, f"abstract_{t}", "").strip()
        for t in targets
        if get_val(row, f"abstract_{t}", "").strip()
    ]

    publication_values = [
        get_val(row, f"publications_{t}", "").strip()
        for t in targets
        if get_val(row, f"publications_{t}", "").strip()
    ]

    # one blank row if none exist
    if not abstract_values:
        abstract_values = [""]

    if not publication_values:
        publication_values = [""]

    rows_needed = 1 + 1 + len(abstract_values) + 1 + len(publication_values)

    table = doc.add_table(rows=rows_needed, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    r = 0

    shade_cell(table.cell(r, 0), "D9EAF7")
    set_cell_text(table.cell(r, 0), "SCHOLARLY ACTIVITIES", bold=True, size=7)
    r += 1

    shade_cell(table.cell(r, 0), "EDEDED")
    set_cell_text(table.cell(r, 0), "ABSTRACT SUBMISSIONS/ACCEPTANCES", bold=True, size=7)
    r += 1

    for value in abstract_values:
        set_cell_text(table.cell(r, 0), value, bold=True, size=7, align="center")
        r += 1

    shade_cell(table.cell(r, 0), "EDEDED")
    set_cell_text(table.cell(r, 0), "PUBLICATIONS", bold=True, size=7)
    r += 1

    for value in publication_values:
        set_cell_text(table.cell(r, 0), value, bold=True, size=7, align="center")
        r += 1

    force_table_width(table, [TABLE_WIDTH])


def add_advocacy_table(doc, row, target):
    targets = get_cumulative_targets(target)

    advocacy_values = [
        get_val(row, f"advocacy_{t}", "").strip()
        for t in targets
        if get_val(row, f"advocacy_{t}", "").strip()
    ]

    # one blank row if none exist
    if not advocacy_values:
        advocacy_values = [""]

    rows_needed = 1 + len(advocacy_values)

    table = doc.add_table(rows=rows_needed, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    shade_cell(table.cell(0, 0), "D9EAF7")
    set_cell_text(table.cell(0, 0), "ADVOCACY", bold=True, size=7)

    for r, value in enumerate(advocacy_values, start=1):
        set_cell_text(table.cell(r, 0), value, bold=True, size=7, align="left")

    force_table_width(table, [TABLE_WIDTH])


def add_signature_section(doc):
    # =========================================
    # HEADER
    # =========================================

    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.style = "Table Grid"
    header_table.autofit = False
    header_table.allow_autofit = False

    set_row_height(header_table, 0.2)
    force_table_width(header_table, [TABLE_WIDTH])

    shade_cell(header_table.cell(0, 0), "D9EAF7")

    set_cell_text(
        header_table.cell(0, 0),
        "SIGNATURES",
        bold=True,
        size=7,
        align="center"
    )

    # =========================================
    # SIGNATURE TABLE
    # =========================================

    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    force_table_width(table, [2.6, 2.6, 2.4])

    signature_rows = [
        "Resident",
        "Program Director",
    ]

    for r_idx, role in enumerate(signature_rows):
        row = table.rows[r_idx]
        row.height = Inches(0.3)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

        # Role
        set_cell_text(
            table.cell(r_idx, 0),
            f"{role} Signature:",
            bold=True,
            size=7,
            align="left"
        )

        # Signature line
        set_cell_text(
            table.cell(r_idx, 1),
            "_____________________________________________",
            bold=False,
            size=7,
            align="left"
        )

        # Date line
        set_cell_text(
            table.cell(r_idx, 2),
            "Date: ____________________________________",
            bold=True,
            size=7,
            align="left"
        )

    doc.add_paragraph()


def add_summary_section(doc, row, target):
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    force_table_width(table, MAIN_WIDTHS_4COL)
    set_row_height(table, 0.2)

    header = table.cell(0, 0)
    for c in range(1, 4):
        header = header.merge(table.cell(0, c))

    shade_cell(header, "D9EAF7")
    
    review_label = "MID-YEAR" if target.endswith("my") else "END-OF-YEAR"
    
    
    
    set_cell_text(
        header,
        f"MEDICINE PEDIATRICS {review_label} SUMMARY",
        bold=True,
        size=7
    )

    rows = [
        ("RESIDENT NAME:", "subject_name", "TRAINING LEVEL:", f"subject_status_{target}"),
        ("COACH NAME:", "coach_name", "REVIEW PERIOD:", f"review_period_{target}"),
        ("CCC MEETING DATE:", f"ccc_meeting_date_{target}", "GRADUATION DATE:", "graduation_date"),
    ]

    for r, item in enumerate(rows, start=1):
        label1, field1, label2, field2 = item

        set_cell_text(table.cell(r, 0), label1, bold=True)
        set_cell_text(table.cell(r, 1), get_val(row, field1))
        set_cell_text(table.cell(r, 2), label2, bold=True)
        set_cell_text(table.cell(r, 3), get_val(row, field2))
        
    # =========================================
    # SPACE AFTER FIRST TABLE
    # =========================================

    #doc.add_paragraph()


    # =========================================
    # ROTATIONS TABLE
    # =========================================

    rotation_table = doc.add_table(rows=2, cols=1)
    rotation_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rotation_table.style = "Table Grid"
    rotation_table.autofit = False

    rotation_table.rows[0].height = Inches(0.2)
    rotation_table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    rotation_table.rows[1].height = Inches(0.4)
    rotation_table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    force_table_width(rotation_table, [TABLE_WIDTH])

    shade_cell(rotation_table.cell(0, 0), "D9EAF7")

    set_cell_text(
        rotation_table.cell(0, 0),
        "ROTATIONS SINCE LAST REVIEW",
        bold=True,
        size=7
    )

    rotation_field = f"rotations_{target}"

    set_cell_text(
        rotation_table.cell(1, 0),
        get_val(row, rotation_field),
        size=7,
        align="left"
    )


def get_milestone_rows_for_target(target, domain):
    current_pgy = int(target[1])
    is_midyear = target.endswith("my")

    if domain == "medicine":
        mid_suffix = "mm"
        end_suffix = "my"
    else:
        mid_suffix = "pm"
        end_suffix = "py"

    rows = []

    for pgy in range(1, current_pgy + 1):

        # Previous years: show end-of-year only
        if pgy < current_pgy:
            rows.append((pgy, end_suffix, f"PGY{pgy} EOY"))

        # Current year: show MY only if midyear
        elif is_midyear:
            rows.append((pgy, mid_suffix, f"PGY{pgy} MY"))

        # Current year: show both MY and EOY if end-of-year
        else:
            rows.append((pgy, mid_suffix, f"PGY{pgy} MY"))
            rows.append((pgy, end_suffix, f"PGY{pgy} EOY"))

    return rows


def add_milestones_table(doc, row, target, domain):
    if domain == "medicine":
        title = "MEDICINE MILESTONES"
        milestone_cols = [
            ("PC",   ["pc1", "pc2", "pc3", "pc4", "pc5", "pc6"]),
            ("MK",   ["mk1", "mk2", "mk3"]),
            ("SBP",  ["sbp1", "sbp2", "sbp3"]),
            ("PBLI", ["pbli1", "pbli2"]),
            ("PRO",  ["pro1", "pro2", "pro3", "pro4"]),
            ("ICS",  ["ics1", "ics2", "ics3"]),
        ]

    else:
        title = "PEDIATRIC MILESTONES"
        milestone_cols = [
            ("PC",   ["pc1", "pc2", "pc3", "pc4", "pc5"]),
            ("MK",   ["mk1", "mk2"]),
            ("SBP",  ["sbp1", "sbp2", "sbp3", "sbp4", "sbp5", "sbp6"]),
            ("PBLI", ["pbli1", "pbli2"]),
            ("PRO",  ["pro1", "pro2", "pro3", "pro4"]),
            ("ICS",  ["ics1", "ics2", "ics3"]),
        ]

    flat_cols = []
    for group, cols in milestone_cols:
        for col in cols:
            flat_cols.append((group, col))

    
    periods = get_milestone_rows_for_target(target, domain)

    data_rows = []

    for pgy, suffix, label in periods:
        values = []

        for _, base in flat_cols:
            field = f"{base}{suffix}_{pgy}"
            values.append(get_val(row, field, default=""))

        # Only keep rows that actually have values
        if any(str(v).strip() != "" for v in values):
            data_rows.append((label, values))

    if not data_rows:
        return

    table = doc.add_table(rows=3 + len(data_rows), cols=1 + len(flat_cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    set_row_height(table, 0.2)

    total_cols = 1 + len(flat_cols)

    header = table.cell(0, 0).merge(table.cell(0, total_cols - 1))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, title, bold=True, size=7)

    set_cell_text(table.cell(1, 0), "")
    shade_cell(table.cell(1, 0), "FFFFFF")

    col_idx = 1
    for group, cols in milestone_cols:
        start = col_idx
        end = col_idx + len(cols) - 1

        group_cell = table.cell(1, start).merge(table.cell(1, end))
        shade_cell(group_cell, "EDEDED")
        set_cell_text(group_cell, group, bold=True, size=6)

        col_idx = end + 1

    set_cell_text(table.cell(2, 0), "")

    col_idx = 1
    for group, cols in milestone_cols:
        for i in range(1, len(cols) + 1):
            shade_cell(table.cell(2, col_idx), "000000")
            set_cell_text(table.cell(2, col_idx), str(i), bold=True, size=6)
            col_idx += 1

    for r_idx, (label, values) in enumerate(data_rows, start=3):
        set_cell_text(table.cell(r_idx, 0), label, bold=True, size=6)

        for c_idx, value in enumerate(values, start=1):
            set_cell_text(
                table.cell(r_idx, c_idx),
                str(value).strip(),
                bold=True,
                size=6
            )

    first_col = 0.6
    remaining = TABLE_WIDTH - first_col
    each_col = remaining / len(flat_cols)

    force_table_width(table, [first_col] + [each_col] * len(flat_cols))


def html_to_clean_text_preserve_paragraphs(html):
    if pd.isna(html) or str(html).strip() == "":
        return ""

    soup = BeautifulSoup(str(html), "html.parser")

    # Convert paragraph-like tags to line breaks
    for tag in soup.find_all(["p", "div", "br", "li"]):
        tag.append("\n\n")

    text = soup.get_text()

    # Clean excess spaces but preserve paragraph breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    text = text.strip()

    return text


def add_duty_hours_summary_table(doc, row, target):
    """Supply the one helper called, but not defined, in the saved notebook.

    Table formatting is transcribed from the supplied Word example. The app
    supplies JSON rows so every dated concern fits the original four columns.
    Legacy numbered notebook fields remain accepted without a fixed row limit.
    """
    prepared = get_val(row, "duty_summary_rows", "")
    if prepared:
        entries = json.loads(prepared)
        if not isinstance(entries, list) or any(not isinstance(entry, dict) for entry in entries):
            raise ValueError("Duty-hour summary rows must be a list of records.")
    else:
        numbers = set()
        pattern = re.compile(r"^dh_(?:date|rule|description|rotation|work_type)_" + re.escape(target) + r"_(\d+)$")
        for field in row.index:
            match = pattern.fullmatch(str(field))
            if match:
                numbers.add(int(match.group(1)))
        entries = []
        for number in sorted(numbers):
            rule = get_val(row, f"dh_rule_{target}_{number}")
            description = get_val(row, f"dh_description_{target}_{number}")
            rotation = get_val(row, f"dh_rotation_{target}_{number}")
            observed = get_val(row, f"dh_date_{target}_{number}")
            work_type = get_val(row, f"dh_work_type_{target}_{number}")
            notes = get_val(row, f"dh_comments_{target}")
            if not any((rule, description, rotation, observed, work_type, notes)):
                continue
            details = []
            if observed:
                details.append("Date: " + observed)
            if work_type:
                details.append("Work type: " + work_type)
            if description:
                details.append(description)
            if notes:
                details.append(notes)
            entries.append({"rotation": rotation, "violation": rule, "notes": "\n".join(details)})
    if not entries:
        entries = [{"review_period": "", "rotation": "", "violation": "", "notes": get_val(row, f"dh_comments_{target}")}]

    table = doc.add_table(rows=2 + len(entries), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    force_table_width(table, MAIN_WIDTHS_4COL)
    set_row_height(table, 0.2)
    header = table.cell(0, 0).merge(table.cell(0, 3))
    shade_cell(header, "D9EAF7")
    set_cell_text(header, "DUTY HOURS SUMMARY", bold=True, size=7)
    headers = ["Review Period", "Rotation", "Violation", "Committee Notes / Resolution"]
    for column, title in enumerate(headers):
        shade_cell(table.cell(1, column), "EDEDED")
        set_cell_text(table.cell(1, column), title, bold=True, size=6)
    period_label = f"PGY{target[1]} {'MY' if target.endswith('my') else 'EOY'}"
    for index, entry in enumerate(entries, start=2):
        values = [entry.get("review_period", period_label), entry.get("rotation", ""),
                  entry.get("violation", ""), entry.get("notes", "")]
        for column, value in enumerate(values):
            set_cell_text(table.cell(index, column), "" if value is None else str(value),
                          bold=column == 0, size=6)


def create_doc(row, index):
    subject_name = get_val(row, "subject_name", f"resident_{index + 1}")

    target = find_resident_target(row)

    if target is None:
        return None

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.2)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.2)
    section.right_margin = Inches(0.2)

    add_summary_section(doc, row, target)
        
    # =========================================
    # COMPETENCY BY TRAINING YEAR
    # =========================================
    add_milestones_table(doc, row, target, domain="medicine")
    add_milestone_trend_summary_table(doc, row, target, domain="medicine")

    add_milestones_table(doc, row, target, domain="pediatrics")
    add_milestone_trend_summary_table(doc, row, target, domain="pediatrics")
    
    # =========================================
    # add_practice_readiness_snapshot
    # =========================================
    add_practice_readiness_snapshot(doc, row, target)
    
    # PROCEDURES LOGGED TO DATE
    add_html_section_with_title(
        doc,
        "PROCEDURES LOGGED TO DATE",
        get_val(row, "procedures"),
        table_width=7.6,
        header_gray=True
    )

    # CONFERENCE ATTENDANCE
    level = int(target[1])

    add_html_section_with_title(
        doc,
        "CONFERENCE ATTENDANCE",
        get_val(row, f"conferences_r{level}"),
        table_width=7.6,
        header_gray=True
    )
    
    # =========================================
    # DUTY HOURS
    # =========================================
    
    add_duty_hours_summary_table(doc, row, target)
    
    add_duty_hour_deep_dive_table(doc, row, target)

    
    # =========================================
    # OBSERVED ENCOUNTERS
    # =========================================
    add_observed_encounters_table(doc, row)
    
    # =========================================
    # Clinic Tally
    # ========================================= 
    add_clinic_tally_table(doc, row, target)
    
    # =========================================
    # INTRAINING SCORES
    # =========================================
    add_intraining_scores_table(doc, row)

    # =========================================
    # USMLE SCORES
    # =========================================
    
    #add_usmle_scores_table(doc, row)
    
    # =========================================
    # Scholarly Activities 
    # =========================================
    add_scholarly_activities_table(doc, row, target)
    
    # =========================================
    # Advocacy
    # =========================================
    add_advocacy_table(doc, row, target)
    
    # =========================================
    # RN 360 EVALUATION
    # =========================================

    level = int(target[1])
    add_simple_summary_table(doc,"360 EVALUATION FROM NURSING",
        html_to_clean_text(get_val(row, f"360_rn_{target}")),
        fill_color="D9EAF7"
    )
    
    # =========================================
    # OVERALL IMPRESSION
    # =========================================
    add_simple_summary_table(
        doc,
        "OVERALL IMPRESSION",
        html_to_clean_text(get_val(row, f"overall_imp_{target}")),
        fill_color="BFEDD2"
    )
    # =========================================
    # AREAS FOR GROWTH
    # =========================================
    
    add_simple_summary_table(
        doc,
        "AREAS FOR GROWTH",
        html_to_clean_text(get_val(row, f"areas_for_growth_{target}")),
        fill_color="F4CCCC"
    )
    # =========================================
    # PATTERNS OVER TIME
    # =========================================
    add_simple_summary_table(
        doc,
        "PATTERNS OVER TIME",
        html_to_clean_text(get_val(row, f"patterns_{target}")),
        fill_color="D9EAF7"
    )
    # =========================================
    # CCC Concerns
    # =========================================
    add_simple_summary_table(doc,"MEDICINE PEDIATRICS CLINICAL COMPETENCY COMMITTEE REVIEW",html_to_clean_text_preserve_paragraphs(get_val(row, f"ccc_concerns_{target}")),fill_color="D9EAF7")


    # =========================================
    # SIGNATURE
    # =========================================
    
    add_signature_section(doc)
    
    # =========================================
    # PAGE BREAK BEFORE EVALUATIONS
    # =========================================

    add_page_break(doc)

    subject_name = get_val(row, "subject_name")

    # =========================================
    # MEDICINE EVALUATIONS
    # =========================================

    medicine_html = get_val(row, f"mp5_{target}")

    add_html_table_to_doc(doc,clean_eval_html_keep_formatting(medicine_html,subject_name),table_width=7.6)

    # =========================================
    # PAGE BREAK BETWEEN EVAL TYPES
    # =========================================

    add_page_break(doc)

    # =========================================
    # PEDIATRIC EVALUATIONS
    # =========================================

    peds_html = get_val(row, f"mp3_{target}")

    add_html_table_to_doc(doc,clean_eval_html_keep_formatting(peds_html,subject_name),table_width=7.6)

    return doc
