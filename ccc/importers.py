"""Bounded, in-memory readers for the program's New Innovations exports.

This module does not persist files, assign resident identities, or compute CCC
decisions.  A successful parse is still an import preview.  Original assessment
questions, responses, comments, dates, and source locations remain available.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import PurePath
from typing import Any

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 400
MAX_SHEETS = 30
MAX_PDF_PAGES = 150
MAX_PDF_TEXT = 4_000_000
KINDS = {
    "evaluations", "self_assessments", "milestones", "procedures",
    "conferences", "duty_hours", "duty_violations",
}
SPECIALTIES = {"Medicine", "Pediatrics", "Combined"}


@dataclass
class ParseResult:
    records: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    profile: str = ""


@dataclass
class _Table:
    label: str
    headers: dict[str, str]
    rows: list[tuple[int, dict[str, Any]]]


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKC", str(value)).lower())


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _get(row: dict, *names: str, default: Any = "") -> Any:
    for name in names:
        value = row.get(_key(name))
        if _text(value) != "":
            return value
    return default


def _has(headers: dict | set, *names: str) -> bool:
    return any(_key(name) in headers for name in names)


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    if hasattr(value, "item"):
        return _json_value(value.item())
    return _text(value)


def _raw(table: _Table, row: dict) -> dict:
    return {table.headers.get(k, k): _json_value(v) for k, v in row.items()}


def _number(value: Any) -> float | None:
    text = _text(value).replace(",", "").strip()
    if text.endswith("%"):
        text = text[:-1].strip()
    if not text or not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        return None
    number = float(text)
    return number if math.isfinite(number) else None


def _datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and 59 <= value <= 100_000:
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    text = _text(value)
    if not text or not re.search(r"\d[/-]\d", text):
        return None
    try:
        import pandas as pd
        result = pd.to_datetime(text, errors="coerce", dayfirst=False)
        if pd.isna(result):
            return None
        return result.to_pydatetime().replace(tzinfo=None)
    except (ValueError, TypeError, OverflowError):
        return None


def _date(value: Any) -> str:
    parsed = _datetime(value)
    return parsed.date().isoformat() if parsed else ""


_DATE_TOKEN = r"(?:\d{4}-\d{1,2}-\d{1,2}|\d{1,2}/\d{1,2}/\d{2,4})"


def _date_range(value: Any) -> tuple[str, str]:
    found = re.findall(_DATE_TOKEN, _text(value))
    return (_date(found[0]), _date(found[1])) if len(found) >= 2 else ("", "")


def _pgy(value: Any) -> int | None:
    text = _text(value)
    match = re.search(r"\b(?:PGY|R|Resident)\s*[- ]?\s*([1-4])\b", text, re.I)
    if not match:
        match = re.search(r"\b([1-4])(?:st|nd|rd|th)\s*Year(?:\s*in\s*Residency)?", text, re.I)
    return int(match.group(1)) if match else None


def _resident(row: dict) -> tuple[str, str]:
    name = _text(_get(row, "resident_name", "Subject Name", "Person", "Name", "Resident"))
    if not name:
        last = _text(_get(row, "ResidentLastName", "Last Name", "LastName"))
        first = _text(_get(row, "ResidentFirstName", "First Name", "FirstName"))
        middle = _text(_get(row, "ResidentMiddleName", "Middle Name", "MiddleName"))
        if last and first:
            name = f"{last}, {' '.join(x for x in (first, middle) if x)}"
    external = _text(_get(row, "external_id", "Subject External ID", "Subject ID", "ResidentID", "Resident ID", "Person ID"))
    return name, external


def _specialty_value(value: Any) -> str:
    text = _key(value)
    if text in {"p", "peds", "pediatric", "pediatrics"}:
        return "Pediatrics"
    if text in {"m", "med", "medicine", "internalmedicine"}:
        return "Medicine"
    return "Combined"


def _specialty(row: dict, selected: str, result: ParseResult) -> str:
    source = _specialty_value(_get(row, "specialty", "eval_type"))
    if selected != "Combined":
        if source != "Combined" and source != selected:
            result.warnings.append(f"The export identifies some rows as {source}, but this upload was set to {selected}. The selected {selected} label will be used. Correct the upload choice or verify this override before saving.")
        return selected
    return source


_COMMON_HEADERS = {
    _key(x) for x in (
        "resident_name", "Subject Name", "Person", "Name", "Resident", "ResidentID",
        "Last Name", "First Name", "ResidentLastName", "ResidentFirstName",
        "Subject Status", "Subject Rotation", "Evaluator", "Evaluator Name",
        "Evaluator Department", "Evaluator Status", "Evaluator Credentials",
        "Evaluation Start Date", "Evaluation End Date", "Date Completed",
        "Start Date", "End Date", "Session Date", "Session", "Session Name",
        "Question", "Comment", "Response", "Overall Comments", "Training Year",
        "Procedure", "DatePerformed", "ResidentStatus", "Category", "Status",
        "% Attended", "Start", "End", "Hours", "Work Type", "Rotation",
        "Training Location", "Last Checked On", "Approved By", "Log Date", "Rule",
        "Description", "Comments", "date", "metric", "milestone", "code", "value",
        "score", "text", "period_start", "period_end", "report_dates", "form",
        "unit", "external_id", "specialty", "Evaluation ID", "eval_type", "record_mode",
    )
}


def _header_score(row: list) -> int:
    keys = {_key(value) for value in row if _text(value)}
    return len(keys & _COMMON_HEADERS) + sum(bool(re.fullmatch(r"question\d+(response|comment)?", k)) for k in keys)


def _matrix_to_table(matrix: list[list], label: str) -> _Table | None:
    if not matrix:
        return None
    if len(matrix) > MAX_ROWS + 50:
        raise ValueError(f"{label}: too many rows; export at most {MAX_ROWS:,} rows per file.")
    if max((len(row) for row in matrix), default=0) > MAX_COLUMNS:
        raise ValueError(f"{label}: more than {MAX_COLUMNS} columns; export a narrower report.")
    candidates = [(i, _header_score(row)) for i, row in enumerate(matrix[:40])]
    header_idx, score = max(candidates, key=lambda item: (item[1], -item[0]))
    if score < 2:
        return None
    headers: dict[str, str] = {}
    keys = []
    for n, value in enumerate(matrix[header_idx]):
        source = _text(value) or f"Column {n + 1}"
        key = _key(source) or f"column{n + 1}"
        if key in headers:
            key = f"{key}_{n + 1}"
        headers[key] = source
        keys.append(key)
    rows = []
    for offset, values in enumerate(matrix[header_idx + 1:], start=header_idx + 2):
        if not any(_text(value) for value in values):
            continue
        if len(values) == len(keys) and all(not _text(value) or _key(value) == _key(matrix[header_idx][n]) for n, value in enumerate(values)):
            continue  # Repeated printed headers are not resident records.
        rows.append((offset, {key: values[n] if n < len(values) else "" for n, key in enumerate(keys)}))
    return _Table(label, headers, rows)


def _decode(data: bytes, result: ParseResult) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if b"\x00" in data[:512]:
        # New Innovations TAB files sometimes omit the UTF-16 BOM.
        return data.decode("utf-16-le")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        result.warnings.append("This text export used a legacy Windows encoding; please check names and punctuation in the preview.")
        return data.decode("cp1252")


def _flat_header(value: Any) -> str:
    if isinstance(value, tuple):
        usable = [_text(x) for x in value if _text(x) and not str(x).startswith("Unnamed:")]
        unique = list(dict.fromkeys(usable))
        return unique[-1] if unique else ""
    return _text(value)


def _read_tables(filename: str, data: bytes, result: ParseResult) -> list[_Table]:
    matrices: list[tuple[str, list[list]]] = []
    suffix = PurePath(filename).suffix.lower()
    if data.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(info.file_size for info in members) > MAX_EXPANDED_BYTES:
                raise ValueError("The workbook expands beyond the supported size; export fewer rows or sheets.")
            if any(info.flag_bits & 1 for info in members):
                raise ValueError("Unlock the password-protected Excel file before importing it. The application encrypts saved data separately.")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("This ZIP file is not an Excel workbook. Upload an individual report file.")
            has_formulas = any(
                re.search(rb"<f(?:\s|>)", archive.read(info.filename))
                for info in members if re.match(r"xl/worksheets/sheet\d+\.xml$", info.filename)
            )
            if has_formulas:
                result.warnings.append("The workbook contains formulas. Only values cached by Excel are read; formulas are never executed. Blank cached values need review.")
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise ValueError(f"The workbook contains more than {MAX_SHEETS} sheets; upload fewer sheets.")
            for sheet in book.worksheets:
                if sheet.max_row and sheet.max_row > MAX_ROWS + 40:
                    raise ValueError(f"{sheet.title}: too many rows; export at most {MAX_ROWS:,} rows.")
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise ValueError(f"{sheet.title}: more than {MAX_COLUMNS} columns.")
                matrix = []
                for cells in sheet.iter_rows():
                    row = []
                    for cell in cells:
                        value = cell.value
                        if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in cell.number_format:
                            value = f"{value * 100:g}%"
                        row.append(value)
                    matrix.append(row)
                    if len(matrix) > MAX_ROWS + 40:
                        raise ValueError("The workbook contains too many rows.")
                matrices.append((sheet.title, matrix))
        finally:
            book.close()
    elif data.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("Legacy .xls support requires xlrd. Save the export as .xlsx or CSV, or install the app's listed dependencies.") from exc
        book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=True)
        try:
            if book.nsheets > MAX_SHEETS:
                raise ValueError(f"The workbook contains more than {MAX_SHEETS} sheets.")
            for sheet in book.sheets():
                if sheet.nrows > MAX_ROWS + 40 or sheet.ncols > MAX_COLUMNS:
                    raise ValueError("The legacy workbook exceeds the supported row or column limit.")
                matrix = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_NUMBER:
                            format_key = book.xf_list[sheet.cell_xf_index(r, c)].format_key
                            format_string = book.format_map.get(format_key)
                            if format_string and "%" in format_string.format_str:
                                value = f"{value * 100:g}%"
                        row.append(value)
                    matrix.append(row)
                matrices.append((sheet.name, matrix))
        finally:
            book.release_resources()
    else:
        text = _decode(data, result)
        if re.search(r"<table\b", text[:200_000], re.I):
            import pandas as pd
            # StringIO prevents pandas from interpreting input text as a URL.
            frames = pd.read_html(io.StringIO(text), flavor="lxml", keep_default_na=False)
            if len(frames) > MAX_SHEETS:
                raise ValueError(f"The report contains more than {MAX_SHEETS} tables.")
            for n, frame in enumerate(frames, 1):
                if len(frame) > MAX_ROWS or len(frame.columns) > MAX_COLUMNS:
                    raise ValueError("The HTML report exceeds the supported row or column limit.")
                matrix = frame.astype(object).values.tolist()
                if not all(isinstance(x, int) for x in frame.columns):
                    matrix.insert(0, [_flat_header(x) for x in frame.columns])
                matrices.append((f"Table {n}", matrix))
            result.profile = "New Innovations HTML / Excel export"
        else:
            if suffix in {".xlsx", ".xlsm"}:
                raise ValueError("The file extension says Excel, but its contents are not a readable workbook.")
            if text.lstrip().startswith("<?xml"):
                raise ValueError("This legacy XML workbook format is not supported. Save it as .xlsx or CSV before uploading.")
            sample = text[:50_000]
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters="\t,;|").delimiter
            except csv.Error:
                delimiter = "\t" if suffix in {".tab", ".tsv"} or "\t" in sample else ","
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            matrix = []
            for row in reader:
                matrix.append(row)
                if len(matrix) > MAX_ROWS + 40:
                    raise ValueError(f"The file contains more than {MAX_ROWS:,} rows. Export a smaller reporting period.")
                if len(row) > MAX_COLUMNS:
                    raise ValueError(f"The file contains more than {MAX_COLUMNS} columns.")
            matrices.append(("Data", matrix))
    if sum(len(matrix) for _, matrix in matrices) > MAX_ROWS + 40 * MAX_SHEETS:
        raise ValueError(f"The file contains more than {MAX_ROWS:,} rows across its sheets.")
    tables = [table for label, matrix in matrices if (table := _matrix_to_table(matrix, label)) is not None]
    if sum(len(table.rows) for table in tables) > MAX_ROWS:
        raise ValueError(f"The file contains more than {MAX_ROWS:,} data rows across its sheets.")
    if not tables:
        raise ValueError("No supported report header was found in the first 40 rows. Upload the original New Innovations export or the documented long CSV template.")
    return tables


def _dates(record: dict, row: dict, selected_start: str, selected_end: str, *, event_names: tuple[str, ...] = (), snapshot: bool = False) -> None:
    range_start, range_end = _date_range(_get(row, "report_dates", "Report Dates", "Report Period"))
    start = _date(_get(row, "period_start", "Evaluation Start Date", "Start Date")) or range_start
    end = _date(_get(row, "period_end", "Evaluation End Date", "End Date")) or range_end
    completed = _date(_get(row, "Date Completed", "Completed Date", "Completion Date"))
    event = _date(_get(row, *event_names)) if event_names else ""
    explicit = _date(_get(row, "date", "Assessment Date", "Report Date"))
    real_date = completed or event or explicit or end or start
    record["date"] = real_date or selected_end
    record["period_start"] = start or selected_start
    record["period_end"] = end or selected_end
    record["completed_date"] = completed
    record["date_is_fallback"] = not bool(real_date)
    record["date_source"] = (
        "completion date" if completed else "event date" if event else "assessment date" if explicit
        else "source period end" if end else "source period start" if start
        else "selected reporting period end" if selected_end else "missing"
    )
    if start and end and start > end:
        record["date_needs_confirmation"] = True
    source_date_fields = {
        _key(name) for name in (
            "period_start", "Evaluation Start Date", "Start Date", "period_end",
            "Evaluation End Date", "End Date", "Date Completed", "Completed Date",
            "Completion Date", "date", "Assessment Date", "Report Date", *event_names,
        )
    }
    record["source_date_needs_review"] = any(
        _text(row.get(key)) and not _date(row[key]) for key in source_date_fields
    )
    if snapshot:
        record["record_mode"] = "snapshot"


def _base(table: _Table, rownum: int, row: dict, kind: str, specialty: str, result: ParseResult) -> dict:
    name, external = _resident(row)
    return {
        "resident_name": name, "external_id": external, "category": kind,
        "external_record_id": _text(_get(row, "external_record_id", "Evaluation ID", "Evaluation Instance ID", "ProcedureLogID", "Procedure Log ID", "LogID", "Log ID", "ViolationID", "Violation ID", "DutyHourID", "ShiftID", "Event ID")),
        "specialty": _specialty(row, specialty, result), "date": "",
        "period_start": "", "period_end": "", "form": "", "metric": "",
        "value": None, "unit": "", "text": "", "questions": [], "values": {},
        "record_mode": "event", "source_locator": f"{table.label}, row {rownum}",
        "evaluator": _text(_get(row, "Evaluator Name", "Evaluator")),
        "evaluator_status": _text(_get(row, "Evaluator Status", "Evaluator Credentials")),
        "rotation": _text(_get(row, "Subject Rotation", "Rotation")),
        "pgy": _pgy(_get(row, "Subject Status", "ResidentStatus", "Status", "Training Year", "training_year", "PGY")),
        "raw": _raw(table, row),
    }


def _questions_wide(row: dict) -> list[dict]:
    parts: dict[int, dict] = {}
    for key, value in row.items():
        match = re.fullmatch(r"question(\d+)(response|comment)?", key)
        if not match:
            continue
        n = int(match.group(1))
        question = parts.setdefault(n, {"number": n, "text": "", "response": "", "comment": ""})
        question[match.group(2) or "text"] = _json_value(value) if match.group(2) == "response" else _text(value)
    return [q for _, q in sorted(parts.items()) if any(_text(q[k]) for k in ("text", "response", "comment"))]


def _unique_text(values: list[Any]) -> list[str]:
    return list(dict.fromkeys(_text(value) for value in values if _text(value)))


def _parse_assessments(table: _Table, kind: str, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    headers = table.headers
    long = _has(headers, "Question") and _has(headers, "Comment", "Response", "Rating", "Answer", "Score")
    wide = any(re.fullmatch(r"question\d+(response|comment)?", k) for k in headers)
    narrative = _has(headers, "Overall Comments") and _has(headers, "Session", "Session Name", "Evaluator Name", "Evaluator")
    has_name = _has(headers, "Subject Name", "resident_name", "Person", "Name") or (_has(headers, "First Name") and _has(headers, "Last Name"))
    if not has_name or not (long or wide or narrative):
        return False
    groups: OrderedDict[tuple, list[tuple[int, dict]]] = OrderedDict()
    excluded = 0
    for rownum, row in table.rows:
        subject_status = _text(_get(row, "Subject Status")).lower()
        if subject_status in {"faculty", "attending", "staff"}:
            excluded += 1
            continue
        name, external = _resident(row)
        if not name:
            result.errors.append(f"{table.label}, row {rownum}: resident name is missing.")
            continue
        if long:
            context = (
                external or name, _text(_get(row, "Evaluation ID", "Evaluation Instance ID")),
                _text(_get(row, "Session Name", "Session", "form")),
                _text(_get(row, "Subject Rotation", "Rotation")),
                _text(_get(row, "Evaluator", "Evaluator Name")),
                _text(_get(row, "Subject Status")),
                _text(_get(row, "Evaluation Start Date", "Start Date", "Session Date")),
                _text(_get(row, "Evaluation End Date", "End Date")),
                _text(_get(row, "Date Completed", "Completed Date")),
                _text(_get(row, "specialty", "eval_type")),
            )
        else:
            context = (rownum,)
        groups.setdefault(context, []).append((rownum, row))
    duplicates = 0
    for rows in groups.values():
        rownum, first = rows[0]
        rec = _base(table, rownum, first, kind, specialty, result)
        rec["form"] = _text(_get(first, "Session Name", "Session", "form")) or ("Self-assessment" if kind == "self_assessments" else "Evaluation")
        rec["metric"] = "Self-assessment" if kind == "self_assessments" else "Evaluation"
        rec["value"] = 1
        rec["unit"] = "assessments"
        _dates(rec, first, start, end, event_names=("Session Date",))
        rec["source_evaluation_id"] = _text(_get(first, "Evaluation ID", "Evaluation Instance ID"))
        rec["source_rows"] = [n for n, _ in rows]
        if len(rows) > 1:
            rec["source_locator"] = f"{table.label}, rows {rows[0][0]}–{rows[-1][0]}"
        questions, seen, overall = [], set(), []
        for _, row in rows:
            overall.append(_get(row, "Overall Comments"))
            if long:
                question = {
                    "number": _text(_get(row, "Question Number", "Question ID")),
                    "text": _text(_get(row, "Question")),
                    "response": _json_value(_get(row, "Response", "Answer", "Rating", "Score", "Value")),
                    "comment": _text(_get(row, "Comment")),
                }
                fingerprint = json.dumps(question, sort_keys=True, ensure_ascii=False)
                if fingerprint in seen:
                    duplicates += 1
                    continue
                seen.add(fingerprint)
                questions.append(question)
            else:
                questions.extend(_questions_wide(row))
        rec["questions"] = questions
        rec["overall_comment"] = "\n\n".join(_unique_text(overall))
        rec["text"] = "\n\n".join(_unique_text([q["comment"] for q in questions] + overall))
        result.records.append(rec)
    if excluded:
        result.warnings.append(f"Excluded {excluded} row(s) explicitly identified as faculty, attending, or staff subjects.")
    if duplicates:
        result.warnings.append(f"Removed {duplicates} exact duplicate question row(s) within this file; distinct responses and comments were preserved.")
    if long and not _has(headers, "Evaluation ID", "Evaluation Instance ID"):
        result.warnings.append("Question rows were grouped by resident, form, rotation, dates, and evaluator because this export has no evaluation ID. Review evaluation counts before saving.")
    result.profile = "New Innovations question-row evaluations" if long else "New Innovations assessment form export"
    return True


def _milestone_code(value: Any) -> str:
    match = re.match(r"^\s*(PC|MK|SBP|PBLI|PROF?|ICS)\s*[-_ ]?\s*(\d+)\b", _text(value), re.I)
    if not match:
        return ""
    prefix = match.group(1).upper()
    return f"{'PROF' if prefix == 'PRO' else prefix}{int(match.group(2))}"


def _score(value: Any) -> tuple[float | None, str]:
    text = _text(value)
    if text.lower() in {"", "x", "na", "n/a", "not assessed", "not yet assessed", "unable to assess", "not applicable", "--", "—"}:
        return None, text or "Not supplied"
    number = _number(text)
    if number is None or not 0 <= number <= 5:
        raise ValueError("milestone score must be a number from 0 through 5 or an explicit not-assessed marker")
    return number, "Assessed"


def _parse_milestones(table: _Table, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not _has(table.headers, "milestone", "code", "metric") or not _has(table.headers, "value", "score", "rating"):
        return False
    if not _has(table.headers, "resident_name", "Subject Name", "Resident", "Name"):
        return False
    for rownum, row in table.rows:
        rec = _base(table, rownum, row, "milestones", specialty, result)
        code = _milestone_code(_get(row, "milestone", "code", "metric"))
        if not rec["resident_name"] or not code:
            result.errors.append(f"{table.label}, row {rownum}: resident name and an explicit milestone code (for example PC1) are required.")
            continue
        if rec["specialty"] == "Combined":
            result.errors.append(f"{table.label}, row {rownum}: select Medicine or Pediatrics for milestone scores, or provide a specialty column.")
            continue
        try:
            rec["value"], status = _score(_get(row, "value", "score", "rating"))
        except ValueError as exc:
            result.errors.append(f"{table.label}, row {rownum}: {exc}.")
            continue
        rec.update(metric=code, unit="milestone level", record_mode="snapshot", form=_text(_get(row, "form")) or "Milestone assessment")
        rec["values"] = {"assessment_status": status, "source_label": _text(_get(row, "milestone", "code", "metric"))}
        _dates(rec, row, start, end, snapshot=True)
        result.records.append(rec)
    result.profile = "Structured milestone scores"
    return True


def _parse_procedures(table: _Table, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not _has(table.headers, "Procedure") or not _has(table.headers, "DatePerformed", "Date Performed", "date"):
        return False
    for rownum, row in table.rows:
        rec = _base(table, rownum, row, "procedures", specialty, result)
        rec.update(metric=_text(_get(row, "Procedure")), value=1, unit="procedures", form="Procedure log")
        if not rec["resident_name"] or not rec["metric"]:
            result.errors.append(f"{table.label}, row {rownum}: procedure and resident name are required.")
            continue
        _dates(rec, row, start, end, event_names=("DatePerformed", "Date Performed", "date"))
        rec["text"] = _text(_get(row, "Comments", "Comment", "Description"))
        rec["values"] = {"resident_status": _text(_get(row, "ResidentStatus", "Status"))}
        result.records.append(rec)
    result.profile = "New Innovations procedure logger"
    return True


def _parse_conferences(table: _Table, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not _has(table.headers, "% Attended", "Percent Attended", "Attendance Percent") or not _has(table.headers, "Person", "resident_name", "Name"):
        return False
    for rownum, row in table.rows:
        rec = _base(table, rownum, row, "conferences", specialty, result)
        value = _number(_get(row, "% Attended", "Percent Attended", "Attendance Percent"))
        if not rec["resident_name"]:
            result.errors.append(f"{table.label}, row {rownum}: resident name is missing.")
            continue
        if value is None or not 0 <= value <= 100:
            result.errors.append(f"{table.label}, row {rownum}: attendance must be a percentage from 0 to 100.")
            continue
        category = _text(_get(row, "Category")) or "Conference attendance"
        rec.update(metric=category, value=value, unit="percent", form="Conference attendance summary", record_mode="snapshot")
        _dates(rec, row, start, end, snapshot=True)
        rec["values"] = {
            "category": category, "attendance_percent": value,
            "status": _text(_get(row, "Status")),
            "attended": _number(_get(row, "Attended", "Sessions Attended", "Number Attended")),
            "offered": _number(_get(row, "Offered", "Sessions Offered", "Total Sessions")),
        }
        result.records.append(rec)
    result.warnings.append("Conference percentages are report snapshots. Keep their reporting dates; overlapping snapshots must not be added or averaged together.")
    result.profile = "New Innovations conference attendance snapshot"
    return True


def _parse_duty(table: _Table, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not all(_has(table.headers, h) for h in ("Start", "End", "Hours")):
        return False
    current_name, current_status = "", ""
    hierarchical = not _has(table.headers, "Name", "Person", "resident_name", "Subject Name")
    if hierarchical and not _has(table.headers, "Work Type"):
        return False
    skipped = 0
    for rownum, source in table.rows:
        row = dict(source)
        work_type = _text(_get(row, "Work Type"))
        if hierarchical:
            if re.fullmatch(r"(?:R|PGY)\s*[1-4]", work_type, re.I):
                current_status, current_name = work_type, ""
                continue
            has_start = bool(_datetime(_get(row, "Start")))
            if not has_start and "," in work_type and not re.search(r"\d|total|hrs\s*:", work_type, re.I):
                current_name = work_type
                continue
            if not has_start:
                skipped += 1
                continue
            row[_key("Name")] = current_name
            row[_key("Status")] = current_status
        if re.search(r"(?:total\s*hrs|hrs\s*:)", " ".join(_text(v) for v in row.values()), re.I):
            skipped += 1
            continue
        rec = _base(table, rownum, row, "duty_hours", specialty, result)
        first, last = _datetime(_get(row, "Start")), _datetime(_get(row, "End"))
        hours = _number(_get(row, "Hours"))
        if not rec["resident_name"]:
            result.errors.append(f"{table.label}, row {rownum}: shift has no resident name or preceding resident header.")
            continue
        if not first or not last or last < first:
            result.errors.append(f"{table.label}, row {rownum}: a valid start and end datetime are required; end must not precede start.")
            continue
        calculated = (last - first).total_seconds() / 3600
        if hours is None or hours < 0:
            result.errors.append(f"{table.label}, row {rownum}: reported Hours is missing or invalid.")
            continue
        if abs(calculated - hours) > 0.25:
            result.warnings.append("Some duty-hour totals differ from elapsed start/end time. Reported Hours was preserved; review these shifts before saving.")
            rec["duration_needs_review"] = True
        rec.update(metric="Duty hours", value=hours, unit="hours", form="Duty-hour log")
        _dates(rec, row, start, end, event_names=("Start",))
        rec["period_start"], rec["period_end"] = first.date().isoformat(), last.date().isoformat()
        rec["values"] = {
            "start_datetime": first.isoformat(), "end_datetime": last.isoformat(),
            "start": first.isoformat(), "end": last.isoformat(),
            "reported_hours": hours, "elapsed_hours": round(calculated, 6), "work_type": work_type,
            "training_location": _text(_get(row, "Training Location")),
            "last_checked_on": _text(_get(row, "Last Checked On")),
            "approved_by": _text(_get(row, "Approved By")),
            "status": _text(_get(row, "Status")),
        }
        result.records.append(rec)
    if skipped:
        result.warnings.append(f"Skipped {skipped} duty-log heading or summary row(s); individual dated shifts were retained.")
    result.profile = "New Innovations duty-hour shifts"
    return True


def _parse_violations(table: _Table, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not _has(table.headers, "Rule") or not _has(table.headers, "Log Date", "date") or not _has(table.headers, "Name", "resident_name", "Person"):
        return False
    for rownum, row in table.rows:
        rec = _base(table, rownum, row, "duty_violations", specialty, result)
        rule = _text(_get(row, "Rule"))
        if not rec["resident_name"] or not rule:
            result.errors.append(f"{table.label}, row {rownum}: resident name and source rule are required.")
            continue
        rec.update(metric=rule, value=1, unit="reported violations", form="New Innovations reported duty-hour violation")
        _dates(rec, row, start, end, event_names=("Log Date", "date"))
        rec["text"] = "\n\n".join(_unique_text([_get(row, "Description"), _get(row, "Comments")]))
        rec["values"] = {"rule": rule, "description": _text(_get(row, "Description")), "comments": _text(_get(row, "Comments")), "status": _text(_get(row, "Status")), "work_type": _text(_get(row, "Work Type"))}
        result.records.append(rec)
    result.profile = "New Innovations reported duty-hour violations"
    return True


def _parse_generic(table: _Table, kind: str, specialty: str, start: str, end: str, result: ParseResult) -> bool:
    if not _has(table.headers, "resident_name") or not _has(table.headers, "metric") or not _has(table.headers, "value", "text"):
        return False
    for rownum, row in table.rows:
        rec = _base(table, rownum, row, kind, specialty, result)
        if not rec["resident_name"] or not _text(_get(row, "metric")):
            result.errors.append(f"{table.label}, row {rownum}: resident_name and metric are required.")
            continue
        source_value = _get(row, "value")
        number = _number(source_value)
        mode = _text(_get(row, "record_mode")).lower()
        if mode and mode not in {"event", "snapshot"}:
            result.errors.append(f"{table.label}, row {rownum}: record_mode must be event or snapshot.")
            continue
        rec.update(metric=_text(_get(row, "metric")), value=number if number is not None else (_json_value(source_value) or None), text=_text(_get(row, "text")), unit=_text(_get(row, "unit")), form=_text(_get(row, "form")))
        _dates(rec, row, start, end, snapshot=kind in {"conferences", "milestones"} or mode == "snapshot")
        rec["questions"] = _questions_wide(row)
        result.records.append(rec)
    result.profile = "Standardized long CSV / spreadsheet"
    return True


def _read_pdf_pages(data: bytes) -> list[str]:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise ValueError(f"The PDF contains more than {MAX_PDF_PAGES} pages. Upload smaller reports.")
        pages, length = [], 0
        for page in pdf.pages:
            text = page.extract_text() or ""
            length += len(text)
            if length > MAX_PDF_TEXT:
                raise ValueError("The PDF contains too much extracted text; upload smaller reports.")
            pages.append(text)
        return pages


_NI_MILESTONE_HEADING = r"Program\s*Subcompetency\s*for\s*Residents"
_PDF_MILESTONE_GROUPS = {
    "Medicine": (("PC", 6), ("MK", 3), ("SBP", 3), ("PBLI", 2), ("PROF", 4), ("ICS", 3)),
    "Pediatrics": (("PC", 5), ("MK", 2), ("SBP", 6), ("PBLI", 2), ("PROF", 4), ("ICS", 3)),
}


def _read_pdf_milestone_tables(data: bytes) -> list[dict]:
    """Read bounded, labeled table cells without retaining the PDF itself.

    The New Innovations report places each domain in its own small grid. Its
    wrapped headers and lower PC/SBP rows are interleaved by ordinary PDF text
    extraction, so a final list of numbers cannot identify the right scores.
    """
    import pdfplumber
    pages, text_length, row_count = [], 0, 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise ValueError(f"The PDF contains more than {MAX_PDF_PAGES} pages. Upload smaller reports.")
        for number, page in enumerate(pdf.pages, 1):
            # Preserve the spaces in the resident's printed name. The report's
            # tiny milestone font otherwise causes pdfplumber to join words.
            text = page.extract_text(x_tolerance=1, y_tolerance=2) or ""
            text_length += len(text)
            if text_length > MAX_PDF_TEXT:
                raise ValueError("The PDF contains too much extracted text; upload smaller reports.")
            tables = page.find_tables() if re.search(_NI_MILESTONE_HEADING, text, re.I) else []
            if len(tables) > MAX_SHEETS:
                raise ValueError("The PDF contains too many tables on one page. Export smaller milestone reports.")
            extracted = []
            for table in tables:
                rows = table.extract(x_tolerance=1, y_tolerance=2)
                row_count += len(rows)
                if row_count > MAX_ROWS or any(len(row) > MAX_COLUMNS for row in rows):
                    raise ValueError("The PDF milestone tables exceed the supported row or column limit.")
                extracted.append(rows)
            pages.append({"page_number": number, "text": text, "tables": extracted})
    return pages


def _parse_pdf_milestone_tables(data: bytes, specialty: str, start: str, end: str, result: ParseResult) -> None:
    """Match explicit milestone header cells to the score cells beneath them."""
    expected = [f"{domain}{number}" for domain, count in _PDF_MILESTONE_GROUPS[specialty]
                for number in range(1, count + 1)]
    expected_codes = set(expected)
    found, seen = [], {}
    duplicate_pages = 0
    for page in _read_pdf_milestone_tables(data):
        page_num, text = page["page_number"], page["text"]
        if not text.strip():
            result.errors.append(f"PDF page {page_num}: no selectable milestone text was found. Re-export this page as a text-based PDF or remove a genuinely blank page before importing.")
            return
        sections = list(re.finditer(_NI_MILESTONE_HEADING, text, re.I))
        if len(sections) != 1:
            result.errors.append(f"PDF page {page_num}: expected one resident's Program Subcompetency milestone section. Export one assessment per page.")
            return
        section = text[sections[0].start():]
        detected = []
        if re.search(r"Internal\s*Medicine\s*Milestones", section, re.I):
            detected.append("Medicine")
        if re.search(r"Pediatrics?\s*Milestones", section, re.I):
            detected.append("Pediatrics")
        if detected != [specialty]:
            result.errors.append(f"PDF page {page_num}: the milestone report does not match the selected {specialty} specialty. Use the matching upload area.")
            return
        header = text[:sections[0].start()]
        ranges = list(re.finditer(rf"({_DATE_TOKEN})\s*(?:-|–|—|to)\s*({_DATE_TOKEN})", header, re.I))
        if len(ranges) != 1:
            result.errors.append(f"PDF page {page_num}: a single resident report date range is required above the milestone table.")
            return
        report_range = ranges[0]
        period_start, period_end = _date(report_range.group(1)), _date(report_range.group(2))
        if not period_start or not period_end or period_start > period_end:
            result.errors.append(f"PDF page {page_num}: the printed milestone report dates are invalid or reversed.")
            return
        names = _unique_text(re.findall(r"^\s*(?:Resident(?:\s+Name)?|Subject(?:\s+Name)?|Name)\s*:\s*(.+)$",
                                       header[:report_range.start()], re.I | re.M))
        previous = [line.strip() for line in header[:report_range.start()].splitlines() if line.strip()]
        name = names[0] if len(names) == 1 else previous[-1] if not names and previous else ""
        if not name or re.search(r"\d|Training\s*Year|Program\s*Subcompetency", name, re.I):
            result.errors.append(f"PDF page {page_num}: a resident name could not be reliably read above the date range. Use a structured milestone CSV for this page.")
            return
        # Rotation names can contain '(PGY4)' even when the resident is PGY1.
        # Read only the training-year row inside the milestone section.
        years = {int(value) for value in re.findall(r"\b([1-4])(?:st|nd|rd|th)\s*Year\s*in\s*Residency", section, re.I)}
        if len(years) != 1:
            result.errors.append(f"PDF page {page_num}: one training year from PGY1 through PGY4 is required in the milestone table.")
            return
        pgy = years.pop()
        scores = {}
        for table_num, rows in enumerate(page["tables"], 1):
            for row_num, row in enumerate(rows):
                for column, cell in enumerate(row):
                    label = " ".join(_text(cell).split())
                    if not re.fullmatch(r"(?:PC|MK|SBP|PBLI|PROF?)\s*[-_ ]?\s*\d+|ICS\s*[-_ ]?\s*\d+", label, re.I):
                        continue
                    code = _milestone_code(label)
                    if code not in expected_codes or code in scores:
                        result.errors.append(f"PDF page {page_num}: the table has an unexpected or repeated milestone label ({code}). Check the selected specialty and export one assessment per page.")
                        return
                    # A None cell is the structural gap created by another
                    # domain's taller header. An empty real cell is not a score.
                    next_cells = [(index, candidate[column]) for index, candidate in enumerate(rows[row_num + 1:], row_num + 1)
                                  if column < len(candidate) and candidate[column] is not None]
                    value_row, supplied = next_cells[0] if next_cells else (-1, "")
                    try:
                        if not _text(supplied):
                            raise ValueError("blank score")
                        value, status = _score(supplied)
                    except ValueError:
                        result.errors.append(f"PDF page {page_num}: {code} has no single valid score directly below its label. Use the original report or a structured CSV to correct this cell.")
                        return
                    scores[code] = (value, status, label, _text(supplied), table_num, row_num + 1, value_row + 1, column + 1)
        if set(scores) != expected_codes:
            result.errors.append(f"PDF page {page_num}: found {len(scores)} labeled milestone scores; the {specialty} report requires {len(expected)}. Re-export the complete table or use the structured milestone CSV.")
            return
        report_dates = f"{report_range.group(1)} - {report_range.group(2)}"
        records = []
        for code in expected:
            value, status, label, supplied, table_num, label_row, value_row, column = scores[code]
            row = {_key("resident_name"): name, _key("training_year"): f"PGY{pgy}",
                   _key("report_dates"): report_dates, _key("milestone"): label, _key("value"): supplied}
            table = _Table("PDF", {key: key for key in row}, [])
            record = _base(table, value_row, row, "milestones", specialty, result)
            record.update(metric=code, value=value, unit="milestone level", form="Milestone assessment",
                          record_mode="snapshot", pgy=pgy, source_period_start=period_start,
                          source_period_end=period_end,
                          source_locator=f"PDF page {page_num}, table {table_num}, label row {label_row}, score row {value_row}, column {column}")
            record["values"] = {"assessment_status": status, "source_label": label}
            record["resident_name_requires_confirmation"] = not bool(names)
            _dates(record, row, start, end, snapshot=True)
            key = (name.casefold(), period_start, period_end, pgy, code)
            score = (value, status.casefold())
            if key in seen:
                if seen[key] != score:
                    result.errors.append("The PDF contains conflicting scores for the same resident and reporting period. Export the intended assessment before saving.")
                    return
            else:
                seen[key] = score
                records.append(record)
        if not records:
            duplicate_pages += 1
        found.extend(records)
    if not found:
        result.errors.append("No complete New Innovations milestone tables were found in this PDF.")
        return
    if duplicate_pages:
        result.warnings.append(f"Skipped {duplicate_pages} repeated milestone report page(s) within this file.")
    result.records.extend(found)
    result.profile = "New Innovations milestone grid PDF"


def _parse_pdf(data: bytes, kind: str, specialty: str, start: str, end: str, result: ParseResult) -> None:
    if kind != "milestones":
        result.errors.append("This PDF report type cannot yet be imported reliably. Upload the corresponding New Innovations Excel, CSV, or TAB export. PDFs with explicit milestone codes and scores are supported in Milestones.")
        return
    if specialty == "Combined":
        result.errors.append("Select Medicine or Pediatrics before importing a milestone PDF.")
        return
    pages = _read_pdf_pages(data)
    text = "\n".join(pages)
    if not text.strip():
        result.errors.append("This PDF has no selectable text. Upload a text-based PDF or a structured milestone CSV; scanned pages need review and transcription.")
        return
    if re.search(_NI_MILESTONE_HEADING, text, re.I):
        _parse_pdf_milestone_tables(data, specialty, start, end, result)
        return
    names = _unique_text(re.findall(r"^\s*(?:Resident(?:\s+Name)?|Subject(?:\s+Name)?|Name)\s*:\s*(.+)$", text, re.I | re.M))
    if len(names) > 1:
        result.errors.append("This PDF appears to contain multiple residents. Upload one resident's milestone report at a time, or use the long CSV template.")
        return
    name = names[0] if names else ""
    date_match = re.search(rf"({_DATE_TOKEN})\s*(?:-|–|—|to)\s*({_DATE_TOKEN})", text, re.I)
    period_start, period_end = (_date(date_match.group(1)), _date(date_match.group(2))) if date_match else ("", "")
    name_inferred = False
    if not name and date_match:
        previous = [line.strip() for line in text[:date_match.start()].splitlines() if line.strip()]
        if previous:
            name = previous[-1]
            name_inferred = True
    if not name:
        result.errors.append("A resident name could not be reliably located in this PDF. Use the structured milestone CSV with resident_name, date, milestone, and value columns.")
        return
    found, ambiguous = [], 0
    for page_num, page in enumerate(pages, 1):
        for line_num, line in enumerate(page.splitlines(), 1):
            match = re.match(r"^\s*((?:PC|MK|SBP|PBLI|PROF?|ICS)\s*[-_ ]?\s*\d+)\b\s*(.*)$", line, re.I)
            if not match:
                continue
            code, tail = _milestone_code(match.group(1)), match.group(2).strip(" :|\t")
            # A row is accepted only when a single explicit scalar or missing
            # marker follows the code/description. Positional pages are refused.
            marker = re.search(r"(?<![\w.])(?:[+-]?\d+(?:\.\d+)?|x|n/a|not assessed|not yet assessed|unable to assess)\s*$", tail, re.I)
            if not marker:
                ambiguous += 1
                continue
            prefix = tail[:marker.start()]
            if re.search(r"(?<![\w.])\d+(?:\.\d+)?(?![\w.])", prefix):
                ambiguous += 1
                continue
            try:
                value, status = _score(marker.group(0).strip())
            except ValueError:
                ambiguous += 1
                continue
            row = {_key("resident_name"): name, _key("period_start"): period_start, _key("period_end"): period_end}
            table = _Table("PDF", {k: k for k in row}, [])
            rec = _base(table, line_num, row, kind, specialty, result)
            rec.update(metric=code, value=value, unit="milestone level", form="New Innovations milestone report", record_mode="snapshot", pgy=_pgy(text), source_locator=f"PDF page {page_num}, line {line_num}")
            rec["values"] = {"assessment_status": status, "source_label": line.strip()}
            rec["resident_name_requires_confirmation"] = name_inferred
            _dates(rec, row, start, end, snapshot=True)
            found.append(rec)
    if not found:
        result.errors.append("No unambiguous labeled milestone score rows were found. The old notebook assigned the last 21 or 22 numbers by position; this app will not guess that mapping. Upload a long CSV with resident_name, date, milestone, and value, or a PDF that prints each code beside its score.")
        return
    by_code: dict[str, set] = {}
    for record in found:
        by_code.setdefault(record["metric"], set()).add(record["value"])
    if any(len(values) > 1 for values in by_code.values()):
        result.errors.append("This PDF shows multiple different scores for the same milestone. Select the intended assessment/date in a structured CSV before importing.")
        return
    if ambiguous:
        result.errors.append(f"{ambiguous} labeled milestone row(s) could not be read unambiguously. Complete the missing scores using the long CSV template before saving this report.")
    if name_inferred:
        result.warnings.append("The PDF name was inferred from the line before its date range. Confirm the resident mapping in the preview.")
    expected_count = 21 if specialty == "Medicine" else 22
    if len(by_code) != expected_count:
        result.warnings.append(f"This PDF yielded {len(by_code)} distinct milestone codes. The supplied {specialty} notebook expects {expected_count}; confirm that the report is intentionally partial and no scores are missing.")
    result.records.extend(found)
    result.profile = "PDF with explicit milestone code / score rows"


_EXPECTED = {
    "evaluations": "Expected a New Innovations evaluation export with resident names, evaluator/form, and question/comment or Overall Comments columns.",
    "self_assessments": "Expected resident names and [Question N], [Question N Response], [Question N Comment] columns, or question-per-row forms.",
    "milestones": "Expected resident_name, milestone (for example PC1), and value; include date or report_dates and select a specialty.",
    "procedures": "Expected Logger_Extract columns ResidentLastName, ResidentFirstName, DatePerformed, and Procedure.",
    "conferences": "Expected Person, Category, Status, and % Attended from the New Innovations conference report.",
    "duty_hours": "Expected Start, End, Hours, and Name, or the hierarchical Work Type report with resident headers.",
    "duty_violations": "Expected Name, Log Date, Rule, Description, and Comments from the New Innovations violation report.",
}


def parse_upload(filename: str, data: bytes, kind: str, specialty: str, period_start: str = "", period_end: str = "") -> ParseResult:
    """Return structured evidence and actionable validation messages.

    Nothing is saved here.  ``errors`` blocks acceptance of the affected file,
    even if some candidate records were extracted.  Missing event dates remain
    explicit; an administrator-selected end date is marked as a fallback.
    """
    result = ParseResult()
    if kind not in KINDS:
        result.errors.append("Choose a supported upload category.")
        return result
    if specialty not in SPECIALTIES:
        result.errors.append("Specialty must be Medicine, Pediatrics, or Combined.")
        return result
    if not isinstance(data, bytes) or not data:
        result.errors.append("The uploaded file is empty or unreadable.")
        return result
    if len(data) > MAX_FILE_BYTES:
        result.errors.append("The file exceeds 15 MB. Export a smaller reporting period or upload individual reports.")
        return result
    start, end = _date(period_start), _date(period_end)
    if (period_start and not start) or (period_end and not end) or (start and end and start > end):
        result.errors.append("Choose valid reporting dates with the start on or before the end.")
        return result
    suffix = PurePath(filename).suffix.lower()
    if suffix not in {".csv", ".tsv", ".tab", ".txt", ".xlsx", ".xls", ".pdf"}:
        result.errors.append("Upload a CSV, TSV, TAB, XLSX, XLS, or PDF report. Notebook/code and ZIP uploads are not accepted.")
        return result
    try:
        if data.lstrip().startswith(b"%PDF") or suffix == ".pdf":
            _parse_pdf(data, kind, specialty, start, end, result)
        else:
            tables = _read_tables(filename, data, result)
            recognized = 0
            for table in tables:
                if kind in {"evaluations", "self_assessments"}:
                    success = _parse_assessments(table, kind, specialty, start, end, result)
                elif kind == "milestones":
                    success = _parse_milestones(table, specialty, start, end, result)
                elif kind == "procedures":
                    success = _parse_procedures(table, specialty, start, end, result)
                elif kind == "conferences":
                    success = _parse_conferences(table, specialty, start, end, result)
                elif kind == "duty_hours":
                    success = _parse_duty(table, specialty, start, end, result)
                else:
                    success = _parse_violations(table, specialty, start, end, result)
                if not success and kind != "milestones":
                    success = _parse_generic(table, kind, specialty, start, end, result)
                if success:
                    recognized += 1
                else:
                    result.warnings.append(f"{table.label} did not match the selected report category and was not imported.")
            if not recognized:
                result.errors.append(_EXPECTED[kind])
    except (ImportError, ValueError, TypeError, KeyError, OSError, csv.Error, UnicodeError, zipfile.BadZipFile) as exc:
        # Controlled parser errors describe format or field names, never echo
        # raw file bytes, credentials, or resident-level rows.
        if isinstance(exc, ValueError) and str(exc).startswith(("The ", "This ", "No ", "Unlock ", "Legacy ")):
            result.errors.append(str(exc))
        else:
            result.errors.append("The report could not be read safely. Check the file format and headers, then try exporting it as CSV or XLSX. " + _EXPECTED[kind])
    except Exception:
        result.errors.append("This report layout is not yet supported. No data was saved. Try a CSV/XLSX export or the documented long CSV template.")
    if not result.records and not result.errors:
        result.errors.append("No resident records were found. Check the report date range and confirm it contains data rows, not only headings.")
    missing_date = sum(not record.get("date") for record in result.records)
    fallback = sum(record.get("date_is_fallback") and bool(record.get("date")) for record in result.records)
    if fallback:
        result.warnings.append(f"{fallback} record(s) had no source event/report date; the selected reporting period end was used and marked as a fallback. Confirm this date before saving.")
    if missing_date:
        result.errors.append(f"{missing_date} record(s) have no source date. Select the actual reporting period or add source dates before saving.")
    if any(record.get("date_needs_confirmation") for record in result.records):
        result.errors.append("Some source period start dates follow their end dates. Correct these dates before saving.")
    invalid_dates = sum(bool(record.get("source_date_needs_review")) for record in result.records)
    if invalid_dates:
        result.warnings.append(f"{invalid_dates} record(s) contain a nonblank source date that could not be read. Another valid source date or the selected period end was used where possible. Check the original dates before saving.")
    combined_evals = sum(record.get("category") == "evaluations" and record.get("specialty") == "Combined" for record in result.records)
    if combined_evals:
        result.warnings.append(f"{combined_evals} evaluation(s) are classified Combined. They remain separate from Medicine and Pediatrics evaluation counts. If these are specialty-specific, choose the correct specialty or add a specialty column before saving.")
    if kind == "milestones" and any(record.get("specialty") == "Combined" for record in result.records):
        result.errors.append("Milestone records require an explicit Medicine or Pediatrics specialty.")
    digest = hashlib.sha256(data).hexdigest()
    for record in result.records:
        record["file_sha256"] = digest
        record["source_name"] = PurePath(filename).name
    result.warnings = list(dict.fromkeys(result.warnings))
    result.errors = list(dict.fromkeys(result.errors))
    return result
