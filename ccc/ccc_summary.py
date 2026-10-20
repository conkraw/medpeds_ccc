"""CCC summaries in the program's supplied, editable table-based Word format.

The visual reference controls table order, sizes, colors, and evaluation page
breaks. No reference document, real resident data, or notebook execution is
needed at runtime. Clinical assessments are entered by the committee; numerical
changes are descriptive only. Every export is constructed in memory.

Pass active evidence for the roster, all saved reviews, and the annual entries
returned by ``annual.progress_for``. This module checks resident ownership and
dates independently. An explicitly selected Final review uses frozen data only;
a live export may use its latest review's narrative without freezing live data.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ccc import model as m
from ccc.evaluation_summary import evaluation_rows
from ccc.reports import _blank, _iso, _number, _text


TABLE_WIDTH = 7.6
BLUE = "D9EAF7"
GREEN = "BFEDD2"
RED = "F4CCCC"
GRAY = "EDEDED"
HEADER_GRAY = "D9D9D9"
MAIN_WIDTHS = [1.8, 2.0, 1.8, 2.0]
MILESTONE_DOMAINS = {
    "Medicine": [("PC", 6), ("MK", 3), ("SBP", 3), ("PBLI", 2), ("PRO", 4), ("ICS", 3)],
    "Pediatrics": [("PC", 5), ("MK", 2), ("SBP", 6), ("PBLI", 2), ("PRO", 4), ("ICS", 3)],
}
READINESS_FIELDS = {
    "Clinical Reasoning": "readiness_clinical_reasoning",
    "Team Leadership": "readiness_team_leadership",
    "Learner Supervision": "readiness_learner_supervision",
    "Communication": "readiness_communication",
    "Professionalism": "readiness_professionalism",
    "Systems-Based Practice": "readiness_systems_based_practice",
}


@dataclass
class _Context:
    resident: dict
    period_start: str
    period_end: str
    evidence: list[dict]
    cumulative: list[dict]
    history: list[dict]
    tracking: list[dict]
    narrative_review: dict | None
    explicit_review: bool
    final: bool
    limited_final_history: bool

    @property
    def fields(self) -> dict:
        fields = (self.narrative_review or {}).get("fields", {})
        return fields if isinstance(fields, dict) else {}

    @property
    def pgy(self) -> int:
        value = (self.narrative_review or {}).get("pgy") if self.explicit_review else None
        try:
            return min(4, max(1, int(value)))
        except (TypeError, ValueError):
            return min(4, max(1, m.current_pgy(self.resident, date.fromisoformat(self.period_end))))

    @property
    def half(self) -> str:
        return (self.narrative_review or {}).get("half", "Year end")


def _validate_date(value: Any, label: str) -> str:
    result = _iso(value)
    if not result:
        raise m.ValidationError(f"{label} must be a calendar date.")
    return result


def _owned(records: list[dict], resident_id: str, start: str, end: str) -> list[dict]:
    return [deepcopy(r) for r in records if isinstance(r, dict)
            and r.get("resident_id") == resident_id and r.get("active", True)
            and m.record_in_window(r, start, end)]


def _context(resident: dict, records: list[dict], reviews: list[dict], start: str,
             end: str, review: dict | None, tracking: list[dict] | None) -> _Context:
    resident_id = _text(resident.get("id"))
    if not resident_id:
        raise m.ValidationError("Choose a resident before preparing a CCC summary.")
    if review and review.get("resident_id") != resident_id:
        raise m.ValidationError("The selected review belongs to another resident.")
    final = bool(review and review.get("status") == "Final")
    if final:
        if not isinstance(review.get("snapshot"), list):
            raise m.ValidationError("This finalized review has no preserved evidence snapshot.")
        if not isinstance(review.get("snapshot_resident"), dict):
            raise m.ValidationError("This finalized review is missing its preserved resident details. Reopen and finalize it after confirming the resident and review data.")
        for key in ("cumulative_evidence_snapshot", "annual_progress_snapshot", "history_snapshot"):
            if key in review and not isinstance(review[key], list):
                raise m.ValidationError("This finalized review contains invalid preserved data. Review the saved record before exporting it.")
        start, end = review.get("period_start"), review.get("period_end")
    start = _validate_date(start, "Reporting start")
    end = _validate_date(end, "Evidence cutoff")
    if start > end:
        raise m.ValidationError("The reporting start must be on or before the evidence cutoff.")

    profile = deepcopy(resident)
    if final and isinstance(review.get("snapshot_resident"), dict):
        profile = deepcopy(review["snapshot_resident"])
        if profile.get("id") != resident_id:
            raise m.ValidationError("The preserved resident details belong to another resident.")
    training_start = _validate_date(profile.get("start_date"), "Training start")
    limited = final and not isinstance(review.get("cumulative_evidence_snapshot"), list)
    if final:
        evidence = _owned(review["snapshot"], resident_id, start, end)
        historical_records = review.get("cumulative_evidence_snapshot", review["snapshot"])
        cumulative = _owned(historical_records, resident_id, training_start, end)
        raw_history = review.get("history_snapshot", [])
        annual = review.get("annual_progress_snapshot", [])
    else:
        evidence = _owned(records, resident_id, start, end)
        cumulative = _owned(records, resident_id, training_start, end)
        raw_history = reviews
        annual = tracking or []
    history_cutoff = _iso((review or {}).get("review_date")) or end
    history = [deepcopy(r) for r in raw_history if isinstance(r, dict)
               and r.get("resident_id") == resident_id
               and _iso(r.get("review_date")) and _iso(r.get("review_date")) <= history_cutoff]
    if review:
        history = [r for r in history if r.get("id") != review.get("id")]
        history.append(deepcopy(review))
    history.sort(key=lambda r: (_iso(r.get("review_date")), str(r.get("id", ""))))
    narrative_review = deepcopy(review) if review else (history[-1] if history else None)
    annual = [deepcopy(r) for r in annual if isinstance(r, dict)
              and r.get("resident_id") == resident_id
              and _iso(r.get("as_of")) and _iso(r.get("as_of")) <= end]
    annual.sort(key=lambda r: (_iso(r.get("period_start")), _iso(r.get("as_of"))))
    latest_annual: dict[tuple, dict] = {}
    for entry in annual:
        latest_annual[(entry.get("pgy"), entry.get("academic_year"), entry.get("period_start"), entry.get("period_end"))] = entry
    return _Context(profile, start, end, evidence, cumulative, history,
                    list(latest_annual.values()), narrative_review, review is not None, final, limited)


def _fmt(value: Any, *, decimals: int | None = None) -> str:
    if _blank(value):
        return ""
    number = _number(value)
    if number is None:
        return _text(value)
    if decimals is not None:
        return f"{number:.{decimals}f}"
    return str(int(number)) if number.is_integer() else f"{number:g}"


def _date_display(value: Any) -> str:
    iso = _iso(value)
    if not iso:
        return _text(value)
    parsed = date.fromisoformat(iso)
    return f"{parsed.month}/{parsed.day}/{parsed.year}"


def _new_document() -> Any:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.2)
    section.left_margin = section.right_margin = Inches(0.2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(7)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1
    properties = doc.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "Medicine Pediatrics CCC Summary"
    properties.subject = "Clinical Competency Committee Summary"
    properties.comments = ""
    return doc


def _table(doc: Any, widths: list[float]) -> Any:
    table = doc.add_table(rows=0, cols=len(widths))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    pr = table._tbl.tblPr
    width = pr.find(qn("w:tblW"))
    width.set(qn("w:w"), str(round(sum(widths) * 1440)))
    width.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
            element.set(qn("w:" + key), value)
        borders.append(element)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, amount in (("top", 0), ("bottom", 0), ("left", 108), ("right", 108)):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:w"), str(amount))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    pr.append(margins)
    return table


def _cell(cell: Any, value: Any, *, size: float = 7, bold: bool = False,
          align: str = "center", fill: str | None = None, color: str = "000000") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                           "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(_text(value))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shading)


def _row(table: Any, values: list[Any], *, spans: list[int] | None = None,
         size: float = 7, bold: bool = False, align: str = "center",
         fill: str | None = None, color: str = "000000", repeat: bool = False,
         height: float = 0.2, keep: bool = False) -> Any:
    row = table.add_row()
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    widths = [column.width for column in table.columns]
    for cell, width in zip(row.cells, widths):
        cell.width = width
    spans = spans or [1] * len(values)
    if len(values) != len(spans) or sum(spans) != len(table.columns):
        raise ValueError("A report table row does not match its column grid.")
    index = 0
    for value, span in zip(values, spans):
        cell = row.cells[index]
        if span > 1:
            cell = cell.merge(row.cells[index + span - 1])
        _cell(cell, value, size=size, bold=bold, align=align, fill=fill, color=color)
        index += span
    if repeat:
        element = OxmlElement("w:tblHeader")
        element.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(element)
    if keep:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return row


def _title(table: Any, title: str, *, fill: str = BLUE, size: float = 7) -> None:
    _row(table, [title], spans=[len(table.columns)], bold=True, fill=fill, repeat=True, size=size)


def _single(doc: Any, title: str, value: Any, *, fill: str = BLUE,
            height: float = 0.2) -> None:
    table = _table(doc, [TABLE_WIDTH])
    _title(table, title, fill=fill)
    _row(table, [value], align="left", height=height)


def _page_break(doc: Any) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def _review_label(review: dict | None, default_pgy: int = 1) -> str:
    review = review or {}
    pgy = review.get("pgy", default_pgy)
    half = "MY" if review.get("half") == "Midyear" else "EOY"
    return f"PGY{pgy} {half}"


def _year_entries(ctx: _Context) -> list[dict]:
    """One dated annual value set per training year, with frozen legacy fallback."""
    legacy: dict[tuple, dict] = {}
    for review in ctx.history:
        pgy, academic = review.get("pgy"), review.get("academic_year")
        if not isinstance(pgy, int) or pgy not in range(1, 5):
            continue
        start = _iso(review.get("period_start"))
        if not start:
            try:
                year = int(str(academic).split("-")[0])
                start = date(year, 7, 1).isoformat()
            except (ValueError, TypeError):
                start = m.add_months(date.fromisoformat(ctx.resident["start_date"]), 12 * (pgy - 1)).isoformat()
        fields = deepcopy(review.get("fields", {}))
        legacy[(pgy, academic or m.academic_year(start))] = {
            "pgy": pgy, "academic_year": academic or m.academic_year(start),
            "period_start": start,
            "period_end": (m.add_months(date.fromisoformat(start), 12) - timedelta(days=1)).isoformat(),
            "as_of": _iso(review.get("period_end")) or _iso(review.get("review_date")),
            "fields": fields, "legacy": True,
        }
    tracked_years = {(entry.get("pgy"), entry.get("academic_year")) for entry in ctx.tracking}
    entries = {(entry.get("pgy"), entry.get("academic_year"), entry.get("period_start"), entry.get("period_end")): entry
               for key, entry in legacy.items() if key not in tracked_years}
    for entry in ctx.tracking:
        entries[(entry.get("pgy"), entry.get("academic_year"), entry.get("period_start"), entry.get("period_end"))] = deepcopy(entry)
    recorded_pgy = {key[0] for key in entries}
    for pgy in range(1, ctx.pgy + 1):
        if pgy not in recorded_pgy:
            start = m.add_months(date.fromisoformat(ctx.resident["start_date"]), 12 * (pgy - 1))
            if start.isoformat() <= ctx.period_end:
                entries[(pgy, m.academic_year(start), start.isoformat(), (m.add_months(start, 12) - timedelta(days=1)).isoformat())] = {
                    "pgy": pgy, "academic_year": m.academic_year(start), "period_start": start.isoformat(),
                    "period_end": (m.add_months(start, 12) - timedelta(days=1)).isoformat(), "fields": {}, "as_of": "",
                }
    return sorted(entries.values(), key=lambda e: (e.get("period_start", ""), e.get("pgy", 0)))


def _year_label(entry: dict, entries: list[dict]) -> str:
    pgy = entry.get("pgy")
    repeated = sum(e.get("pgy") == pgy for e in entries) > 1
    same_year = sum(e.get("pgy") == pgy and e.get("academic_year") == entry.get("academic_year") for e in entries) > 1
    label = f"PGY{pgy}" + (f"\n{entry.get('academic_year', '')}" if repeated else "")
    if same_year:
        label += f"\n{entry.get('period_start', '')} to {entry.get('period_end', '')}"
    return label


def _record_year(record: dict, ctx: _Context, entries: list[dict]) -> tuple[int, str]:
    observed = _iso(record.get("date")) or _iso(record.get("period_end"))
    for entry in reversed(entries):
        first, last = _iso(entry.get("period_start")), _iso(entry.get("period_end"))
        if first and last and first <= observed <= last:
            return int(entry["pgy"]), str(entry.get("academic_year", m.academic_year(first)))
    when = date.fromisoformat(observed)
    pgy = min(4, max(1, m.current_pgy(ctx.resident, when)))
    return pgy, m.academic_year(when)


def _identity(doc: Any, ctx: _Context) -> None:
    table = _table(doc, MAIN_WIDTHS)
    half = "MID-YEAR" if ctx.half == "Midyear" else "END-OF-YEAR"
    _title(table, f"MEDICINE PEDIATRICS {half} SUMMARY")
    period = f"{ctx.period_start} to {ctx.period_end}"
    if ctx.limited_final_history:
        period += "\nHistorical data limited to this saved review"
    rows = [
        ["RESIDENT NAME:", ctx.resident.get("name", ""), "TRAINING LEVEL:", f"R{ctx.pgy}"],
        ["COACH NAME:", ctx.fields.get("coach_name", ctx.resident.get("coach_name", "")), "REVIEW PERIOD:", period],
        ["CCC MEETING DATE:", (ctx.narrative_review or {}).get("review_date", ""), "GRADUATION DATE:", ctx.resident.get("graduation_date", "")],
    ]
    for values in rows:
        row = _row(table, values)
        for i in (0, 2):
            _cell(row.cells[i], values[i], bold=True)
    rotations = ctx.fields.get("rotation_summary", "")
    if _blank(rotations):
        rotations = "; ".join(dict.fromkeys(_text(r.get("rotation")) for r in ctx.evidence if _text(r.get("rotation"))))
    _single(doc, "ROTATIONS SINCE LAST REVIEW", rotations, height=0.4)


def _clinic(doc: Any, entries: list[dict]) -> None:
    table = _table(doc, MAIN_WIDTHS)
    _title(table, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    for entry in entries or [{"pgy": 1, "fields": {}}]:
        fields = entry.get("fields", {})
        label = _year_label(entry, entries)
        values = [f"ADULT {label}", _fmt(fields.get("adult_clinic_count")),
                  f"PEDIATRIC {label}", _fmt(fields.get("pediatric_clinic_count"))]
        row = _row(table, values)
        for i in (0, 2):
            _cell(row.cells[i], values[i], bold=True)


def _procedures(doc: Any, ctx: _Context) -> None:
    title = _table(doc, [TABLE_WIDTH])
    _title(title, "PROCEDURES LOGGED TO DATE")
    table = _table(doc, [TABLE_WIDTH / 10] * 10)
    _row(table, ["PROCEDURE", "LAST DATE LOGGED", "NUMBER LOGGED"], spans=[4, 3, 3],
         size=6, bold=True, fill=HEADER_GRAY, repeat=True)
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for record in ctx.cumulative:
        if record.get("category") == "procedures":
            groups[(_text(record.get("specialty")), _text(record.get("metric")), record.get("record_mode", "event"))].append(record)
    items = sorted(groups.items(), key=lambda pair: max(r.get("date", "") for r in pair[1]), reverse=True)
    for (specialty, name, mode), records in items:
        if mode == "snapshot":
            for record in records:
                value = _fmt(record.get("value"))
                label = f"{name} ({specialty}; snapshot)"
                _row(table, [label, _date_display(record.get("date")), value], spans=[4, 3, 3], size=6)
        else:
            name_specialties = {s for s, n, _ in groups if n == name}
            label = name + (f" ({specialty})" if len(name_specialties) > 1 else "")
            _row(table, [label, _date_display(max(r.get("date", "") for r in records)), str(len(records))],
                 spans=[4, 3, 3], size=6)
        notes = [_text(r.get("text")) for r in records if _text(r.get("text"))]
        if notes:
            _row(table, ["\n".join(notes)], spans=[10], size=6, align="left")
    if not items:
        _row(table, ["", "", ""], spans=[4, 3, 3], size=6)


def _conferences(doc: Any, ctx: _Context, entries: list[dict]) -> None:
    title = _table(doc, [TABLE_WIDTH])
    _title(title, "CONFERENCE ATTENDANCE")
    years = [(int(e["pgy"]), str(e.get("academic_year", ""))) for e in entries]
    grouped: dict[tuple, dict[tuple, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for record in ctx.cumulative:
        if record.get("category") != "conferences":
            continue
        year = _record_year(record, ctx, entries)
        if year not in years:
            years.append(year)
        grouped[(_text(record.get("specialty")), _text(record.get("metric")))][year].append(record)
    years = sorted(set(years), key=lambda y: (y[1], y[0])) or [(ctx.pgy, "")]
    # The reference uses six grid units for category and one for each PGY.
    table = _table(doc, [TABLE_WIDTH / (6 + len(years))] * (6 + len(years)))
    repeated = Counter(year[0] for year in years)
    labels = [f"PGY{pgy}" + (f"\n{academic}" if repeated[pgy] > 1 else "") for pgy, academic in years]
    _row(table, ["CATEGORY", *labels], spans=[6] + [1] * len(years), size=6, bold=True, fill=HEADER_GRAY, repeat=True)
    for (specialty, metric), periods in sorted(grouped.items()):
        cells = []
        for year in years:
            values = []
            for record in sorted(periods.get(year, []), key=lambda r: (r.get("period_start", ""), r.get("date", ""))):
                unit = _text(record.get("unit"))
                percentage = unit.casefold() in {"%", "percent", "percentage"}
                value = _fmt(record.get("value"), decimals=2 if percentage else None)
                if _number(record.get("value")) is not None and percentage:
                    value += "%"
                elif value and unit:
                    value += " " + unit
                elif not value:
                    value = _text(record.get("values", {}).get("status"))
                source_start, source_end = _iso(record.get("period_start")), _iso(record.get("period_end"))
                annual_entry = next((entry for entry in entries
                                     if (entry.get("pgy"), entry.get("academic_year")) == year), None)
                complete_year = bool(annual_entry and source_start == _iso(annual_entry.get("period_start"))
                                     and source_end == _iso(annual_entry.get("period_end"))
                                     and source_end == (m.add_months(date.fromisoformat(source_start), 12) - timedelta(days=1)).isoformat())
                if source_start and source_end and not complete_year:
                    value += f"\n{source_start} to {source_end}"
                elif not source_start and not source_end and _iso(record.get("date")):
                    value += f"\nAs of {record['date']}"
                values.append(value)
            cells.append("\n\n".join(values))
        label = f"{specialty} :: {metric}" if specialty else metric
        _row(table, [label, *cells], spans=[6] + [1] * len(years), size=6)
    if not grouped:
        _row(table, [""] * (1 + len(years)), spans=[6] + [1] * len(years), size=6)


def _duty_hours(doc: Any, ctx: _Context) -> None:
    table = _table(doc, MAIN_WIDTHS)
    _title(table, "DUTY HOURS SUMMARY")
    _row(table, ["Review Period", "Rotation", "Violation", "Committee Notes / Resolution"], bold=True, fill=GRAY, repeat=True)
    violations = [r for r in ctx.evidence if r.get("category") == "duty_violations"]
    for record in sorted(violations, key=lambda r: r.get("date", "")):
        values = record.get("values", {})
        note_parts = [_text(record.get("text"))]
        for key in ("description", "comments"):
            value = _text(values.get(key))
            if value and not any(value in part for part in note_parts):
                note_parts.append(value)
        manual = _text(ctx.fields.get("duty_resolution_notes"))
        if manual:
            note_parts.append("Committee: " + manual)
        _row(table, [_review_label(ctx.narrative_review, ctx.pgy) + "\n" + _date_display(record.get("date")),
                     record.get("rotation", values.get("training_location", "")),
                     record.get("metric", values.get("rule", "")), "\n".join(p for p in note_parts if p)], size=6)
    if not violations:
        _row(table, ["", "", "", ctx.fields.get("duty_resolution_notes", "")], size=6)
    deep = _table(doc, [1.7, 2.95, 2.95])
    _title(deep, "DUTY HOUR DEEP DIVE")
    latest = max((r.get("date", "") for r in ctx.evidence if r.get("category") == "duty_hours"), default="")
    row = _row(deep, ["LAST DUTY HOUR LOG:", ctx.fields.get("duty_last_log", latest)], spans=[1, 2], size=6)
    _cell(row.cells[0], "LAST DUTY HOUR LOG:", size=6, bold=True, fill=GRAY)
    _row(deep, ["MINIMUM DAYS OFF", "DAY OFF WINDOW", "ROTATIONS"], size=6, bold=True, fill="000000", color="FFFFFF")
    _row(deep, [ctx.fields.get(key, "") for key in ("duty_min_days_off", "duty_days_off_window", "duty_days_off_rotations")], size=6)
    _row(deep, ["MAX HOURS IN 4 WEEK PERIOD", "80 HOUR WINDOW", "ROTATIONS"], size=6, bold=True, fill="000000", color="FFFFFF")
    _row(deep, [ctx.fields.get(key, "") for key in ("duty_max_hours_4weeks", "duty_80_hour_window", "duty_80_hour_rotations")], size=6)


def _milestone_groups(ctx: _Context, specialty: str, entries: list[dict]) -> list[dict]:
    groups: dict[tuple, dict] = {}
    for record in ctx.cumulative:
        if record.get("category") != "milestones" or record.get("specialty") != specialty:
            continue
        key = (record.get("date", ""), record.get("period_start", ""), record.get("period_end", ""), record.get("form", ""))
        if key not in groups:
            pgy, academic = _record_year(record, ctx, entries)
            entry = next((e for e in entries if e.get("pgy") == pgy and e.get("academic_year") == academic), None)
            first = _iso((entry or {}).get("period_start"))
            if not first:
                first = m.add_months(date.fromisoformat(ctx.resident["start_date"]), 12 * (pgy - 1)).isoformat()
            midpoint = m.add_months(date.fromisoformat(first), 6).isoformat()
            half = "MY" if key[0] <= midpoint else "EOY"
            groups[key] = {"label": f"PGY{pgy} {half}", "date": key[0], "academic_year": academic,
                           "form": key[3], "scores": defaultdict(list), "notes": []}
        metric = re.sub(r"\s+", "", _text(record.get("metric"))).upper()
        metric = re.sub(r"^PROF(?=\d)", "PRO", metric)
        status = record.get("values", {}).get("assessment_status", "")
        value = record.get("value") if not _blank(record.get("value")) else status
        groups[key]["scores"][metric].append(value)
        if _text(record.get("text")):
            groups[key]["notes"].append(f"{metric}: {_text(record['text'])}")
    output = [groups[key] for key in sorted(groups)]
    labels = Counter(g["label"] for g in output)
    for group in output:
        if labels[group["label"]] > 1:
            group["label"] += "\n" + group["date"]
    return output


def _milestones(doc: Any, ctx: _Context, specialty: str, entries: list[dict]) -> None:
    domains = MILESTONE_DOMAINS[specialty]
    codes = [f"{domain}{number}" for domain, count in domains for number in range(1, count + 1)]
    groups = _milestone_groups(ctx, specialty, entries)
    # Unrecognized imported codes remain visible in additional columns.
    extra_codes = sorted({code for group in groups for code in group["scores"]} - set(codes))
    codes += extra_codes
    if extra_codes:
        domains = [*domains, ("OTHER", len(extra_codes))]
    table = _table(doc, [0.6] + [(TABLE_WIDTH - 0.6) / len(codes)] * len(codes))
    title = "MEDICINE" if specialty == "Medicine" else "PEDIATRIC"
    _title(table, f"{title} MILESTONES")
    _row(table, [""] + [domain for domain, _ in domains], spans=[1] + [count for _, count in domains],
         size=6, bold=True, fill=GRAY, repeat=True)
    number_values = [""] + [str(number) for domain, count in MILESTONE_DOMAINS[specialty] for number in range(1, count + 1)] + extra_codes
    row = _row(table, number_values, size=6, bold=True, fill="000000", color="FFFFFF", repeat=True)
    _cell(row.cells[0], "", size=6, fill="FFFFFF")
    for group in groups:
        values = ["\n".join(_fmt(v, decimals=2) if _number(v) is not None else _text(v)
                             for v in group["scores"].get(code, [])) for code in codes]
        _row(table, [group["label"], *values], size=6, bold=True)
        if group["notes"]:
            _row(table, ["\n".join(group["notes"])], spans=[1 + len(codes)], size=6, align="left")
    if not groups:
        _row(table, [""] * (1 + len(codes)), size=6)

    period_count = max(1, len(groups))
    trend = _table(doc, [0.75] + [4.65 / period_count] * period_count + [0.65, 1.55])
    _title(trend, f"{title} MILESTONE TREND SUMMARY")
    _row(trend, ["Domain", *([g["label"] for g in groups] or [""]), "Change", "Trend"],
         size=6, bold=True, fill=GRAY, repeat=True)
    for domain, count in MILESTONE_DOMAINS[specialty]:
        means = []
        for group in groups:
            scores = []
            conflict = False
            for number in range(1, count + 1):
                values = group["scores"].get(f"{domain}{number}", [])
                numeric = [_number(value) for value in values]
                numeric = [value for value in numeric if value is not None]
                # Conflicting scores are shown above; do not silently average revisions.
                if len(set(numeric)) > 1:
                    conflict = True
                elif numeric:
                    scores.append(numeric[0])
            means.append(None if conflict or not scores else sum(scores) / len(scores))
        present = [score for score in means if score is not None]
        change = present[-1] - present[0] if len(present) >= 2 else None
        change_label = "" if change is None else f"{change:+.2f}"
        direction = "" if change is None else ("Increase" if change > 0.00001 else "Decrease" if change < -0.00001 else "No change")
        manual = ctx.fields.get(f"{specialty.lower()}_milestone_trend_{domain.lower()}", "")
        display = manual or direction
        _row(trend, [domain, *([_fmt(value, decimals=2) for value in means] or [""]), change_label, display],
             size=6, bold=True)


def _readiness(doc: Any, ctx: _Context) -> None:
    table = _table(doc, [2.2, 5.4])
    _title(table, "PRACTICE READINESS AND DEVELOPMENT SNAPSHOT")
    for label, key in READINESS_FIELDS.items():
        row = _row(table, [label, ctx.fields.get(key, "")])
        _cell(row.cells[0], label, bold=True, fill=GRAY)


def _encounters(doc: Any, entries: list[dict]) -> None:
    entries = entries or [{"pgy": 1, "fields": {}}]
    # The four ordinary training years remain columns, with extra columns only
    # for a repeated PGY so an earlier year is never overwritten.
    columns = list(entries)
    present = {e.get("pgy") for e in columns}
    columns.extend({"pgy": pgy, "fields": {}} for pgy in range(1, 5) if pgy not in present)
    columns.sort(key=lambda e: (e.get("pgy", 0), e.get("academic_year", "")))
    table = _table(doc, [1.5] + [6.1 / len(columns)] * len(columns))
    _title(table, "OBSERVED ENCOUNTERS COMPLETION")
    _row(table, ["NUMBER", *[_year_label(entry, columns) for entry in columns]], bold=True, fill=GRAY, repeat=True)
    logged = [number for number in range(1, 5)
              if any(entry.get("fields", {}).get(f"observed_encounter_{number}") is not None for entry in columns)]
    for number in logged or [1]:
        values = []
        for entry in columns:
            value = entry.get("fields", {}).get(f"observed_encounter_{number}")
            values.append("X" if value is True else "Not completed" if value is False else "")
        _row(table, [str(number), *values], bold=True)
    legacy = [entry for entry in columns if _text(entry.get("fields", {}).get("observed_encounters_completed"))]
    if legacy:
        note = "\n".join(f"{_year_label(entry, columns)} recorded total: {entry['fields']['observed_encounters_completed']}"
                         for entry in legacy)
        _row(table, [note], spans=[1 + len(columns)], align="left")


def _ite_value(fields: dict, specialty: str) -> str:
    if specialty == "Medicine":
        score = _fmt(fields.get("medicine_ite_score"))
        percentile = _fmt(fields.get("medicine_ite_percentile"))
        return score + (f" ({percentile}%tile)" if percentile else "")
    score = _fmt(fields.get("pediatrics_ite_score"))
    expected = _fmt(fields.get("pediatrics_ite_expected_pass_rate"))
    old_percentile = _fmt(fields.get("pediatrics_ite_percentile"))
    if expected:
        score += f" ({expected}%_PR)"
    if old_percentile:
        score += f" ({old_percentile}%tile)"
    return score


def _ite(doc: Any, entries: list[dict]) -> None:
    table = _table(doc, [0.55, 3.525, 3.525])
    _title(table, "IN-TRAINING EXAM SCORES")
    _row(table, ["", "MED IN-TRAINING EXAM SCORE:", "PED IN-TRAINING EXAM SCORE:"], bold=True, fill=GRAY, repeat=True)
    rows = [(entry, _ite_value(entry.get("fields", {}), "Medicine"), _ite_value(entry.get("fields", {}), "Pediatrics"))
            for entry in entries]
    rows = [row for row in rows if row[1] or row[2]]
    for entry, med, ped in rows:
        _row(table, [_year_label(entry, entries), med, ped], bold=True)
    if not rows:
        _row(table, ["PGY1", "", ""], bold=True)


def _history_values(ctx: _Context, field: str) -> list[str]:
    return [_text(review.get("fields", {}).get(field)) for review in ctx.history
            if _text(review.get("fields", {}).get(field))]


def _scholarship(doc: Any, ctx: _Context) -> None:
    table = _table(doc, [TABLE_WIDTH])
    _title(table, "SCHOLARLY ACTIVITIES")
    for label, key in (("ABSTRACT SUBMISSIONS/ACCEPTANCES", "abstracts"), ("PUBLICATIONS", "publications")):
        _row(table, [label], bold=True, fill=GRAY)
        for value in _history_values(ctx, key) or [""]:
            _row(table, [value], bold=True)
    advocacy = _table(doc, [TABLE_WIDTH])
    _title(advocacy, "ADVOCACY")
    for value in _history_values(ctx, "advocacy") or [""]:
        _row(advocacy, [value], bold=True, align="left")


def _narratives(doc: Any, ctx: _Context) -> None:
    fields = ctx.fields
    _single(doc, "360 EVALUATION FROM NURSING", fields.get("nursing_feedback", ""))
    overall = _text(fields.get("overall_impression"))
    if _text(fields.get("strengths")):
        overall += ("\n\n" if overall else "") + "Strengths\n" + _text(fields["strengths"])
    _single(doc, "OVERALL IMPRESSION", overall, fill=GREEN)
    _single(doc, "AREAS FOR GROWTH", fields.get("areas_for_growth", ""), fill=RED)
    _single(doc, "PATTERNS OVER TIME", fields.get("patterns_over_time", ""))
    committee = [_text(fields.get("ccc_summary"))]
    for key, label in (("ccc_decision", "CCC assessment or decision"), ("action_plan", "Action plan"),
                       ("mentor", "Mentor or responsible person"), ("follow_up_date", "Follow-up"),
                       ("next_review_date", "Next review"), ("resident_comments", "Resident comments"),
                       ("observed_encounters_notes", "Observed encounter notes")):
        if _text(fields.get(key)):
            committee.append(label + "\n" + _text(fields[key]))
    _single(doc, "MEDICINE PEDIATRICS CLINICAL COMPETENCY COMMITTEE REVIEW", "\n\n".join(value for value in committee if value))
    header = _table(doc, [TABLE_WIDTH])
    _title(header, "SIGNATURES")
    signature = _table(doc, [2.6, 2.6, 2.4])
    for role in ("Resident", "Program Director"):
        row = _row(signature, [f"{role} Signature:", "_____________________________________________",
                               "Date: ____________________________________"], height=0.3, align="left")
        for i in (0, 2):
            _cell(row.cells[i], row.cells[i].text, bold=True, align="left")


def _evaluations(doc: Any, records: list[dict], title: str, *, fill: str = BLUE) -> None:
    _page_break(doc)
    table = _table(doc, [TABLE_WIDTH / 10] * 10)
    _title(table, title, fill=fill, size=6)
    rows = evaluation_rows(records)
    for entry in rows:
        metadata = _row(table, [_date_display(entry.get("date")), entry.get("rotation", ""),
                               entry.get("evaluator", "")], spans=[2, 2, 6], size=6, keep=True)
        _cell(metadata.cells[0], _date_display(entry.get("date")), size=6, align="left")
        narrative = _text(entry.get("narrative"))
        if narrative:
            # Long narrative rows are deliberately allowed to continue across
            # pages. No fixed heights or comment limits can truncate evidence.
            for cell in {id(cell._tc): cell for cell in metadata.cells}.values():
                cell.paragraphs[0].paragraph_format.keep_with_next = True
            _row(table, [narrative], spans=[10], size=6, align="left")
    if not rows:
        _row(table, [""], spans=[10], size=6, align="left")


def _apply_pagination(doc: Any) -> None:
    """Keep section headings attached without fixing the height of narratives."""
    for table_index, table in enumerate(doc.tables):
        for row in table.rows:
            repeated = bool(row._tr.xpath("./w:trPr/w:tblHeader"))
            if repeated:
                for cell in {id(c._tc): c for c in row.cells}.values():
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
        if table_index == 0:
            # This is the report title, not a continuation column heading.
            # Suppress an extra title when an office renderer joins adjacent
            # tables for pagination, while keeping the first data row with it.
            for element in table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
                element.getparent().remove(element)
        if table.cell(0, 0).text not in {"MEDICINE MILESTONES", "PEDIATRIC MILESTONES"}:
            continue
        # A usual MY/EOY matrix has only a few rows. Keep that short table
        # intact, while allowing a longer longitudinal record to paginate.
        text_size = sum(len(c.text) for row in table.rows for c in {id(cell._tc): cell for cell in row.cells}.values())
        if len(table.rows) <= 10 and text_size < 4000:
            for row in table.rows[:-1]:
                for cell in {id(c._tc): c for c in row.cells}.values():
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True


def build_ccc_summary(
    resident: dict,
    records: list[dict],
    reviews: list[dict],
    *,
    period_start: str,
    period_end: str,
    review: dict | None = None,
    program_name: str = "Medicine Pediatrics Residency",
    tracking: list[dict] | None = None,
) -> bytes:
    """Return an ordinary editable DOCX matching the supplied CCC Summary.

    ``review=None`` means live evidence for the requested dates. The most recent
    saved review through the cutoff supplies narrative fields only. An explicit
    Final review always uses its saved dates, resident profile, current-period
    snapshot, and (when available) cumulative/annual snapshots. Missing fields
    remain blank; missing frozen evidence is never filled from live records.
    """
    ctx = _context(resident, records, reviews, period_start, period_end, review, tracking)
    doc = _new_document()
    doc.core_properties.subject = _text(program_name) + " Clinical Competency Committee Summary"
    entries = _year_entries(ctx)
    _identity(doc, ctx)
    _clinic(doc, entries)
    _procedures(doc, ctx)
    _conferences(doc, ctx, entries)
    _duty_hours(doc, ctx)
    _milestones(doc, ctx, "Medicine", entries)
    _milestones(doc, ctx, "Pediatrics", entries)
    _readiness(doc, ctx)
    _encounters(doc, entries)
    _ite(doc, entries)
    _scholarship(doc, ctx)
    _narratives(doc, ctx)
    _evaluations(doc, [r for r in ctx.evidence if r.get("category") == "evaluations" and r.get("specialty") == "Medicine"],
                 "MEDICINE EVALUATION SUMMARY")
    _evaluations(doc, [r for r in ctx.evidence if r.get("category") == "evaluations" and r.get("specialty") == "Pediatrics"],
                 "PEDIATRIC EVALUATION SUMMARY", fill=GREEN)
    combined = [r for r in ctx.evidence if r.get("category") == "evaluations" and r.get("specialty") == "Combined"]
    if combined:
        _evaluations(doc, combined, "COMBINED EVALUATION SUMMARY")
    self_assessments = [r for r in ctx.evidence if r.get("category") == "self_assessments"]
    if self_assessments:
        _evaluations(doc, self_assessments, "RESIDENT SELF-ASSESSMENT")
    _apply_pagination(doc)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
