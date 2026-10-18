"""Editable Word exports built in memory from reviewed resident evidence.

The export layer does not infer entrustment, readiness, promotion, or compliance.
Final reviews use their saved evidence and review period; an export never silently
substitutes the current evidence for a missing final-review snapshot.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from html import unescape
from html.parser import HTMLParser
from io import BytesIO
import math
import re
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DEFAULT_PROGRAM_NAME = "Medicine Pediatrics Residency"
CONTENT_WIDTH = 7.1

REVIEW_FIELD_LABELS = {
    "coach_name": "Coach",
    "rotation_summary": "Rotations during the review period",
    "adult_clinic_count": "Adult clinic patients recorded",
    "pediatric_clinic_count": "Pediatric clinic patients recorded",
    "observed_encounters_completed": "Observed encounters recorded as completed",
    "observed_encounters_notes": "Observed encounter notes",
    "medicine_ite_score": "Medicine in training examination score",
    "medicine_ite_percentile": "Medicine in training examination percentile",
    "pediatrics_ite_score": "Pediatrics in training examination score",
    "pediatrics_ite_percentile": "Pediatrics in training examination percentile",
    "usmle_step1": "USMLE Step 1",
    "usmle_step2ck": "USMLE Step 2 CK",
    "usmle_step3": "USMLE Step 3",
    "abstracts": "Abstracts",
    "publications": "Publications",
    "advocacy": "Advocacy",
    "nursing_feedback": "Nursing feedback",
    "overall_impression": "Overall impression",
    "strengths": "Strengths",
    "areas_for_growth": "Areas for growth",
    "patterns_over_time": "Patterns over time",
    "ccc_summary": "Committee summary",
    "ccc_decision": "Recorded committee decision",
    "action_plan": "Action plan",
    "mentor": "Mentor or responsible person",
    "follow_up_date": "Follow up date",
    "next_review_date": "Next review date",
    "resident_comments": "Resident comments",
}

CATEGORY_LABELS = {
    "evaluations": "Evaluations",
    "self_assessments": "Self assessments",
    "milestones": "Milestone scores",
    "procedures": "Procedure entries",
    "conferences": "Attendance snapshots",
    "duty_hours": "Duty hour entries",
    "duty_violations": "Reported duty hour violations",
}


class _PlainHTML(HTMLParser):
    """Extract text without fetching resources or carrying HTML into the DOCX."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in {"script", "style"}:
            self.hidden += 1
        elif not self.hidden and tag in {"p", "div", "br", "li", "tr"}:
            self.parts.append("\n")
        elif not self.hidden and tag in {"td", "th"}:
            self.parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden and tag in {"p", "div", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, value: str) -> None:
        if not self.hidden:
            self.parts.append(value)


def _blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and not math.isfinite(value):
        return True
    return isinstance(value, str) and not value.strip()


def _text(value: Any) -> str:
    if _blank(value):
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple)):
        return "; ".join(_text(item) for item in value if not _blank(item))
    if isinstance(value, dict):
        return "; ".join(f"{_text(k)}: {_text(v)}" for k, v in value.items())
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    if re.search(r"</?(?:p|br|div|li|table|tr|td|th|span|b|strong|i|em|ul|ol|script|style)\b", text, re.I):
        parser = _PlainHTML()
        parser.feed(text)
        parser.close()
        text = "".join(parser.parts)
    text = unescape(text)
    # XML 1.0 cannot contain these control/surrogate/noncharacter code points.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _number(value: Any) -> float | None:
    if _blank(value) or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _display(value: Any, missing: str = "Not recorded") -> str:
    if _blank(value):
        return missing
    number = _number(value)
    if number is not None:
        return str(int(number)) if number.is_integer() else f"{number:g}"
    return _text(value) or missing


def _iso(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text:
        return ""
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return ""


def _date(value: Any) -> str:
    iso = _iso(value)
    return date.fromisoformat(iso).strftime("%b %d, %Y") if iso else "Not recorded"


def _date_or_text(value: Any) -> str:
    """Keep an entered follow-up instruction when it is not a calendar date."""
    return _date(value) if _iso(value) else (_text(value) or "Not recorded")


def _timestamp(value: Any) -> str:
    try:
        parsed = datetime.fromisoformat(_text(value).replace("Z", "+00:00"))
        if parsed.tzinfo:
            parsed = parsed.astimezone(timezone.utc)
            return parsed.strftime("%b %d, %Y at %H:%M UTC")
        return parsed.strftime("%b %d, %Y at %H:%M")
    except ValueError:
        return _text(value)


def _period(start: Any, end: Any) -> str:
    if not _iso(start) and not _iso(end):
        return "Not recorded"
    if _iso(start) == _iso(end):
        return _date(start)
    return f"{_date(start)} to {_date(end)}"


def _record_date(record: dict) -> str:
    return _iso(record.get("date")) or _iso(record.get("period_end"))


def _record_period(record: dict) -> str:
    if _iso(record.get("period_start")) or _iso(record.get("period_end")):
        return _period(record.get("period_start"), record.get("period_end"))
    return _date(record.get("date"))


def _record_value(record: dict) -> str:
    if _blank(record.get("value")):
        values = record.get("values") if isinstance(record.get("values"), dict) else {}
        return _text(values.get("assessment_status")) or "Not recorded"
    text = _display(record.get("value"))
    unit = _text(record.get("unit"))
    if text == "Not recorded" or not unit:
        return text
    return f"{text}%" if unit == "%" and not text.endswith("%") else f"{text} {unit}"


def _source(record: dict) -> str:
    return " | ".join(_text(record.get(key)) for key in ["source_name", "source_locator"] if _text(record.get(key))) or "Not recorded"


def _fields(review: dict | None) -> dict:
    return review.get("fields", {}) if review and isinstance(review.get("fields"), dict) else {}


def _checkpoint(review: dict | None) -> str:
    if not review:
        return "Review not selected"
    pgy = review.get("pgy")
    half = _text(review.get("half"))
    return " ".join(part for part in [f"PGY{pgy}" if pgy else "", half] if part) or "Review"


def _validate_resident(resident: dict, review: dict | None) -> str:
    resident_id = _text(resident.get("id"))
    if not resident_id:
        raise ValueError("A resident ID is required to export a report.")
    if review and _text(review.get("resident_id")) != resident_id:
        raise ValueError("The selected review does not belong to this resident.")
    return resident_id


def _select_records(
    resident: dict, records: list[dict], review: dict | None, start: str, end: str
) -> tuple[list[dict], list[dict], str, str]:
    from .model import record_in_window

    resident_id = _validate_resident(resident, review)
    final = bool(review and str(review.get("status", "")).casefold() == "final")
    if final:
        if not isinstance(review.get("snapshot"), list):
            raise ValueError("This finalized review has no saved evidence. Restore its saved version before exporting.")
        records = review["snapshot"]
        start, end = review.get("period_start", ""), review.get("period_end", "")
    start, end = _iso(start), _iso(end)
    if not start or not end or start > end:
        raise ValueError("Choose a valid report start date and end date.")
    own = [record for record in records if isinstance(record, dict) and _text(record.get("resident_id")) == resident_id]
    included = [record for record in own if record_in_window(record, start, end)]
    return included, own, start, end


def _set_black_font(style: Any, size: float, bold: bool = False) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        rfonts.set(qn(f"w:{key}"), "Arial")
    for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        rfonts.attrib.pop(qn(f"w:{key}"), None)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for key in ["themeColor", "themeTint", "themeShade"]:
            color.attrib.pop(qn(f"w:{key}"), None)
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)


def _new_document(title: str, program_name: str) -> Any:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.header_distance = section.footer_distance = Inches(0.3)
    for style, size, bold in [("Normal", 11, False), ("Title", 21, True), ("Subtitle", 11, False), ("Heading 1", 14, True), ("Heading 2", 11.5, True), ("Heading 3", 11, True)]:
        _set_black_font(doc.styles[style], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.1
    normal.widow_control = True
    for style in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style].paragraph_format.keep_with_next = True
        doc.styles[style].paragraph_format.space_before = Pt(11)
        doc.styles[style].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].paragraph_format.space_after = Pt(4)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(10)
    doc.core_properties.title = title
    doc.core_properties.subject = "Resident educational review"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(_text(program_name) or DEFAULT_PROGRAM_NAME, "Subtitle")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Confidential educational record | Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(70, 70, 70)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_props = OxmlElement("w:rPr")
    for tag, attributes in [("rFonts", {"ascii": "Arial", "hAnsi": "Arial"}), ("sz", {"val": "16"}), ("color", {"val": "464646"})]:
        item = OxmlElement(f"w:{tag}")
        for key, value in attributes.items():
            item.set(qn(f"w:{key}"), value)
        field_props.append(item)
    field_run.append(field_props)
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)
    return doc


def _paragraphs(doc: Any, value: Any, *, style: str | None = None) -> None:
    text = _text(value)
    if not text:
        return
    for part in re.split(r"\n\s*\n", text):
        if part.strip():
            doc.add_paragraph(part.strip(), style=style)


def _labeled(doc: Any, label: str, value: Any, *, missing: bool = False, keep_with_next: bool = False) -> None:
    text = _text(value)
    if not text and not missing:
        return
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep_with_next
    p.add_run(label + ": ").bold = True
    p.add_run(text or "Not recorded")


def _heading_content(doc: Any, title: str, value: Any, *, level: int = 2) -> None:
    if not _text(value):
        return
    doc.add_heading(title, level)
    _paragraphs(doc, value)


def _table(doc: Any, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[float] | None = None) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = widths or [CONTENT_WIDTH / len(headers)] * len(headers)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{edge}")
        for name, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            border.set(qn(f"w:{name}"), value)
        borders.append(border)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ["top", "bottom", "left", "right"]:
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), "80")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    pr.append(margins)
    header_repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    all_rows = [headers] + [list(row) for row in rows]
    for row_index, values in enumerate(all_rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = row_index == 0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(_text(value))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "29465B" if row_index == 0 else ("F3F6F8" if row_index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _review_section(doc: Any, review: dict | None) -> None:
    doc.add_heading("Committee review", 1)
    if not review:
        doc.add_paragraph("No committee review has been selected for this report.")
        return
    fields = _fields(review)
    status = _text(review.get("status")) or "Draft"
    _labeled(doc, "Review status", status)
    _labeled(doc, "Committee meeting date", _date(review.get("review_date")))
    if status.casefold() == "final":
        if review.get("finalized_at"):
            _labeled(doc, "Finalized on", _timestamp(review["finalized_at"]))
        _labeled(doc, "Finalized by", review.get("finalized_by"))
    if review.get("version") is not None:
        _labeled(doc, "Review version", review["version"])
    for key in ["ccc_decision", "ccc_summary", "overall_impression", "strengths", "areas_for_growth", "patterns_over_time", "action_plan"]:
        _heading_content(doc, REVIEW_FIELD_LABELS[key], fields.get(key))
    if not any(_text(fields.get(key)) for key in ["ccc_summary", "overall_impression", "ccc_decision"]):
        doc.add_paragraph("A committee summary or decision has not been entered.")
    for key in ["mentor", "follow_up_date", "next_review_date"]:
        if not _blank(fields.get(key)):
            value = _date_or_text(fields[key]) if key.endswith("_date") else fields[key]
            _labeled(doc, REVIEW_FIELD_LABELS[key], value)
    _heading_content(doc, "Resident comments", fields.get("resident_comments"))


def _coverage(doc: Any, records: list[dict], own: list[dict], start: str, end: str) -> None:
    doc.add_heading("Evidence received", 1)
    doc.add_paragraph("Counts refer to imported entries. Dates show the earliest and latest included records and do not establish complete coverage. Each evaluation or self assessment is counted once, regardless of its number of questions.")
    rows = []
    for category, label in CATEGORY_LABELS.items():
        items = [record for record in records if record.get("category") == category]
        dates = sorted(_record_date(record) for record in items if _record_date(record))
        rows.append((label, len(items), _period(dates[0], dates[-1]) if dates else "No dated entries included"))
    _table(doc, ["Evidence", "Entries", "Record dates"], rows, [3.1, 0.65, 3.35])
    for category in ["evaluations", "self_assessments"]:
        counts = Counter(_text(record.get("specialty")) or "Unspecified" for record in records if record.get("category") == category)
        if counts:
            _labeled(doc, CATEGORY_LABELS[category] + " by specialty", "; ".join(f"{key}: {value}" for key, value in sorted(counts.items())))
    no_dates = [record for record in own if not _record_date(record)]
    if no_dates:
        doc.add_paragraph(f"{len(no_dates)} assigned entries have no usable evidence date and are excluded from the period totals.")
    excluded_periods: set[tuple[str, str, str]] = set()
    included_ids = {id(record) for record in records}
    for record in own:
        if id(record) in included_ids or record.get("record_mode") != "snapshot":
            continue
        first, last = _iso(record.get("period_start")), _iso(record.get("period_end"))
        if first and last and first <= end and last >= start and not (start <= first <= last <= end):
            excluded_periods.add((CATEGORY_LABELS.get(record.get("category"), "Evidence"), _period(first, last), _text(record.get("source_locator"))))
    if excluded_periods:
        doc.add_paragraph("These source summaries overlap the selected dates but cover a different interval. They are excluded rather than prorated.")
        _table(doc, ["Evidence", "Source period", "Source location"], sorted(excluded_periods), [1.7, 2.5, 2.9])


def _manual_activity(doc: Any, review: dict | None) -> None:
    fields = _fields(review)
    keys = ["adult_clinic_count", "pediatric_clinic_count", "observed_encounters_completed"]
    rows = [(REVIEW_FIELD_LABELS[key], _display(fields[key])) for key in keys if not _blank(fields.get(key))]
    if rows or _text(fields.get("observed_encounters_notes")):
        doc.add_heading("Clinic and observed encounters", 1)
        if rows:
            _table(doc, ["Measure entered for this review", "Value"], rows, [5.1, 2.0])
        _paragraphs(doc, fields.get("observed_encounters_notes"))
    exams = []
    for specialty, prefix in [("Medicine", "medicine"), ("Pediatrics", "pediatrics")]:
        score, percentile = fields.get(f"{prefix}_ite_score"), fields.get(f"{prefix}_ite_percentile")
        if not _blank(score) or not _blank(percentile):
            exams.append((f"{specialty} in training examination", _display(score), _display(percentile)))
    for key in ["usmle_step1", "usmle_step2ck", "usmle_step3"]:
        if not _blank(fields.get(key)):
            exams.append((REVIEW_FIELD_LABELS[key], _display(fields[key]), "Not applicable"))
    if exams:
        doc.add_heading("Examination results", 1)
        _table(doc, ["Examination", "Entered result", "Entered percentile"], exams, [3.55, 1.75, 1.8])
    for key in ["abstracts", "publications", "advocacy", "nursing_feedback"]:
        _heading_content(doc, REVIEW_FIELD_LABELS[key], fields.get(key), level=1)


def _milestones(doc: Any, records: list[dict]) -> None:
    items = [record for record in records if record.get("category") == "milestones"]
    if not items:
        return
    doc.add_heading("Milestone scores", 1)
    doc.add_paragraph("Scores are presented by specialty and source form. Blank scores remain unrecorded; no readiness or promotion decision is inferred from a score.")
    groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for record in items:
        groups[(_text(record.get("specialty")) or "Unspecified specialty", _text(record.get("form")) or "Form not recorded")].append(record)
    for (specialty, form), group in sorted(groups.items()):
        doc.add_heading(specialty, 2)
        _labeled(doc, "Source form", form, keep_with_next=True)
        rows = []
        for record in sorted(group, key=lambda item: (_record_date(item), _text(item.get("metric")), _text(item.get("source_locator")))):
            values = record.get("values") if isinstance(record.get("values"), dict) else {}
            context = values.get("rating_type") or values.get("score_type") or values.get("stage") or ""
            metric = _text(record.get("metric"))
            if context:
                metric = f"{metric} ({_text(context)})"
            rows.append((_date(_record_date(record)), metric, _record_value(record), _source(record)))
        _table(doc, ["Date", "Milestone", "Reported value", "Source location"], rows, [1.1, 1.85, 1.15, 3.0])


def _procedures(doc: Any, records: list[dict], include_details: bool) -> None:
    items = [record for record in records if record.get("category") == "procedures"]
    if not items:
        return
    doc.add_heading("Procedures", 1)
    events = [record for record in items if record.get("record_mode", "event") != "snapshot"]
    snapshots = [record for record in items if record.get("record_mode") == "snapshot"]
    if events:
        groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
        for record in events:
            groups[(_text(record.get("specialty")) or "Combined", _text(record.get("metric")) or "Procedure not specified")].append(record)
        rows = [(specialty, metric, len(group), _date(max(_record_date(record) for record in group))) for (specialty, metric), group in sorted(groups.items())]
        _table(doc, ["Specialty", "Procedure", "Log entries", "Most recent"], rows, [1.05, 3.95, 0.8, 1.3])
    if snapshots:
        doc.add_paragraph("Source summary counts are listed separately and are not added to individual procedure entries or to other summaries.")
        rows = [(_text(record.get("specialty")), _text(record.get("metric")), _record_value(record), _record_period(record)) for record in sorted(snapshots, key=_record_date)]
        _table(doc, ["Specialty", "Procedure", "Source value", "Source period"], rows, [1.0, 2.4, 1.1, 2.6])
    if include_details and events:
        doc.add_heading("Procedure log", 2)
        rows = [(_date(_record_date(record)), _text(record.get("metric")), _source(record)) for record in sorted(events, key=lambda item: (_record_date(item), _text(item.get("metric"))))]
        _table(doc, ["Date", "Procedure", "Source location"], rows, [1.1, 2.4, 3.6])


def _conferences(doc: Any, records: list[dict]) -> None:
    items = [record for record in records if record.get("category") == "conferences"]
    if not items:
        return
    doc.add_heading("Conference attendance", 1)
    doc.add_paragraph("Each row is a separate source summary for its stated period. Percentages are not added together or averaged across different source periods.")
    rows = []
    for record in sorted(items, key=lambda item: (_text(item.get("specialty")), _text(item.get("metric")), _record_date(item))):
        values = record.get("values") if isinstance(record.get("values"), dict) else {}
        attendance = _record_value(record)
        if not _blank(values.get("attended")) or not _blank(values.get("offered")):
            attendance += f"\n{_display(values.get('attended'))} attended / {_display(values.get('offered'))} offered"
        rows.append((_text(record.get("specialty")) or "Combined", _text(record.get("metric")) or "Category not recorded", attendance, _record_period(record)))
    _table(doc, ["Specialty", "Category", "Reported attendance", "Source period"], rows, [1.0, 2.0, 1.3, 2.8])


def _duty_hours(doc: Any, records: list[dict], include_details: bool) -> None:
    hours = [record for record in records if record.get("category") == "duty_hours"]
    violations = [record for record in records if record.get("category") == "duty_violations"]
    if not hours and not violations:
        return
    doc.add_heading("Duty hours", 1)
    events = [record for record in hours if record.get("record_mode", "event") != "snapshot"]
    summaries = [record for record in hours if record.get("record_mode") == "snapshot"]
    if events:
        groups: dict[str, list[dict]] = defaultdict(list)
        for record in events:
            values = record.get("values") if isinstance(record.get("values"), dict) else {}
            work_type = _text(values.get("work_type")) or _text(record.get("metric")) or "Work type not recorded"
            groups[work_type].append(record)
        rows = []
        for work_type, group in sorted(groups.items()):
            numbers = [_number(record.get("value")) for record in group]
            available = [value for value in numbers if value is not None]
            total = _display(sum(available)) if available else "Not recorded"
            rows.append((work_type, len(group), total, sum(value is None for value in numbers)))
        _table(doc, ["Source work type", "Entries", "Sum of reported hours", "Entries missing hours"], rows, [3.6, 0.65, 1.5, 1.35])
        doc.add_paragraph("Hour totals use the numeric values recorded in the source. Gaps in logging do not establish time off, and these summaries do not make a compliance determination.")
        if include_details:
            doc.add_heading("Duty hour log", 2)
            rows = []
            for record in sorted(events, key=_record_date):
                values = record.get("values") if isinstance(record.get("values"), dict) else {}
                rows.append((_date(_record_date(record)), _text(values.get("work_type")) or _text(record.get("metric")), _display(record.get("value")), _text(record.get("rotation")), _source(record)))
            _table(doc, ["Date", "Work type", "Hours", "Rotation", "Source location"], rows, [1.0, 1.35, 0.6, 1.65, 2.5])
    if summaries:
        doc.add_paragraph("Duty hour source summaries are displayed separately and are not added to individual logged entries.")
        _table(doc, ["Source measure", "Reported value", "Source period"], [(_text(record.get("metric")), _record_value(record), _record_period(record)) for record in sorted(summaries, key=_record_date)], [2.8, 1.5, 2.8])
    doc.add_heading("Reported duty hour violations", 2)
    if not violations:
        doc.add_paragraph("No violation entries were included for the selected period.")
    else:
        doc.add_paragraph(f"{len(violations)} violation entries were reported by the imported source.")
        for index, record in enumerate(sorted(violations, key=_record_date), 1):
            values = record.get("values") if isinstance(record.get("values"), dict) else {}
            _labeled(doc, f"Entry {index}", f"{_date(_record_date(record))} | {_text(record.get('metric')) or 'Rule not recorded'}")
            _labeled(doc, "Rotation", record.get("rotation"))
            _labeled(doc, "Work type", values.get("work_type"))
            _paragraphs(doc, record.get("text"))
            for key, label in [("description", "Source description"), ("comments", "Source comments")]:
                value = values.get(key)
                if _text(value) and _text(value) != _text(record.get("text")):
                    _labeled(doc, label, value)
            _labeled(doc, "Source location", _source(record))


def _review_history(doc: Any, resident: dict, reviews: list[dict], review: dict | None, cutoff: str) -> None:
    doc.add_heading("Recorded review history", 1)
    own = review.get("history_snapshot") if review and isinstance(review.get("history_snapshot"), list) else reviews
    own = [item for item in own if isinstance(item, dict) and _text(item.get("resident_id")) == _text(resident.get("id"))]
    if review:
        selected_key = (review.get("academic_year"), review.get("pgy"), review.get("half"))
        own = [item for item in own if (item.get("academic_year"), item.get("pgy"), item.get("half")) != selected_key] + [review]
    rows = []
    selected_pgy = int(review.get("pgy") or 4) if review else 4
    selected_half = review.get("half") if review else None
    for pgy in range(1, 5):
        for half in ["Midyear", "Year end"]:
            matches = [item for item in own if item.get("pgy") == pgy and item.get("half") == half]
            if review and (pgy > selected_pgy or (pgy == selected_pgy and selected_half == "Midyear" and half == "Year end")):
                matches = []
            matches = [item for item in matches if item is review or not _iso(item.get("period_end")) or _iso(item.get("period_end")) <= cutoff]
            if not matches:
                rows.append((f"PGY{pgy} {half}", "", "", "Not included", ""))
                continue
            for item in sorted(matches, key=lambda item: (_text(item.get("academic_year")), _iso(item.get("review_date")), _text(item.get("updated_at")))):
                rows.append((f"PGY{pgy} {half}", _text(item.get("academic_year")), _date(item.get("review_date")), _text(item.get("status")) or "Draft", _text(_fields(item).get("ccc_decision"))))
    _table(doc, ["Checkpoint", "Academic year", "Review date", "Status", "Recorded decision"], rows, [1.25, 1.0, 1.05, 0.9, 2.9])


def _detailed_narratives(doc: Any, records: list[dict]) -> None:
    for category, title, noun in [("evaluations", "Evaluation details", "evaluation"), ("self_assessments", "Self assessment details", "self assessment")]:
        items = [record for record in records if record.get("category") == category]
        if not items:
            continue
        doc.add_heading(title, 1)
        for index, record in enumerate(sorted(items, key=lambda item: (_text(item.get("specialty")), _record_date(item), _text(item.get("form")), _text(item.get("evaluator")))), 1):
            specialty = _text(record.get("specialty")) or "Combined"
            doc.add_heading(f"{specialty} {noun} {index}", 2)
            _labeled(doc, "Date", _date(_record_date(record)), keep_with_next=True)
            if _iso(record.get("period_start")) or _iso(record.get("period_end")):
                _labeled(doc, "Source period", _record_period(record), keep_with_next=True)
            for key, label in [("form", "Form"), ("rotation", "Rotation"), ("evaluator", "Evaluator")]:
                _labeled(doc, label, record.get(key), keep_with_next=True)
            questions = record.get("questions") if isinstance(record.get("questions"), list) else []
            values = record.get("values") if isinstance(record.get("values"), dict) else {}
            overall = record.get("overall_comment") or values.get("overall_comments")
            # NI's text field is often a convenience concatenation of the same
            # question comments and overall comment. Avoid printing it twice,
            # while retaining independent narrative in other source formats.
            narrative = _text(record.get("text"))
            individual = [_text(question.get("comment")) for question in questions if isinstance(question, dict) and _text(question.get("comment"))]
            if _text(overall):
                individual.extend(re.split(r"\n\s*\n", _text(overall)))
            unique_individual = list(dict.fromkeys(individual))
            if not questions or narrative != "\n\n".join(unique_individual):
                _paragraphs(doc, narrative)
            for number, question in enumerate(questions, 1):
                if not isinstance(question, dict):
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.add_run(f"Question {_text(question.get('number')) or number}: ").bold = True
                p.add_run(_text(question.get("text")) or "Question text not recorded")
                _labeled(doc, "Response", question.get("response"), missing=True, keep_with_next=bool(_text(question.get("comment"))))
                _labeled(doc, "Comment", question.get("comment"))
            if _text(overall) and (questions or _text(overall) != narrative):
                _labeled(doc, "Overall comments", overall)
            if not _text(record.get("text")) and not questions and not _text(overall):
                doc.add_paragraph("No narrative or question responses were recorded in this entry.")
            _labeled(doc, "Source location", _source(record))


def _source_locations(doc: Any, records: list[dict]) -> None:
    sources = Counter((CATEGORY_LABELS.get(record.get("category"), "Evidence"), _text(record.get("source_name")) or _text(record.get("source_locator"))) for record in records if _text(record.get("source_name")) or _text(record.get("source_locator")))
    if sources:
        doc.add_heading("Source files", 1)
        _table(doc, ["Evidence", "Source file or location", "Included entries"], [(category, source, count) for (category, source), count in sorted(sources.items())], [1.8, 4.3, 1.0])


def _bytes(doc: Any) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_resident_report(
    resident: dict,
    records: list[dict],
    reviews: list[dict],
    *,
    period_start: str,
    period_end: str,
    review: dict | None = None,
    program_name: str = DEFAULT_PROGRAM_NAME,
    include_details: bool = True,
) -> bytes:
    """Return an editable resident review DOCX without writing plaintext to disk.

    A Final review's saved evidence and dates take precedence over live inputs.
    Callers must supply already active evidence for draft/current reports; the
    selected resident ID and the date window are independently checked here.
    """
    selected, own, start, end = _select_records(resident, records, review, period_start, period_end)
    doc = _new_document("Resident competency review", program_name)
    _labeled(doc, "Resident", resident.get("name"), missing=True)
    _labeled(doc, "Review", _checkpoint(review))
    _labeled(doc, "Academic year", review.get("academic_year") if review else None)
    _labeled(doc, "Reporting period", _period(start, end))
    _labeled(doc, "Coach", _fields(review).get("coach_name"))
    _labeled(doc, "Expected graduation", _date(resident["graduation_date"]) if _iso(resident.get("graduation_date")) else None)
    if review and str(review.get("status", "")).casefold() == "final":
        doc.add_paragraph("This report uses the evidence saved with the finalized review.")
    else:
        doc.add_paragraph("This report brings the imported evidence and entered committee review together for the selected dates.")
    _heading_content(doc, "Rotations during the review period", _fields(review).get("rotation_summary"))
    _review_section(doc, review)
    _coverage(doc, selected, own, start, end)
    _manual_activity(doc, review)
    _milestones(doc, selected)
    _procedures(doc, selected, include_details)
    _conferences(doc, selected)
    _duty_hours(doc, selected, include_details)
    _review_history(doc, resident, reviews, review, end)
    if include_details:
        _detailed_narratives(doc, selected)
    else:
        if any(record.get("category") in {"evaluations", "self_assessments"} for record in selected):
            doc.add_paragraph("Detailed evaluation and self assessment responses were not selected for this export.")
    _source_locations(doc, selected)
    return _bytes(doc)


def build_ccc_letter(
    resident: dict,
    review: dict,
    *,
    program_name: str = DEFAULT_PROGRAM_NAME,
) -> bytes:
    """Return a letter using only the explicitly entered committee content."""
    _validate_resident(resident, review)
    fields = _fields(review)
    doc = _new_document("Clinical competency committee letter", program_name)
    _labeled(doc, "Review", _checkpoint(review))
    _labeled(doc, "Academic year", review.get("academic_year"))
    _labeled(doc, "Review status", _text(review.get("status")) or "Draft")
    if _iso(review.get("review_date")):
        _labeled(doc, "Review date", _date(review["review_date"]))
    if _iso(review.get("period_start")) or _iso(review.get("period_end")):
        _labeled(doc, "Review period", _period(review.get("period_start"), review.get("period_end")))
    doc.add_paragraph(f"Dear {_text(resident.get('name')) or 'Resident'},")
    doc.add_paragraph(f"The following records your {_checkpoint(review)} review and the feedback entered by the Clinical Competency Committee.")
    content_keys = ["ccc_summary", "overall_impression", "strengths", "areas_for_growth", "patterns_over_time", "ccc_decision", "action_plan"]
    for key in content_keys:
        _heading_content(doc, REVIEW_FIELD_LABELS[key], fields.get(key))
    if not any(_text(fields.get(key)) for key in content_keys):
        doc.add_paragraph("Committee feedback has not yet been entered for this review.")
    for key in ["mentor", "follow_up_date", "next_review_date"]:
        if not _blank(fields.get(key)):
            value = _date_or_text(fields[key]) if key.endswith("_date") else fields[key]
            _labeled(doc, REVIEW_FIELD_LABELS[key], value)
    _heading_content(doc, "Resident comments", fields.get("resident_comments"))
    # A stored account name is audit metadata, not an authorized letter signature.
    # Signature text is included only when a program explicitly supplies it.
    if _text(fields.get("signatory_name")):
        doc.add_paragraph("Sincerely,")
        _paragraphs(doc, fields["signatory_name"])
        _paragraphs(doc, fields.get("signatory_title"))
    return _bytes(doc)
