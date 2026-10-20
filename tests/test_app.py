"""Production Streamlit flows with fictional data and a memory-only fake store.

AppTest executes the actual app.py entry point, including login and role gates.
Every network request is blocked during these tests. No real token is needed.
"""

from __future__ import annotations

import copy
import hashlib
import io
from datetime import date
from pathlib import Path

from docx import Document
from zipfile import ZipFile
import pytest
import requests
from cryptography.fernet import Fernet
from streamlit.testing.v1 import AppTest

from ccc import model as m
from ccc import annual, ccc_summary, security, ui
from ccc.storage import ConflictError, StorageError

APP = Path(__file__).resolve().parents[1] / "app.py"
PASSWORD = "Fictional app test sign-in passphrase 567!"


def by_label(elements, label):
    matches = [element for element in elements if element.label == label]
    assert len(matches) == 1, f"Expected one {label!r}; available: {[e.label for e in elements]}"
    return matches[0]


def no_exceptions(app):
    assert not app.exception, [error.message for error in app.exception]


@pytest.fixture(scope="module")
def password_hash():
    return security.hash_password(PASSWORD)


class MemoryBackend:
    def __init__(self):
        initial = m.new_state("Fictional App Test Program")
        initial = m.save_resident(initial, {"id": "resident-alex", "name": "Alex Fictional",
                                             "external_id": "NI-101", "start_date": "2026-07-01",
                                             "aliases": [], "status": "Active"}, "setup")
        self.destinations = {"private-data": initial, "alternate-data": m.new_state("Different Fictional Program")}
        self.shas = {name: hashlib.sha1(name.encode()).hexdigest() for name in self.destinations}
        self.constructed = []
        self.loads = []
        self.saves = []
        self.next_failure = None
        self.counter = 0

    def store_class(self):
        backend = self

        class FakeStore:
            def __init__(self, config, keys):
                self.destination = config["repo"]
                backend.constructed.append(self.destination)

            def load(self):
                backend.loads.append(self.destination)
                return copy.deepcopy(backend.destinations[self.destination]), backend.shas[self.destination]

            def save(self, state, expected_sha):
                if backend.next_failure is not None:
                    exc, backend.next_failure = backend.next_failure, None
                    raise exc
                if expected_sha != backend.shas[self.destination]:
                    raise ConflictError("Another synthetic session saved first.")
                backend.counter += 1
                sha = hashlib.sha1(f"synthetic-save-{backend.counter}".encode()).hexdigest()
                backend.destinations[self.destination] = copy.deepcopy(state)
                backend.shas[self.destination] = sha
                backend.saves.append((self.destination, copy.deepcopy(state)))
                return sha

            def write_blob(self, *_args, **_kwargs):
                raise AssertionError("These app checks do not persist original files.")

            def close(self):
                pass

        return FakeStore


@pytest.fixture
def app_factory(monkeypatch, password_hash):
    backend = MemoryBackend()
    monkeypatch.setattr(ui, "GitHubStore", backend.store_class())
    monkeypatch.setattr(ui, "_today", lambda _state=None: date(2026, 10, 17))
    monkeypatch.setattr(security, "DEFAULT_RATE_LIMITER", security.LoginRateLimiter())

    def reject_network(*args, **kwargs):
        raise AssertionError("A test attempted a real network request.")

    monkeypatch.setattr(requests.sessions.Session, "request", reject_network)
    ephemeral_key = Fernet.generate_key().decode("ascii")
    config = {
        "users": {role: {"display_name": f"Fictional {role}", "role": role, "password_hash": password_hash}
                  for role in ("admin", "reviewer", "viewer")},
        "github": {"owner": "example", "repo": "private-data", "branch": "main",
                   "path": "encrypted/ccc.enc", "token": "github_pat_SYNTHETIC_NOT_A_REAL_TOKEN_123456"},
        "encryption": {"key": ephemeral_key, "previous_keys": []},
        "app": {"session_timeout_minutes": 30, "max_session_hours": 8},
    }

    def create(*, configured=True):
        app = AppTest.from_file(APP, default_timeout=15)
        if configured:
            app.secrets.update(copy.deepcopy(config))
        app.run()
        no_exceptions(app)
        return app

    return create, backend


def login(app, role="admin", password=PASSWORD):
    by_label(app.text_input, "Username").set_value(role)
    by_label(app.text_input, "Password").set_value(password)
    by_label(app.button, "Sign in").click().run()
    no_exceptions(app)
    return app


def navigate(app, page):
    by_label(app.radio, "Go to").set_value(page).run()
    no_exceptions(app)
    return app


def enter_new_resident(app):
    navigate(app, "Administration")
    by_label(app.text_input, "Resident name").set_value("Casey Fictional")


def test_missing_configuration_and_unsigned_session_never_open_storage(app_factory):
    create, backend = app_factory
    missing = create(configured=False)
    assert any("Setup" in item.value for item in missing.info)
    configured = create()
    assert by_label(configured.button, "Sign in")
    assert backend.constructed == [] and backend.loads == []
    assert "state" not in configured.session_state


def test_wrong_password_never_constructs_or_loads_the_store(app_factory):
    create, backend = app_factory
    app = login(create(), password="Wrong synthetic passphrase")
    assert backend.constructed == [] and backend.loads == []
    assert "identity" not in app.session_state
    assert any("could not be verified" in error.value for error in app.error)
    assert by_label(app.text_input, "Password").value == ""
    assert app.session_state.get("login_password", "") == ""


def test_successful_login_loads_saved_state_and_refresh_restores_server_version(app_factory):
    create, backend = app_factory
    app = login(create())
    assert backend.loads == ["private-data"]
    assert app.session_state["identity"]["username"] == "admin"
    assert app.session_state["state"]["residents"][0]["name"] == "Alex Fictional"
    backend.destinations["private-data"]["settings"]["program_name"] = "Updated by another fictional session"
    by_label(app.button, "Refresh saved data").click().run()
    no_exceptions(app)
    assert len(backend.loads) == 2
    assert app.session_state["state"]["settings"]["program_name"] == "Updated by another fictional session"


@pytest.mark.parametrize("role,can_upload,can_admin,can_review", [
    ("viewer", False, False, False), ("reviewer", False, False, True), ("admin", True, True, True)])
def test_roles_control_navigation_and_review_actions(app_factory, role, can_upload, can_admin, can_review):
    create, _backend = app_factory
    app = login(create(), role)
    navigation = by_label(app.radio, "Go to")
    assert ("Upload files" in navigation.options) is can_upload
    assert ("Administration" in navigation.options) is can_admin
    navigate(app, "Resident review")
    assert any(button.label == "Save draft" for button in app.button) is can_review
    assert any(button.label == "Finalize review" for button in app.button) is can_review
    if not can_admin:
        # Request a destination through the app's own navigation mechanism;
        # AppTest correctly rejects injecting an invalid radio option itself.
        app.session_state["nav_target"] = "Administration"
        app.run()
        no_exceptions(app)
        assert by_label(app.radio, "Go to").value == "Resident review"
        assert not any(button.label == "Save resident" for button in app.button)


def test_admin_roster_save_is_persisted_and_available_to_a_new_session(app_factory):
    create, backend = app_factory
    app = login(create())
    enter_new_resident(app)
    by_label(app.button, "Save resident").click().run()
    no_exceptions(app)
    assert len(backend.saves) == 1
    assert {r["name"] for r in app.session_state["state"]["residents"]} == {"Alex Fictional", "Casey Fictional"}
    added = next(r for r in app.session_state["state"]["residents"] if r["name"] == "Casey Fictional")
    assert added["external_id"] == ""
    assert added["id"]
    assert added["id"] != "Casey Fictional"
    reopened = login(create(), "viewer")
    assert {r["name"] for r in reopened.session_state["state"]["residents"]} == {"Alex Fictional", "Casey Fictional"}
    assert next(r for r in reopened.session_state["state"]["residents"] if r["name"] == "Casey Fictional")["id"] == added["id"]


def test_remove_resident_requires_confirmation_and_persists_to_new_sessions(app_factory):
    create, backend = app_factory
    app = login(create())
    enter_new_resident(app)
    by_label(app.button, "Save resident").click().run()
    removed_id = next(r["id"] for r in app.session_state["state"]["residents"] if r["name"] == "Casey Fictional")
    by_label(app.selectbox, "Add or edit a resident").set_value(removed_id).run()
    no_exceptions(app)
    assert by_label(app.button, "Remove resident").disabled
    assert any("no saved evidence or CCC reviews" in item.value for item in app.markdown)
    app.session_state["selected_resident"] = removed_id
    app.session_state["upload_preview"] = [{"data": b"fictional pending upload"}]
    app.session_state["upload_preview_context"] = ("fictional",)
    app.session_state["generated_export"] = {"data": b"fictional prepared report"}
    by_label(app.checkbox, "I understand and want to remove this resident").check().run()
    by_label(app.button, "Remove resident").click().run()
    no_exceptions(app)
    assert len(backend.saves) == 2
    assert [r["name"] for r in app.session_state["state"]["residents"]] == ["Alex Fictional"]
    assert by_label(app.selectbox, "Add or edit a resident").value == "__new__"
    for key in ("selected_resident", "upload_preview", "upload_preview_context", "generated_export"):
        assert key not in app.session_state
    assert app.session_state["state"]["audit"][-1]["action"] == "Resident removed"
    reopened = login(create(), "viewer")
    assert [r["name"] for r in reopened.session_state["state"]["residents"]] == ["Alex Fictional"]


@pytest.mark.parametrize("failure_type", [StorageError, ConflictError])
def test_failed_resident_removal_keeps_saved_roster_and_uses_existing_recovery(app_factory, failure_type):
    create, backend = app_factory
    app = login(create())
    navigate(app, "Administration")
    by_label(app.selectbox, "Add or edit a resident").set_value("resident-alex").run()
    before = copy.deepcopy(app.session_state["state"])
    app.session_state["selected_resident"] = "resident-alex"
    backend.next_failure = failure_type("Synthetic removal save failure.")
    by_label(app.checkbox, "I understand and want to remove this resident").check().run()
    by_label(app.button, "Remove resident").click().run()
    no_exceptions(app)
    assert app.session_state["state"] == before
    assert backend.destinations["private-data"] == before
    assert app.session_state["selected_resident"] == "resident-alex"
    assert app.session_state["pending_save"]["state"]["residents"] == []
    assert not backend.saves
    app.run()
    if failure_type is StorageError:
        by_label(app.button, "Retry this save").click().run()
        no_exceptions(app)
        assert app.session_state["state"]["residents"] == []
        assert "selected_resident" not in app.session_state
        assert "pending_save" not in app.session_state
        reopened = login(create())
        assert reopened.session_state["state"]["residents"] == []
        navigate(reopened, "Resident review")
        navigate(reopened, "Word reports")
        no_exceptions(reopened)
    else:
        assert not any(button.label == "Retry this save" for button in app.button)
        by_label(app.checkbox, "I have kept a copy of any changes I need").check().run()
        by_label(app.button, "Reload saved records").click().run()
        no_exceptions(app)
        assert app.session_state["state"] == before
        assert "pending_save" not in app.session_state


class FictionalUpload:
    def __init__(self, filename, data):
        self.name = filename
        self.data = data
        self.size = len(data)

    def getvalue(self):
        return self.data


@pytest.fixture
def supplied_uploads(monkeypatch):
    """Supply bytes at the uploader boundary, while rendering the real widgets.

    AppTest has no browser file-picker control. Parsing, matching, confirmation,
    batch preparation, and dashboard persistence still use production code.
    """
    selected = {}
    actual_uploader = ui.st.file_uploader

    def uploader(*args, **kwargs):
        value = actual_uploader(*args, **kwargs)
        key = kwargs.get("key", "")
        return selected.get(key, []) if key.startswith("files_") else value

    monkeypatch.setattr(ui.st, "file_uploader", uploader)
    return selected


def fictional_example(filename, *, extra=b""):
    data = (APP.parent / "examples" / filename).read_bytes().replace(b"Example", b"Fictional") + extra
    return FictionalUpload(filename, data)


def confirm_and_save_upload(app):
    by_label(app.checkbox, "Keep an encrypted copy of each original file").uncheck()
    by_label(app.checkbox, "The included files, resident matches, dates, and any warnings are correct").check().run()
    assert not by_label(app.button, "Save to resident dashboards").disabled
    by_label(app.button, "Save to resident dashboards").click().run()
    no_exceptions(app)


@pytest.mark.parametrize("group,filename,specialty,category", [
    ("eval", "evaluations_medicine.csv", "Medicine", "evaluations"),
    ("milestones", "milestones.csv", "Combined", "milestones"),
    ("activity", "procedures.csv", "Medicine", "procedures"),
    ("duty", "duty_hours.csv", "Combined", "duty_hours"),
])
def test_each_upload_area_accepts_one_file_with_every_other_area_empty(app_factory, supplied_uploads, group, filename, specialty, category):
    create, backend = app_factory
    app = login(create())
    navigate(app, "Upload files")
    assert by_label(app.button, "Process files").disabled
    supplied_uploads[f"files_{group}_0"] = [fictional_example(filename)]
    by_label(app.radio, "These files are for").set_value(specialty).run()
    assert not by_label(app.button, "Process 1 file").disabled
    by_label(app.button, "Process 1 file").click().run()
    no_exceptions(app)
    assert len(app.session_state["upload_preview"]) == 1
    assert not app.error
    assert by_label(app.button, "Save to resident dashboards").disabled
    assert not backend.saves
    confirm_and_save_upload(app)
    state = app.session_state["state"]
    assert len(backend.saves) == 1 and len(state["sources"]) == 1
    assert state["records"] and {r["category"] for r in state["records"]} == {category}
    assert {r["resident_id"] for r in state["records"]} == {"resident-alex"}
    assert "upload_preview" not in app.session_state
    assert by_label(app.button, "Process files").disabled
    reopened = login(create(), "viewer")
    assert reopened.session_state["state"]["records"] == state["records"]


def test_later_upload_adds_evidence_while_duplicate_only_files_do_not_block_it(app_factory, supplied_uploads):
    create, backend = app_factory
    supplied_uploads["files_eval_0"] = [fictional_example("evaluations_medicine.csv")]
    app = login(create())
    navigate(app, "Upload files")
    by_label(app.button, "Process 1 file").click().run()
    confirm_and_save_upload(app)
    earlier = copy.deepcopy(app.session_state["state"]["records"])

    supplied_uploads["files_eval_1"] = [fictional_example("evaluations_medicine.csv", extra=b"\n")]
    supplied_uploads["files_activity_1"] = [fictional_example("procedures.csv")]
    app.run()
    by_label(app.button, "Process 2 files").click().run()
    no_exceptions(app)
    confirm_and_save_upload(app)
    state = app.session_state["state"]
    assert state["records"][:len(earlier)] == earlier
    assert len(state["sources"]) == 2
    assert sum(r["category"] == "procedures" for r in state["records"]) == 2
    assert any("Skipped 1 file with no new entries" in item.value for item in app.success)

    # Another changed-byte export of the same evidence is a no-op, not an error.
    before = copy.deepcopy(state)
    supplied_uploads["files_eval_2"] = [fictional_example("evaluations_medicine.csv", extra=b"\n\n")]
    app.run()
    by_label(app.button, "Process 1 file").click().run()
    no_exceptions(app)
    assert not any(button.label == "Save to resident dashboards" for button in app.button)
    assert app.session_state["state"] == before
    assert len(backend.saves) == 2
    assert not app.error
    assert any("No new entries to save" in item.value for item in app.info)


def test_identical_file_can_complete_an_earlier_import_after_a_resident_was_skipped(app_factory, supplied_uploads):
    create, backend = app_factory
    data = ("resident_name,date,metric,value,text,form,evaluator,rotation\n"
            "Alex Fictional,2026-09-01,Overall,3,Alex observation,Faculty,Teacher,Ward\n"
            "Brooke Fictional,2026-09-01,Overall,4,Brooke observation,Faculty,Teacher,Ward\n").encode()
    file = FictionalUpload("two-residents.csv", data)
    supplied_uploads["files_eval_0"] = [file]
    app = login(create())
    navigate(app, "Upload files")
    by_label(app.button, "Process 1 file").click().run()
    by_label(app.selectbox, "Brooke Fictional").set_value("__skip__").run()
    confirm_and_save_upload(app)
    earlier = copy.deepcopy(app.session_state["state"]["records"])
    assert len(earlier) == 1 and earlier[0]["resident_id"] == "resident-alex"

    navigate(app, "Administration")
    by_label(app.text_input, "Resident name").set_value("Brooke Fictional")
    by_label(app.button, "Save resident").click().run()
    brooke_id = next(r["id"] for r in app.session_state["state"]["residents"] if r["name"] == "Brooke Fictional")
    supplied_uploads["files_eval_1"] = [file]
    navigate(app, "Upload files")
    by_label(app.button, "Process 1 file").click().run()
    no_exceptions(app)
    assert any("This file was uploaded before" in item.value for item in app.info)
    assert by_label(app.selectbox, "Brooke Fictional").value == brooke_id
    confirm_and_save_upload(app)
    state = app.session_state["state"]
    assert state["records"][:1] == earlier
    assert len(state["records"]) == len(state["sources"]) == 2
    assert state["records"][-1]["resident_id"] == brooke_id
    assert state["sources"][0]["sha256"] == state["sources"][1]["sha256"]

    before = copy.deepcopy(state)
    supplied_uploads["files_eval_2"] = [file]
    app.run()
    by_label(app.button, "Process 1 file").click().run()
    no_exceptions(app)
    assert any("No new entries to save" in item.value for item in app.info)
    assert not any(button.label == "Save to resident dashboards" for button in app.button)
    assert app.session_state["state"] == before
    assert len(backend.saves) == 3


def test_changed_entries_across_new_files_require_an_available_confirmation(app_factory, supplied_uploads):
    create, backend = app_factory
    header = "resident_name,date,metric,value,text,form,evaluator,rotation\n"
    original = (header + "Alex Fictional,2026-09-01,Overall,3,Original observation,Faculty,Teacher,Ward\n").encode()
    changed = (header + "Alex Fictional,2026-09-01,Overall,4,Changed observation,Faculty,Teacher,Ward\n").encode()
    supplied_uploads["files_eval_0"] = [FictionalUpload("first.csv", original), FictionalUpload("changed.csv", changed)]
    app = login(create())
    navigate(app, "Upload files")
    by_label(app.button, "Process 2 files").click().run()
    no_exceptions(app)
    overlap = "I checked overlapping entries and want to keep any remaining revisions as additional evidence"
    assert not by_label(app.checkbox, overlap).disabled
    assert any("1 entry overlaps" in item.value for item in app.warning)
    by_label(app.checkbox, "The included files, resident matches, dates, and any warnings are correct").check().run()
    assert by_label(app.button, "Save to resident dashboards").disabled
    assert not backend.saves
    by_label(app.checkbox, overlap).check().run()
    confirm_and_save_upload(app)
    assert len(app.session_state["state"]["records"]) == 2
    assert app.session_state["state"]["records"][1]["possible_revision"] is True
    assert len(backend.saves) == 1


def test_reviewer_saves_draft_then_finalizes_its_evidence_snapshot(app_factory):
    create, backend = app_factory
    app = login(create(), "reviewer")
    navigate(app, "Resident review")
    by_label(app.text_area, "CCC assessment or decision").set_value("A synthetic committee decision for testing.")
    by_label(app.button, "Save draft").click().run()
    no_exceptions(app)
    saved = app.session_state["state"]["reviews"][0]
    assert saved["status"] == "Draft"
    by_label(app.selectbox, "Start or open a review").set_value(saved["id"]).run()
    no_exceptions(app)
    by_label(app.checkbox, "I reviewed the evidence and this is the CCC's recorded assessment").check()
    by_label(app.button, "Finalize review").click().run()
    no_exceptions(app)
    final = app.session_state["state"]["reviews"][0]
    assert final["status"] == "Final" and isinstance(final["snapshot"], list)
    assert final["finalized_by"] == "reviewer"
    assert len(backend.saves) == 2


def test_default_ccc_summary_opens_without_a_passphrase_and_uses_start_through_today(app_factory, monkeypatch):
    create, backend = app_factory
    state = backend.destinations["private-data"]
    state["residents"][0]["start_date"] = "2023-07-01"
    app = login(create(), "viewer")
    navigate(app, "Word reports")
    assert by_label(app.radio, "Document").value == "CCC summary"
    assert by_label(app.date_input, "Reporting period starts").value == date(2023, 7, 1)
    assert by_label(app.date_input, "Evidence through").value == date(2026, 10, 17)
    assert not any("passphrase" in item.label.lower() for item in app.text_input)
    by_label(app.button, "Prepare Word download").click().run()
    no_exceptions(app)
    generated = app.session_state["generated_export"]
    assert generated["protected"] is False
    assert generated["filename"] == "2026-10-17_Alex_Fictional_CCC_Summary.docx"
    doc = Document(io.BytesIO(generated["data"]))
    assert len(doc.tables) == 27
    assert doc.tables[-2].cell(0, 0).text == "MEDICINE EVALUATION SUMMARY"
    assert doc.tables[-1].cell(0, 0).text == "PEDIATRIC EVALUATION SUMMARY"
    assert "2023-07-01 to 2026-10-17" in "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert not backend.saves
    assert app.session_state.get("export_password", "") == ""
    assert app.session_state.get("export_confirmation", "") == ""

    # A failed second export must remove the previous resident download.
    def fail_report(*_args, **_kwargs):
        raise m.ValidationError("Synthetic report preparation failure.")

    monkeypatch.setattr(ccc_summary, "build_ccc_summary", fail_report)
    by_label(app.button, "Prepare Word download").click().run()
    no_exceptions(app)
    assert "generated_export" not in app.session_state
    assert not app.download_button
    assert any("Synthetic report preparation failure" in error.value for error in app.error)


def test_yearly_progress_saves_optional_values_and_keeps_other_years_separate(app_factory):
    create, backend = app_factory
    app = login(create(), "reviewer")
    navigate(app, "Yearly progress")
    assert by_label(app.selectbox, "Residency year").value == 1
    assert all(item.value is None for item in app.number_input)
    by_label(app.number_input, "Pediatrics ITE score").set_value(71)
    by_label(app.number_input, "Pediatrics expected pass rate (%)").set_value(93)
    by_label(app.number_input, "Medicine ITE score").set_value(75)
    by_label(app.number_input, "Medicine ITE percentile").set_value(68)
    by_label(app.number_input, "Adult clinic patient tally").set_value(0)
    by_label(app.selectbox, "Observed encounter 1").set_value(True)
    by_label(app.selectbox, "Observed encounter 2").set_value(False)
    by_label(app.button, "Save yearly progress").click().run()
    no_exceptions(app)
    entry = app.session_state["state"]["annual_progress"][0]
    assert entry["pgy"] == 1 and entry["as_of"] == "2026-10-17"
    assert entry["fields"] == {"pediatrics_ite_score": 71, "pediatrics_ite_expected_pass_rate": 93,
                               "medicine_ite_score": 75, "medicine_ite_percentile": 68,
                               "adult_clinic_count": 0, "pediatric_clinic_count": None,
                               "observed_encounter_1": True, "observed_encounter_2": False,
                               "observed_encounter_3": None, "observed_encounter_4": None}
    by_label(app.number_input, "Pediatric clinic patient tally").set_value(25)
    by_label(app.button, "Save yearly progress").click().run()
    no_exceptions(app)
    updated = app.session_state["state"]["annual_progress"][0]
    assert updated["fields"]["pediatric_clinic_count"] == 25
    assert updated["fields"]["medicine_ite_score"] == 75
    assert updated["versions"][0]["fields"]["pediatric_clinic_count"] is None
    by_label(app.selectbox, "Residency year").set_value(2).run()
    no_exceptions(app)
    assert all(item.value is None for item in app.number_input)
    by_label(app.number_input, "Adult clinic patient tally").set_value(44)
    by_label(app.button, "Save yearly progress").click().run()
    no_exceptions(app)
    entries = app.session_state["state"]["annual_progress"]
    assert len(entries) == 2 and entries[0]["fields"]["adult_clinic_count"] == 0
    assert entries[1]["pgy"] == 2 and entries[1]["fields"]["adult_clinic_count"] == 44
    assert len(backend.saves) == 3
    # Future-year values are saved, but do not appear in a report before that year.
    navigate(app, "Word reports")
    by_label(app.button, "Prepare Word download").click().run()
    no_exceptions(app)
    doc = Document(io.BytesIO(app.session_state["generated_export"]["data"]))
    ite = next(table for table in doc.tables if table.cell(0, 0).text == "IN-TRAINING EXAM SCORES")
    assert "75 (68%tile)" in "\n".join(cell.text for row in ite.rows for cell in row.cells)
    clinic = next(table for table in doc.tables if table.cell(0, 0).text == "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    assert "44" not in "\n".join(cell.text for row in clinic.rows for cell in row.cells)


def test_viewer_sees_yearly_progress_but_cannot_change_it(app_factory):
    create, backend = app_factory
    state = backend.destinations["private-data"]
    entry = annual.default_progress(state["residents"][0], 1, "2026-10-01")
    entry["fields"]["medicine_ite_score"] = 70
    backend.destinations["private-data"] = annual.save_progress(state, "resident-alex", entry, "setup")
    app = login(create(), "viewer")
    navigate(app, "Yearly progress")
    assert app.dataframe[0].value.iloc[0]["Medicine score"] == 70
    assert not app.number_input
    assert not any(button.label == "Save yearly progress" for button in app.button)
    assert not backend.saves


def test_resident_to_yearly_progress_link_keeps_the_selected_resident(app_factory):
    create, backend = app_factory
    backend.destinations["private-data"] = m.save_resident(backend.destinations["private-data"],
        {"id": "resident-casey", "name": "Casey Fictional", "start_date": "2024-07-01"}, "setup")
    app = login(create())
    navigate(app, "Resident review")
    by_label(app.selectbox, "Find a resident").set_value("resident-casey").run()
    by_label(app.button, "Open yearly exam scores, clinic tallies and encounters").click().run()
    no_exceptions(app)
    assert by_label(app.radio, "Go to").value == "Yearly progress"
    assert by_label(app.selectbox, "Find a resident").value == "resident-casey"


def test_bulk_ccc_summaries_default_to_each_residents_own_start_date(app_factory):
    create, backend = app_factory
    backend.destinations["private-data"] = m.save_resident(backend.destinations["private-data"],
        {"id": "resident-casey", "name": "Casey Fictional", "start_date": "2024-07-01"}, "setup")
    app = login(create(), "viewer")
    navigate(app, "Word reports")
    by_label(app.radio, "Who is this for?").set_value("Several residents").run()
    no_exceptions(app)
    assert by_label(app.radio, "Report start dates").value == "Each resident's start date"
    assert by_label(app.radio, "Group report basis").value == "Current dashboard date range"
    by_label(app.button, "Prepare Word download").click().run()
    no_exceptions(app)
    with ZipFile(io.BytesIO(app.session_state["generated_export"]["data"])) as archive:
        assert len(archive.namelist()) == 2
        for name, start in [("Alex_Fictional", "2026-07-01"), ("Casey_Fictional", "2024-07-01")]:
            path = next(path for path in archive.namelist() if name in path)
            doc = Document(io.BytesIO(archive.read(path)))
            assert start + " to 2026-10-17" in "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert not backend.saves


def test_failed_save_preserves_saved_state_and_conflict_recovery_is_available(app_factory):
    create, backend = app_factory
    app = login(create())
    saved_before = copy.deepcopy(app.session_state["state"])
    enter_new_resident(app)
    backend.next_failure = ConflictError("Another synthetic session saved first.")
    by_label(app.button, "Save resident").click().run()
    no_exceptions(app)
    assert app.session_state["state"] == saved_before
    assert backend.destinations["private-data"] == saved_before
    assert app.session_state["pending_save"]["state"]["residents"][-1]["name"] == "Casey Fictional"
    app.run()
    no_exceptions(app)
    assert not any(button.label == "Retry this save" for button in app.button)
    assert by_label(app.button, "Reload saved records").disabled
    assert by_label(app.download_button, "Download encrypted pending changes")
    by_label(app.checkbox, "I have kept a copy of any changes I need").check().run()
    no_exceptions(app)
    by_label(app.button, "Reload saved records").click().run()
    no_exceptions(app)
    assert "pending_save" not in app.session_state
    assert app.session_state["state"] == saved_before


def test_recoverable_storage_error_can_retry_without_losing_pending_changes(app_factory):
    create, backend = app_factory
    app = login(create())
    enter_new_resident(app)
    backend.next_failure = StorageError("Synthetic temporary storage outage.")
    by_label(app.button, "Save resident").click().run()
    no_exceptions(app)
    app.run()
    no_exceptions(app)
    by_label(app.button, "Retry this save").click().run()
    no_exceptions(app)
    assert "pending_save" not in app.session_state
    assert any(r["name"] == "Casey Fictional" for r in app.session_state["state"]["residents"])
    assert len(backend.saves) == 1


def test_sign_out_clears_resident_data_and_returns_to_the_login_gate(app_factory):
    create, backend = app_factory
    app = login(create())
    app.session_state["upload_preview"] = [{"data": b"synthetic private source"}]
    app.session_state["generated_export"] = {"data": b"synthetic private report"}
    by_label(app.button, "Sign out").click().run()
    no_exceptions(app)
    for key in ("state", "sha", "identity", "upload_preview", "generated_export", "encryption_keys"):
        assert key not in app.session_state
    assert by_label(app.button, "Sign in")
    assert backend.loads == ["private-data"]


def test_storage_destination_change_invalidates_the_old_session_database(app_factory):
    create, backend = app_factory
    app = login(create())
    app.session_state["upload_preview"] = [{"data": b"synthetic private source"}]
    app.secrets["github"]["repo"] = "alternate-data"
    app.run()
    no_exceptions(app)
    assert app.session_state["state"]["settings"]["program_name"] == "Different Fictional Program"
    assert app.session_state["state"]["residents"] == []
    assert "upload_preview" not in app.session_state
    assert backend.loads == ["private-data", "alternate-data"]
    assert backend.saves == []
