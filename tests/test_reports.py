"""Regression checks for evidence fidelity and review-export isolation."""

from copy import deepcopy
from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from ccc.reports import build_ccc_letter, build_resident_report


RESIDENT = {"id": "resident-example", "name": "Alex Example", "start_date": "2026-07-01"}
REVIEW = {
    "id": "review-example", "resident_id": RESIDENT["id"], "pgy": 1,
    "half": "Midyear", "academic_year": "2026-2027", "status": "Draft",
    "period_start": "2026-07-01", "period_end": "2026-12-31",
    "review_date": "2026-12-15", "fields": {},
}


def record(**changes):
    result = {
        "id": "entry-example", "resident_id": RESIDENT["id"],
        "resident_name": RESIDENT["name"], "category": "evaluations",
        "specialty": "Medicine", "date": "2026-10-01",
        "record_mode": "event", "form": "Example evaluation",
        "metric": "Evaluation", "value": 1, "unit": "assessments",
        "text": "Entered evaluation narrative", "questions": [], "values": {},
        "source_name": "example.csv", "source_locator": "Sheet 1, row 2",
    }
    result.update(changes)
    return result


def document_text(data):
    doc = Document(BytesIO(data))
    return "\n".join(node.text or "" for node in doc.element.body.iter(qn("w:t")))


def report(records, review=None, details=True, reviews=None):
    return build_resident_report(
        deepcopy(RESIDENT), deepcopy(records), deepcopy(reviews or []),
        period_start="2026-07-01", period_end="2026-12-31",
        review=deepcopy(review), include_details=details,
    )


class ReportEvidenceTests(unittest.TestCase):
    def test_date_boundaries_and_resident_isolation(self):
        entries = [
            record(id="first", date="2026-07-01", text="First boundary evidence"),
            record(id="last", date="2026-12-31", text="Last boundary evidence"),
            record(id="old", date="2026-06-30", text="Outside old evidence"),
            record(id="future", date="2027-01-01", text="Outside future evidence"),
            record(id="other", resident_id="other-person", text="Other resident evidence"),
            record(id="unassigned", resident_id="", text="Unassigned evidence"),
        ]
        text = document_text(report(entries))
        self.assertIn("First boundary evidence", text)
        self.assertIn("Last boundary evidence", text)
        for excluded in ["Outside old evidence", "Outside future evidence", "Other resident evidence", "Unassigned evidence"]:
            self.assertNotIn(excluded, text)

    def test_final_review_uses_saved_evidence_and_saved_dates(self):
        review = deepcopy(REVIEW)
        review.update(status="Final", snapshot=[record(text="Saved evidence at finalization")], version=2)
        data = build_resident_report(
            RESIDENT, [record(text="Later live evidence")], [],
            period_start="2030-01-01", period_end="2030-02-01", review=review,
        )
        text = document_text(data)
        self.assertIn("Saved evidence at finalization", text)
        self.assertNotIn("Later live evidence", text)
        self.assertIn("Jul 01, 2026 to Dec 31, 2026", text)
        self.assertNotIn("2030", text)
        self.assertIn("Review version:", text)

    def test_empty_final_snapshot_is_not_replaced_with_live_data(self):
        review = dict(REVIEW, status="Final", snapshot=[])
        self.assertNotIn("Live evidence", document_text(report([record(text="Live evidence")], review)))

    def test_missing_final_snapshot_fails_closed(self):
        with self.assertRaisesRegex(ValueError, "saved evidence"):
            report([record()], dict(REVIEW, status="Final"))

    def test_other_resident_review_fails_closed(self):
        with self.assertRaisesRegex(ValueError, "does not belong"):
            report([], dict(REVIEW, resident_id="different-person"))

    def test_eval_count_is_not_question_count_and_has_no_fifty_entry_cap(self):
        entries = [record(
            id=f"evaluation-{i}", text=f"Evaluation narrative number {i}",
            questions=[
                {"number": "1", "text": "Question one", "response": 3, "comment": f"Comment A {i}"},
                {"number": "2", "text": "Question two", "response": None, "comment": f"Comment B {i}"},
            ],
        ) for i in range(61)]
        data = report(entries)
        text = document_text(data)
        self.assertIn("Evaluation narrative number 60", text)
        self.assertIn("Comment B 60", text)
        doc = Document(BytesIO(data))
        coverage = next(table for table in doc.tables if table.cell(0, 0).text == "Evidence" and table.cell(0, 1).text == "Entries")
        row = next(row for row in coverage.rows if row.cells[0].text == "Evaluations")
        self.assertEqual("61", row.cells[1].text)

    def test_distinct_overall_comment_and_question_response_are_preserved_once(self):
        entry = record(
            text="Specific question comment\n\nDistinct overall comment",
            overall_comment="Distinct overall comment",
            questions=[{"number": "4", "text": "Original question wording", "response": "Always", "comment": "Specific question comment"}],
        )
        text = document_text(report([entry]))
        for value in ["Original question wording", "Always", "Specific question comment", "Distinct overall comment"]:
            self.assertEqual(1, text.count(value), value)

    def test_null_milestone_is_not_zero_and_specialties_are_separate(self):
        entries = [
            record(category="milestones", specialty="Medicine", metric="PC1", value=None, unit="milestone level", text=""),
            record(category="milestones", specialty="Pediatrics", metric="PC1", value=0, unit="milestone level", text=""),
        ]
        doc = Document(BytesIO(report(entries)))
        tables = [table for table in doc.tables if table.cell(0, 0).text == "Date" and table.cell(0, 1).text == "Milestone"]
        self.assertEqual(2, len(tables))
        self.assertEqual("Not recorded", tables[0].cell(1, 2).text)
        self.assertEqual("0 milestone level", tables[1].cell(1, 2).text)
        text = document_text(report(entries))
        self.assertNotIn("Ready for Independent Practice", text)

    def test_source_observation_status_is_preserved_for_missing_score(self):
        entry = record(category="milestones", metric="PC2", value=None, values={"assessment_status": "Not observed"})
        self.assertIn("Not observed", document_text(report([entry])))

    def test_partial_attendance_snapshot_is_excluded_without_prorating(self):
        entry = record(
            category="conferences", metric="Morning report", value=66, unit="%",
            record_mode="snapshot", date="2026-10-01",
            period_start="2026-05-01", period_end="2026-10-01",
        )
        doc = Document(BytesIO(report([entry])))
        coverage = next(table for table in doc.tables if table.cell(0, 0).text == "Evidence" and table.cell(0, 1).text == "Entries")
        row = next(row for row in coverage.rows if row.cells[0].text == "Attendance snapshots")
        self.assertEqual("0", row.cells[1].text)
        self.assertIn("excluded rather than prorated", document_text(report([entry])))

    def test_duty_snapshot_is_not_added_to_events_and_missing_hours_remain_missing(self):
        entries = [
            record(category="duty_hours", metric="Clinical", value=8, values={"work_type": "Clinical"}),
            record(category="duty_hours", metric="Clinical", value=None, values={"work_type": "Clinical"}),
            record(category="duty_hours", metric="Monthly hours", value=100, record_mode="snapshot", period_start="2026-07-01", period_end="2026-10-01"),
        ]
        doc = Document(BytesIO(report(entries)))
        table = next(table for table in doc.tables if table.cell(0, 0).text == "Source work type")
        self.assertEqual("8", table.cell(1, 2).text)
        self.assertEqual("1", table.cell(1, 3).text)

    def test_all_violation_narratives_remain_beyond_legacy_five_entry_limit(self):
        entries = [record(category="duty_violations", metric="Source rule", text=f"Violation detail {i}") for i in range(8)]
        text = document_text(report(entries, details=False))
        self.assertIn("8 violation entries", text)
        self.assertIn("Violation detail 7", text)

    def test_invalid_report_window_rejected(self):
        with self.assertRaisesRegex(ValueError, "valid report"):
            build_resident_report(RESIDENT, [], [], period_start="2027-01-01", period_end="2026-12-31")

    def test_html_and_invalid_xml_chars_cannot_embed_remote_content(self):
        entry = record(text='<p>Safe narrative\x00</p><script>Unwanted script</script><img src="https://example.invalid/private.png"><p>Second paragraph</p>')
        data = report([entry])
        text = document_text(data)
        self.assertIn("Safe narrative", text)
        self.assertIn("Second paragraph", text)
        self.assertNotIn("Unwanted script", text)
        with ZipFile(BytesIO(data)) as archive:
            relationships = "\n".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".rels"))
        self.assertNotIn('TargetMode="External"', relationships)
        self.assertNotIn("example.invalid", relationships)

    def test_review_history_keeps_eight_checkpoints_and_repeated_training_year(self):
        earlier = dict(REVIEW, id="earlier", academic_year="2025-2026", period_start="2025-07-01", period_end="2025-12-31", review_date="2025-12-15")
        doc = Document(BytesIO(report([], REVIEW, reviews=[earlier, REVIEW])))
        history = next(table for table in doc.tables if table.cell(0, 0).text == "Checkpoint")
        # Header + eight checkpoint rows + the additional recorded academic year.
        self.assertEqual(10, len(history.rows))
        self.assertIn("2025-2026", " ".join(cell.text for row in history.rows for cell in row.cells))


class LetterTests(unittest.TestCase):
    def test_letter_never_invents_concerns_or_progression_decision(self):
        for half in ["Midyear", "Year end"]:
            text = document_text(build_ccc_letter(RESIDENT, dict(REVIEW, half=half)))
            self.assertNotIn("No Concerns", text)
            self.assertNotIn("No concerns", text)
            self.assertNotIn("Ready for Independent Practice", text)
            self.assertNotIn("promotion", text.lower())
            self.assertIn("not yet been entered", text)

    def test_letter_contains_only_entered_feedback_and_no_assumed_signatory(self):
        review = deepcopy(REVIEW)
        review.update(finalized_by="Account operator", fields={
            "ccc_decision": "Continue current training plan.",
            "strengths": "Clear family communication.",
            "areas_for_growth": "Improve handoff prioritization.",
            "action_plan": "Review two observed handoffs with a coach.",
            "mentor": "Designated faculty coach",
            "follow_up_date": "2027-01-15",
            "next_review_date": "2027-06-10",
        })
        text = document_text(build_ccc_letter(RESIDENT, review))
        for value in ["Continue current training plan.", "Clear family communication.", "Improve handoff prioritization.", "Review two observed handoffs with a coach.", "Jan 15, 2027", "Jun 10, 2027"]:
            self.assertIn(value, text)
        self.assertNotIn("Account operator", text)
        self.assertNotIn("Sincerely", text)

    def test_other_resident_letter_rejected(self):
        with self.assertRaisesRegex(ValueError, "does not belong"):
            build_ccc_letter(RESIDENT, dict(REVIEW, resident_id="other-person"))

    def test_entered_followup_instruction_is_not_discarded_if_not_iso_date(self):
        review = dict(REVIEW, fields={"follow_up_date": "At the next coaching meeting"})
        self.assertIn("At the next coaching meeting", document_text(build_ccc_letter(RESIDENT, review)))
        self.assertIn("At the next coaching meeting", document_text(report([], review)))


if __name__ == "__main__":
    unittest.main()
