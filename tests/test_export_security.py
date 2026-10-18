"""Synthetic Word encryption checks; no exported resident data or live services."""

import io
import zipfile

import msoffcrypto
import olefile
import pytest
from docx import Document
from msoffcrypto.exceptions import InvalidKeyError

from ccc import export_security


@pytest.fixture(scope="module")
def plain_docx():
    document = Document()
    document.add_heading("Synthetic CCC report", level=1)
    document.add_paragraph("Jordan Example: synthetic learning plan for encryption verification.")
    document.core_properties.title = "Synthetic confidential metadata marker"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture(scope="module")
def protected(plain_docx):
    return export_security.protect_docx(plain_docx, "Synthetic separate export passphrase 789!")


def test_protected_word_is_encrypted_and_roundtrips(plain_docx, protected):
    assert protected.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    assert not zipfile.is_zipfile(io.BytesIO(protected))
    office = msoffcrypto.OfficeFile(io.BytesIO(protected))
    assert office.is_encrypted()
    office.load_key(password="Synthetic separate export passphrase 789!", verify_password=True)
    decrypted = io.BytesIO()
    office.decrypt(decrypted, verify_integrity=True)
    assert decrypted.getvalue() == plain_docx
    document = Document(io.BytesIO(decrypted.getvalue()))
    assert "Jordan Example" in document.paragraphs[1].text


def test_report_text_metadata_and_password_do_not_appear_in_protected_bytes(protected):
    for marker in ("Jordan Example", "Synthetic confidential metadata marker",
                   "Synthetic separate export passphrase 789!"):
        assert marker.encode("utf-8") not in protected
        assert marker.encode("utf-16le") not in protected


def test_wrong_password_is_rejected(protected):
    office = msoffcrypto.OfficeFile(io.BytesIO(protected))
    with pytest.raises(InvalidKeyError):
        office.load_key(password="The wrong synthetic export password", verify_password=True)


def test_password_protection_uses_fresh_randomness(plain_docx, protected):
    other = export_security.protect_docx(plain_docx, "Synthetic separate export passphrase 789!")
    assert other != protected


def test_modified_encrypted_payload_fails_integrity_check(protected):
    tampered = io.BytesIO(protected)
    with olefile.OleFileIO(tampered, write_mode=True) as compound:
        payload = bytearray(compound.openstream("EncryptedPackage").read())
        payload[200] ^= 1
        compound.write_stream("EncryptedPackage", bytes(payload))
    office = msoffcrypto.OfficeFile(io.BytesIO(tampered.getvalue()))
    office.load_key(password="Synthetic separate export passphrase 789!", verify_password=True)
    with pytest.raises(InvalidKeyError):
        office.decrypt(io.BytesIO(), verify_integrity=True)


@pytest.mark.parametrize("password", ["", "short", " " * 14, "x" * 1025, None])
def test_bad_export_passphrases_are_rejected(plain_docx, password):
    with pytest.raises(export_security.ExportSecurityError):
        export_security.protect_docx(plain_docx, password)


def test_invalid_word_input_is_rejected():
    with pytest.raises(export_security.ExportSecurityError):
        export_security.protect_docx(b"not a Word report", "A valid synthetic passphrase")


def test_encryption_failure_cannot_fall_back_to_readable_word(plain_docx, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("Synthetic private document details")

    monkeypatch.setattr(export_security.OOXMLFile, "encrypt", fail)
    with pytest.raises(export_security.ExportSecurityError) as caught:
        export_security.protect_docx(plain_docx, "A valid synthetic passphrase")
    assert "Synthetic private document details" not in str(caught.value)
    assert caught.value.__suppress_context__


def test_failed_post_encryption_verification_releases_nothing(plain_docx, monkeypatch):
    original = export_security.msoffcrypto.OfficeFile

    def invalid_check(stream):
        check = original(stream)
        check.decrypt = lambda output, **kwargs: output.write(b"unexpected contents")
        return check

    monkeypatch.setattr(export_security.msoffcrypto, "OfficeFile", invalid_check)
    with pytest.raises(export_security.ExportSecurityError, match="could not be verified"):
        export_security.protect_docx(plain_docx, "A valid synthetic passphrase")
