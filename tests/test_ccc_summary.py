"""Template and data-boundary checks using fictional residents only."""

from __future__ import annotations

import copy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from ccc.ccc_summary import build_ccc_summary


@pytest.fixture
def resident():
    return {"id": "resident-example", "name": "Alex Example", "start_date": "2022-07-01",
            "graduation_date": "2026-06-30", "coach_name": "Example Coach", "status": "Active"}


def record(label, category="evaluations", specialty="Medicine", observed="2026-01-15", **fields):
    result = {"id": label, "resident_id": "resident-example", "source_id": "fictional-source",
              "category": category, "specialty": specialty, "date": observed, "record_mode": "event",
              "form": "Fictional form", "metric": "Fictional metric", "value": None,
              "rotation": "Example rotation", "evaluator": "Fictional Faculty", "text": "",
              "questions": [], "source_locator": "Fictional source row"}
    result.update(fields)
    return result


def review(**fields):
    result = {"id": "review-example", "resident_id": "resident-example", "pgy": 4, "half": "Year end",
              "academic_year": "2025-2026", "period_start": "2025-07-01", "period_end": "2026-06-30",
              "review_date": "2026-06-30", "status": "Draft", "fields": {"ccc_decision": "Fictional entered decision."}}
    result.update(fields)
    return result


def annual(pgy, academic, **fields):
    year = int(academic.split("-")[0])
    return {"id": f"annual-{pgy}-{academic}", "resident_id": "resident-example", "pgy": pgy,
            "academic_year": academic, "period_start": f"{year}-07-01", "period_end": f"{year+1}-06-30",
            "as_of": f"{year+1}-06-30", "fields": fields}


def build(resident, records=None, reviews=None, **kwargs):
    return build_ccc_summary(resident, records or [], reviews or [],
                             period_start=kwargs.pop("period_start", "2025-07-01"),
                             period_end=kwargs.pop("period_end", "2026-06-30"), **kwargs)


def document_text(doc):
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in {id(c._tc): c for c in row.cells}.values())
    return "\n".join(texts)


def named_table(doc, heading):
    return next(table for table in doc.tables if table.cell(0, 0).text == heading)


def test_reference_structure_fonts_colors_widths_and_editable_docx(resident):
    med = record("med", text="Fictional Medicine narrative.")
    ped = record("ped", specialty="Pediatrics", text="Fictional Pediatrics narrative.")
    payload = build(resident, [med, ped], [review()])
    assert payload.startswith(b"PK")
    doc = Document(BytesIO(payload))
    assert len(doc.tables) == 27
    assert len(doc.sections) == 1
    section = doc.sections[0]
    assert section.page_width.inches == 8.5 and section.page_height.inches == 11
    assert all(getattr(section, edge + "_margin").inches == 0.2 for edge in ("top", "bottom", "left", "right"))
    titles = [table.cell(0, 0).text for table in doc.tables]
    assert titles == [
        "MEDICINE PEDIATRICS END-OF-YEAR SUMMARY", "ROTATIONS SINCE LAST REVIEW",
        "ADULT AND PEDIATRIC CLINIC PATIENT TALLY", "PROCEDURES LOGGED TO DATE", "PROCEDURE",
        "CONFERENCE ATTENDANCE", "CATEGORY", "DUTY HOURS SUMMARY", "DUTY HOUR DEEP DIVE",
        "MEDICINE MILESTONES", "MEDICINE MILESTONE TREND SUMMARY", "PEDIATRIC MILESTONES",
        "PEDIATRIC MILESTONE TREND SUMMARY", "PRACTICE READINESS AND DEVELOPMENT SNAPSHOT",
        "OBSERVED ENCOUNTERS COMPLETION", "IN-TRAINING EXAM SCORES", "SCHOLARLY ACTIVITIES",
        "ADVOCACY", "360 EVALUATION FROM NURSING", "OVERALL IMPRESSION", "AREAS FOR GROWTH",
        "PATTERNS OVER TIME", "MEDICINE PEDIATRICS CLINICAL COMPETENCY COMMITTEE REVIEW",
        "SIGNATURES", "Resident Signature:", "MEDICINE EVALUATION SUMMARY", "PEDIATRIC EVALUATION SUMMARY",
    ]
    assert [column.width.inches for column in doc.tables[0].columns] == [1.8, 2.0, 1.8, 2.0]
    assert len(doc.tables[9].columns) == 22
    assert len(doc.tables[11].columns) == 23
    assert named_table(doc, "OVERALL IMPRESSION").cell(0, 0).paragraphs[0].paragraph_format.keep_with_next is True
    assert all(cell.paragraphs[0].paragraph_format.keep_with_next is True
               for row in doc.tables[9].rows[:-1] for cell in row.cells)
    for title, fill in (("OVERALL IMPRESSION", "BFEDD2"), ("AREAS FOR GROWTH", "F4CCCC"),
                        ("MEDICINE EVALUATION SUMMARY", "D9EAF7"), ("PEDIATRIC EVALUATION SUMMARY", "BFEDD2")):
        cell = named_table(doc, title).cell(0, 0)
        assert cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == fill
    for title in ("MEDICINE EVALUATION SUMMARY", "PEDIATRIC EVALUATION SUMMARY"):
        table = named_table(doc, title)
        assert len(table.columns) == 10 and len(table.rows) == 3
        assert [int(tc.tcPr.gridSpan.val) for tc in table.rows[1]._tr.tc_lst] == [2, 2, 6]
        assert int(table.rows[2]._tr.tc_lst[0].tcPr.gridSpan.val) == 10
        assert all(run.font.name == "Arial" and run.font.size.pt == 6
                   for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs if run.text)
    with ZipFile(BytesIO(payload)) as archive:
        xml = archive.read("word/document.xml").decode()
        assert xml.count('w:type="page"') == 2
        assert "documentProtection" not in archive.read("word/settings.xml").decode()
        assert "Fictional Medicine narrative." in xml and "Fictional Pediatrics narrative." in xml


def test_live_evidence_uses_current_range_even_when_latest_narrative_review_is_final(resident):
    old = review(period_end="2025-12-31", review_date="2026-01-10", status="Final",
                 snapshot=[], fields={"overall_impression": "Fictional historical narrative."})
    records = [record("later", observed="2026-03-20", text="Later current evidence."),
               record("future", observed="2026-07-01", text="Future evidence excluded."),
               record("other", resident_id="someone-else", text="Other resident excluded.")]
    text = document_text(Document(BytesIO(build(resident, records, [old]))))
    assert "Later current evidence." in text and "Fictional historical narrative." in text
    assert "Future evidence excluded." not in text and "Other resident excluded." not in text
    assert "2025-07-01 to 2026-06-30" in text


def test_explicit_final_uses_all_frozen_sources_and_profile_without_live_leaks(resident):
    saved = review(status="Final", snapshot=[record("frozen", text="Frozen evaluation.")],
                   cumulative_evidence_snapshot=[record("old-procedure", "procedures", observed="2023-06-01", metric="Frozen procedure")],
                   annual_progress_snapshot=[annual(4, "2025-2026", adult_clinic_count=0, pediatric_clinic_count=None)],
                   snapshot_resident={**resident, "name": "Alex Frozen Example", "coach_name": "Frozen coach"},
                   history_snapshot=[])
    before = copy.deepcopy(saved)
    payload = build(resident, [record("live", text="LIVE EVIDENCE MUST NOT LEAK")], [review(fields={"abstracts": "LIVE ABSTRACT"})],
                    review=saved, tracking=[annual(4, "2025-2026", adult_clinic_count=999)],
                    period_start="2026-04-01", period_end="2026-05-01")
    doc = Document(BytesIO(payload))
    text = document_text(doc)
    assert "Frozen evaluation." in text and "Frozen procedure" in text
    assert "Alex Frozen Example" in text and "Frozen coach" in text
    assert "LIVE EVIDENCE MUST NOT LEAK" not in text and "LIVE ABSTRACT" not in text and "999" not in text
    assert "2025-07-01 to 2026-06-30" in text
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    assert clinic.rows[-1].cells[1].text == "0" and clinic.rows[-1].cells[3].text == ""
    assert saved == before


def test_old_final_never_fills_missing_cumulative_or_annual_snapshots_from_live_data(resident):
    saved = review(status="Final", snapshot=[record("saved", text="Preserved comment.")],
                   fields={"medicine_ite_score": "70", "medicine_ite_percentile": "60"},
                   snapshot_resident=copy.deepcopy(resident), history_snapshot=[])
    text = document_text(Document(BytesIO(build(resident,
        [record("new", "procedures", metric="NEW PROCEDURE")], [review(fields={"publications": "NEW PUBLICATION"})],
        review=saved, tracking=[annual(4, "2025-2026", medicine_ite_score=999)]))))
    assert "Preserved comment." in text and "70 (60%tile)" in text
    assert "Historical data limited to this saved review" in text
    assert "NEW PROCEDURE" not in text and "NEW PUBLICATION" not in text and "999" not in text


@pytest.mark.parametrize("bad_review, message", [
    (review(status="Final"), "no preserved evidence"),
    (review(resident_id="other"), "another resident"),
])
def test_review_boundary_errors_fail_closed(resident, bad_review, message):
    with pytest.raises(ValueError, match=message):
        build(resident, review=bad_review)


def test_annual_values_are_dated_authoritative_and_distinguish_pediatrics_measure(resident):
    legacy = review(fields={"adult_clinic_count": "999", "pediatrics_ite_score": "60", "pediatrics_ite_percentile": "80"})
    entries = [annual(4, "2025-2026", adult_clinic_count=None, pediatric_clinic_count=0,
                       medicine_ite_score=72, medicine_ite_percentile=65,
                       pediatrics_ite_score=70, pediatrics_ite_expected_pass_rate=92,
                       observed_encounter_1=True, observed_encounter_2=False, observed_encounter_3=None)]
    doc = Document(BytesIO(build(resident, reviews=[legacy], tracking=entries)))
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    assert clinic.rows[-1].cells[1].text == "" and clinic.rows[-1].cells[3].text == "0"
    text = document_text(doc)
    assert "72 (65%tile)" in text and "70 (92%_PR)" in text
    assert "999" not in text and "80%tile" not in text
    encounters = named_table(doc, "OBSERVED ENCOUNTERS COMPLETION")
    assert encounters.rows[2].cells[-1].text == "X"
    assert encounters.rows[3].cells[-1].text == "Not completed"
    assert len(encounters.rows) == 4


def test_legacy_pediatrics_percentile_is_not_relabelled_expected_pass_rate(resident):
    old = review(fields={"pediatrics_ite_score": "66", "pediatrics_ite_percentile": "88"})
    text = document_text(Document(BytesIO(build(resident, reviews=[old]))))
    assert "66 (88%tile)" in text and "expected pass rate" not in text


def test_repeated_pgy_retains_separate_dated_annual_columns(resident):
    entries = [annual(2, "2023-2024", adult_clinic_count=11, observed_encounter_1=True, medicine_ite_score=51),
               annual(2, "2024-2025", adult_clinic_count=22, observed_encounter_1=False, medicine_ite_score=62)]
    checkpoint = review(pgy=2, academic_year="2024-2025", period_start="2024-07-01",
                        period_end="2025-06-30", review_date="2025-06-30")
    doc = Document(BytesIO(build(resident, reviews=[checkpoint], tracking=entries)))
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    text = "\n".join(cell.text for row in clinic.rows for cell in row.cells)
    assert "2023-2024" in text and "2024-2025" in text and "11" in text and "22" in text
    encounters = named_table(doc, "OBSERVED ENCOUNTERS COMPLETION")
    assert len(encounters.columns) == 6
    assert "2023-2024" in encounters.rows[1].cells[2].text
    assert "2024-2025" in encounters.rows[1].cells[3].text


def test_milestone_matrix_keeps_zero_missing_status_and_specialties_separate(resident):
    records = [record("med-zero", "milestones", observed="2025-09-01", metric="PC1", value=0),
               record("med-null", "milestones", observed="2025-09-01", metric="PC2", values={"assessment_status": "Not assessed"}),
               record("med-next", "milestones", observed="2026-02-01", metric="PC1", value=2),
               record("ped-score", "milestones", specialty="Pediatrics", observed="2026-02-01", metric="PC1", value=4.5)]
    doc = Document(BytesIO(build(resident, records, [review()])))
    med = named_table(doc, "MEDICINE MILESTONES")
    ped = named_table(doc, "PEDIATRIC MILESTONES")
    assert med.rows[3].cells[1].text == "0.00" and med.rows[3].cells[2].text == "Not assessed"
    assert med.rows[4].cells[1].text == "2.00" and ped.rows[3].cells[1].text == "4.50"
    trend = named_table(doc, "MEDICINE MILESTONE TREND SUMMARY")
    assert trend.rows[2].cells[-2].text == "+2.00" and trend.rows[2].cells[-1].text == "Increase"
    readiness = named_table(doc, "PRACTICE READINESS AND DEVELOPMENT SNAPSHOT")
    assert all(row.cells[1].text == "" for row in readiness.rows[1:])
    text = document_text(doc)
    assert "Exemplary" not in text and "Ready" not in text and "Strong upward progression" not in text


def test_comments_are_uncapped_and_combined_self_assessment_have_distinct_pages(resident):
    long_comment = "Fictional long comment. " * 1500 + "END OF FULL COMMENT"
    rows = [record("med", questions=[{"number": "1", "text": "Fictional question?", "response": "Progressing", "comment": long_comment}]),
            record("combined", specialty="Combined", text="Combined evidence retained."),
            record("self", "self_assessments", "Combined", text="Self reflection retained.")]
    doc = Document(BytesIO(build(resident, rows)))
    text = document_text(doc)
    assert long_comment in text and "Fictional question?" in text and "Progressing" in text
    assert "COMBINED EVALUATION SUMMARY" in text and "Combined evidence retained." in text
    assert "RESIDENT SELF-ASSESSMENT" in text and "Self reflection retained." in text
    assert len(doc.element.xpath('.//w:br[@w:type="page"]')) == 4


def test_conference_snapshots_keep_actual_periods_and_are_not_summed(resident):
    records = [record("first", "conferences", observed="2025-12-31", value=40, metric="Fictional conference", record_mode="snapshot",
                      unit="%", period_start="2025-07-01", period_end="2025-12-31"),
               record("second", "conferences", observed="2026-06-30", value=60, metric="Fictional conference", record_mode="snapshot",
                      unit="%", period_start="2025-07-01", period_end="2026-06-30")]
    doc = Document(BytesIO(build(resident, records, [review()])))
    table = named_table(doc, "CATEGORY")
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    assert "40.00%" in text and "60.00%" in text and "100.00%" not in text
    assert "2025-07-01 to 2025-12-31" in text


def test_procedure_counts_and_duty_notes_use_saved_records_without_compliance_inference(resident):
    records = [record("procedure-one", "procedures", observed="2023-01-01", metric="Fictional procedure"),
               record("procedure-two", "procedures", observed="2026-03-01", metric="Fictional procedure"),
               record("shift", "duty_hours", observed="2026-05-01", value=10),
               record("violation", "duty_violations", metric="Source reported rule", text="Source explanation.",
                      values={"description": "Source explanation.", "comments": "Source comments retained."})]
    saved = review(fields={"duty_resolution_notes": "Committee entered resolution.", "readiness_communication": "Entered assessment."})
    doc = Document(BytesIO(build(resident, records, [saved])))
    procedures = named_table(doc, "PROCEDURE")
    assert procedures.rows[1].cells[-1].text == "2"
    text = document_text(doc)
    assert "Source comments retained." in text and "Committee entered resolution." in text and "Entered assessment." in text
    deep = named_table(doc, "DUTY HOUR DEEP DIVE")
    assert deep.rows[1].cells[1].text == "2026-05-01"
    assert all(cell.text == "" for row in (deep.rows[3], deep.rows[5]) for cell in row.cells)


def test_future_tracking_and_narrative_are_excluded_and_inputs_unchanged(resident):
    entries = [annual(4, "2025-2026", adult_clinic_count=10), annual(4, "2026-2027", adult_clinic_count=999)]
    reviews = [review(), review(id="future", review_date="2027-01-01", fields={"overall_impression": "Future narrative."})]
    before = copy.deepcopy((resident, reviews, entries))
    text = document_text(Document(BytesIO(build(resident, reviews=reviews, tracking=entries))))
    assert "Future narrative." not in text and "999" not in text
    assert (resident, reviews, entries) == before


def test_distinct_periods_in_same_pgy_and_academic_year_are_preserved(resident):
    first = annual(4, "2025-2026", adult_clinic_count=11, medicine_ite_score=55, observed_encounter_1=True)
    first.update(period_end="2025-12-31", as_of="2025-12-31")
    second = annual(4, "2025-2026", adult_clinic_count=22, medicine_ite_score=66, observed_encounter_1=False)
    second.update(id="annual-second-half", period_start="2026-01-01")
    doc = Document(BytesIO(build(resident, tracking=[first, second])))
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    fourth_year = [row for row in clinic.rows if row.cells[0].text.startswith("ADULT PGY4")]
    assert len(fourth_year) == 2
    assert [row.cells[1].text for row in fourth_year] == ["11", "22"]
    assert "2025-07-01 to 2025-12-31" in fourth_year[0].cells[0].text
    assert "2026-01-01 to 2026-06-30" in fourth_year[1].cells[0].text
    encounters = named_table(doc, "OBSERVED ENCOUNTERS COMPLETION")
    assert len(encounters.columns) == 6


def test_same_start_but_different_annual_end_dates_are_not_collapsed(resident):
    first = annual(4, "2025-2026", pediatric_clinic_count=13)
    first.update(period_end="2025-12-31", as_of="2025-12-31")
    second = annual(4, "2025-2026", pediatric_clinic_count=26)
    doc = Document(BytesIO(build(resident, tracking=[first, second])))
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    fourth_year = [row for row in clinic.rows if row.cells[0].text.startswith("ADULT PGY4")]
    assert len(fourth_year) == 2
    assert [row.cells[3].text for row in fourth_year] == ["13", "26"]


def test_live_header_and_year_defaults_use_cutoff_pgy_after_old_review(resident):
    old = review(pgy=1, academic_year="2022-2023", period_start="2022-07-01", period_end="2023-06-30",
                 review_date="2023-06-30", fields={"overall_impression": "Prior reviewed narrative."})
    doc = Document(BytesIO(build(resident, reviews=[old], period_start="2024-07-01", period_end="2024-12-31")))
    assert doc.tables[0].rows[1].cells[3].text == "R3"
    clinic = named_table(doc, "ADULT AND PEDIATRIC CLINIC PATIENT TALLY")
    assert len(clinic.rows) == 4 and clinic.rows[-1].cells[0].text == "ADULT PGY3"
    assert "Prior reviewed narrative." in document_text(doc)


def test_final_missing_preserved_resident_profile_fails_with_clear_validation_error(resident):
    from ccc.model import ValidationError
    saved = review(status="Final", snapshot=[])
    with pytest.raises(ValidationError, match="preserved resident details"):
        build(resident, review=saved)


def test_conference_source_units_are_preserved_instead_of_inventing_percentages(resident):
    evidence = record("sessions", "conferences", metric="Fictional teaching conference", value=4, unit="sessions",
                      record_mode="snapshot", period_start="2025-07-01", period_end="2026-01-15")
    text = document_text(Document(BytesIO(build(resident, [evidence]))))
    assert "4 sessions" in text and "4.00%" not in text
